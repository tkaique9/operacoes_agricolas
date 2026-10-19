"""Atualiza a reunião semanal de Tratos Culturais (index.html).

Coloque na pasta dados/ as duas planilhas da semana e rode:

    python atualizar_tratos.py

- "Operações Tratos Culturais - 2025-2026" (.xlsb): aba Semanal-Acumulado, com o
  realizado por dia de cada atividade, as metas e a semana de referência.
- "Operações Tratos Culturais - SAFRA 2026" (.xlsx): as abas de cada programa
  (pré-emergente, catações, 3ª molécula, adubação), com o realizado e o saldo a
  realizar por fazenda e por bloco.

O script acha cada planilha pelas abas que ela tem, não pelo nome do arquivo.
Grava window.TRATOS entre os marcadores DADOS-INICIO e DADOS-FIM do index.html;
todo o resto (totais, percentuais, gráficos) é calculado no navegador.
Precisa de: pip install openpyxl pyxlsb
"""
import datetime, glob, json, os, re, sys
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(AQUI, "index.html")
EPOCA = datetime.datetime(1899, 12, 30)

# nome curto de cada frente (o bloco da aba Semanal-Acumulado)
FRENTES = [
    ("APLICAÇÃO MANUAL", "manual", "Manual e semi-mecanizada"),
    ("UNIPORT", "uniport", "Uniport"),
    ("AÉREA", "aerea", "Aérea e monitoramento"),
    ("QUEBRA LOMBO", "lombo", "Quebra lombo"),
    ("COCHO", "cocho", "Adubação com cocho"),
]
# frentes que não entram na apresentação (nem nas páginas, nem nos totais). Quebra
# lombo saiu a pedido em 13/10/2026: não é aplicação desta operação. Para voltar,
# tire o id daqui.
FRENTES_FORA = {"lombo"}
# meta de hectares por pessoa/dia que vale no lugar da que vem da planilha. Cata Mato:
# a linha "Pessoas Dia" traz 8 ha (120 ha ÷ 15 pessoas), mas a meta é 4 ha/pessoa
# (corrigido em 13/10/2026). Para voltar à da planilha, tire a atividade daqui.
META_PESSOA = {"Cata Mato": 4.0}
# nº de equipamentos que vale no lugar do da planilha; a meta do dia, da semana e do
# mês acompanham (equipamentos × meta por equipamento). Quadriciclo: a planilha traz
# 1, mas operam 2, com meta de 10 ha cada: 20 ha/dia (corrigido em 13/10/2026).
# Catação de Formiga: 2 equipamentos de 10 ha cada (a planilha traz 1 de 20 ha), mesma
# meta de 20 ha/dia (corrigido em 13/10/2026).
EQUIP = {"Quadriciclo": 2, "Catação de Formiga": 2}
# meta de hectares por equipamento/dia que vale no lugar da planilha; aplicada antes
# de EQUIP, que refaz a meta do dia com ela
META_EQUIP = {"Catação de Formiga": 10.0}
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


def s(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def data_xl(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, (int, float)) and 40000 < v < 60000:
        return EPOCA + datetime.timedelta(days=float(v))
    return None


def iso(d):
    return d.strftime("%Y-%m-%d")


def r2(v):
    return None if v is None else round(v, 2)


# ------------------------------------------------------------------ leitura
def le_xlsb(path):
    from pyxlsb import open_workbook
    out = {}
    with open_workbook(path) as wb:
        for nome in wb.sheets:
            with wb.get_sheet(nome) as sh:
                out[nome] = [[c.v for c in r] for r in sh.rows()]
    return out


def le_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def le(path):
    return le_xlsb(path) if path.lower().endswith(".xlsb") else le_xlsx(path)


def cel(linha, j):
    return linha[j] if j < len(linha) else None


# ------------------------------------------------------- semana (xlsb)
def semana(ws):
    """Aba Semanal-Acumulado: atividades, metas e realizado por dia."""
    # cabeçalho com as datas (linha 'Atividade') e o quadro de resumo à direita
    hdr_i = next(i for i, r in enumerate(ws) if s(cel(r, 2)) == "Atividade")
    hdr = ws[hdr_i]
    colData = {j: data_xl(v) for j, v in enumerate(hdr) if data_xl(v)}
    rot = {s(v): j for j, v in enumerate(hdr) if isinstance(v, str)}
    C = {
        "equip": rot.get("N° Equipamentos/Pessoas"), "metaUnid": rot.get("Meta - Equipamentos/Pessoas"),
        "metaDia": rot.get("Meta Dia"), "metaSemana": rot.get("Meta Semana"), "metaMes": rot.get("Meta Mês"),
    }
    # dia de referência e quantidade de dias da semana (quadro "DIA REF" / "QTD DIAS")
    dia_ref = qtd = None
    for r in ws[:hdr_i]:
        for j, v in enumerate(r):
            if s(v) == "DIA REF:":
                dia_ref = data_xl(cel(r, j + 1))
            if s(v) == "QTD DIAS":
                qtd = int(num(cel(r, j + 1)) or 0)
    if not dia_ref:
        sys.exit("Semanal-Acumulado: não achei o DIA REF")
    seg = dia_ref - datetime.timedelta(days=dia_ref.weekday())
    semana_toda = [seg + datetime.timedelta(days=k) for k in range(7)]
    col_de = {iso(d): j for j, d in colData.items()}

    def valores(r, ds):
        return [num(cel(r, col_de.get(iso(d), -1))) if iso(d) in col_de else None for d in ds]

    frentes, atual, ativs = [], None, []
    linhas_pess = []      # (índice da atividade dona, pessoas[], rendimento[])
    for i in range(hdr_i + 1, len(ws)):
        r = ws[i]
        nome = s(cel(r, 2))
        if not nome or nome == "Atividade" or re.match(r"^\d{2}/\d{2}/\d{4} a", nome) or nome.startswith("Observação"):
            continue
        fr = next((f for f in FRENTES if f[0] in nome.upper()), None)
        if fr and nome.upper() == nome:                      # cabeçalho de frente
            atual = {"id": fr[1], "nome": fr[2], "titulo": nome, "atividades": []}
            frentes.append(atual)
            continue
        if atual is None:
            continue
        if nome in ("Pessoas Dia", "Rendimento H/Dia"):
            if nome == "Pessoas Dia":
                rend = ws[i + 1] if i + 1 < len(ws) and s(cel(ws[i + 1], 2)).startswith("Rendimento") else None
                linhas_pess.append((atual, r, rend))
            continue
        a = {
            "nome": nome, "linha": i + 1,
            "equip": num(cel(r, C["equip"])) if C["equip"] is not None else None,
            "metaUnid": num(cel(r, C["metaUnid"])) if C["metaUnid"] is not None else None,
            "metaDia": num(cel(r, C["metaDia"])) or 0.0,
            "metaSemana": num(cel(r, C["metaSemana"])) or 0.0,
            "metaMes": num(cel(r, C["metaMes"])) or 0.0,
            "dias": valores(r, semana_toda),
            "ant": valores(r, [d - datetime.timedelta(days=7) for d in semana_toda]), "pessoas": None,
            "pessoasPlan": None, "metaPessoa": None,
            "_r": r,
        }
        if nome in META_EQUIP:
            a["metaUnid"] = META_EQUIP[nome]
        if nome in EQUIP and a["metaDia"] and a["metaUnid"]:
            fator = EQUIP[nome] * a["metaUnid"] / a["metaDia"]
            a["equip"] = float(EQUIP[nome])
            for k in ("metaDia", "metaSemana", "metaMes"):
                a[k] *= fator
        atual["atividades"].append(a)

    # os valores vêm como a planilha soma: a semana vai de segunda ao DIA REF e leva também
    # os dias seguintes da semana que tiveram produção (o domingo trabalhado entra na
    # semana e no mês, sem meta: a meta da semana segue a da planilha)
    fim = max(k for k, d in enumerate(semana_toda)
              if d <= dia_ref or any((a["dias"][k] or 0) > 0 for fr in frentes if fr["id"] not in FRENTES_FORA
                                     for a in fr["atividades"]))
    dias = semana_toda[:fim + 1]
    dias_ant = [d - datetime.timedelta(days=7) for d in dias]
    for fr in frentes:
        for a in fr["atividades"]:
            a["dias"], a["ant"] = a["dias"][:fim + 1], a["ant"][:fim + 1]

    # pessoas/dia: a linha pertence à atividade da mesma frente cujo realizado
    # dividido pelas pessoas bate com a linha de rendimento logo abaixo
    avisos = []
    for fr, rp, rr in linhas_pess:
        pess = valores(rp, dias)
        rend = valores(rr, dias) if rr else [None] * len(dias)
        melhor, erro_min, invertido = None, None, False
        for a in fr["atividades"]:
            e1, e2 = [], []
            for v, p, rd in zip(a["dias"], pess, rend):
                if v and p and rd:
                    e1.append(abs(v / p - rd))          # ha por pessoa (o certo)
                    e2.append(abs(p / v - rd))          # pessoas por ha (fórmula invertida)
            if e1:
                e = min(sum(e1), sum(e2)) / len(e1)
                if erro_min is None or e < erro_min:
                    melhor, erro_min, invertido = a, e, sum(e2) < sum(e1)
        if melhor is None:      # sem rendimento para conferir: a atividade com valor mais próxima acima
            melhor = next((a for a in reversed(fr["atividades"]) if any(a["dias"])), None)
        if melhor is not None:
            melhor["pessoas"] = pess
            # planejado de pessoas: a linha "Pessoas Dia" pode trazer o próprio quadro de
            # meta (Cata Mato: 15 pessoas); sem ele, vale o da atividade. META_PESSOA
            # tem a palavra final sobre a meta por pessoa
            pp, pm = num(cel(rp, C["equip"])), num(cel(rp, C["metaUnid"]))
            melhor["pessoasPlan"] = r2(pp) if pp else melhor["equip"]
            melhor["metaPessoa"] = r2(pm) if pm else melhor["metaUnid"]
            if melhor["nome"] in META_PESSOA:
                melhor["metaPessoa"] = META_PESSOA[melhor["nome"]]
            if melhor["nome"] in EQUIP:
                melhor["pessoasPlan"] = float(EQUIP[melhor["nome"]])
            if invertido:
                avisos.append(f"{melhor['nome']}: o \"Rendimento H/Dia\" da planilha está como pessoas ÷ ha; "
                              f"na apresentação o rendimento é ha ÷ pessoa.")

    # dias em que uma atividade repete a linha de pessoas ou de rendimento: é linha
    # colada fora do lugar na planilha. Esses valores saem da conta e viram aviso.
    refs = []
    for _, rp, rr in linhas_pess:
        refs.append(rp)
        if rr:
            refs.append(rr)
    meses = sorted({(d.year, d.month) for d in colData.values() if d <= dias[-1]})
    for fr in frentes:
        for a in fr["atividades"]:
            r = a.pop("_r")
            ruins = set()
            for ref in refs:
                if ref is r:
                    continue
                seq = []
                for j in sorted(colData):
                    v, w = num(cel(r, j)), num(cel(ref, j))
                    if v and w and abs(v - w) < 1e-9:
                        seq.append(j)
                    else:
                        if len(seq) >= 3:
                            ruins.update(seq)
                        seq = []
                if len(seq) >= 3:
                    ruins.update(seq)
            if ruins:
                ds = sorted(colData[j] for j in ruins)
                avisos.append(f"{a['nome']}: de {ds[0]:%d/%m} a {ds[-1]:%d/%m} a linha repete pessoas/rendimento "
                              f"de outra atividade ({len(ds)} dias) — valores fora da conta, conferir a planilha.")
                for k, d in enumerate(dias):
                    if col_de.get(iso(d)) in ruins:
                        a["dias"][k] = None
                for k, d in enumerate(dias_ant):
                    if col_de.get(iso(d)) in ruins:
                        a["ant"][k] = None
            mes = {}
            for (y, m) in meses:
                tot = sum(num(cel(r, j)) or 0 for j, d in colData.items()
                          if d.year == y and d.month == m and d <= dias[-1] and j not in ruins)
                mes[f"{y}-{m:02d}"] = round(tot, 2)
            a["mes"] = mes
            a["dias"] = [r2(v) for v in a["dias"]]
            a["ant"] = [r2(v) for v in a["ant"]]
            if a["pessoas"]:
                a["pessoas"] = [r2(v) for v in a["pessoas"]]
            for k in ("equip", "metaUnid", "metaDia", "metaSemana", "metaMes"):
                a[k] = r2(a[k])
    frentes = [fr for fr in frentes if fr["id"] not in FRENTES_FORA]
    # atividade sem meta e sem nada realizado no ano não entra
    for fr in frentes:
        fr["atividades"] = [a for a in fr["atividades"] if a["metaSemana"] or any(a["mes"].values())]
    return {
        "diaRef": iso(dia_ref), "qtdDias": qtd or len(dias),
        "dias": [iso(d) for d in dias], "diasAnt": [iso(d) for d in dias_ant],
        "meses": [f"{y}-{m:02d}" for (y, m) in meses],
    }, frentes, avisos


# ---------------------------------------------------- programas (xlsx)
def tabela_fazendas(ws, cab_req):
    """Linhas de fazenda entre o cabeçalho (que tem as colunas pedidas) e o 'Total'."""
    for i, r in enumerate(ws):
        rot = {s(v): j for j, v in enumerate(r) if isinstance(v, str)}
        if all(c in rot for c in cab_req):
            lin = []
            for r2_ in ws[i + 1:]:
                if s(cel(r2_, 0)) in ("Total", "TOTAL", "TOTAL GERAL"):
                    break
                cod = num(cel(r2_, rot["Cód Faz"]))
                if cod and 9000 <= cod < 10000:
                    lin.append(r2_)
            return rot, lin, i
    return None, [], None


# o mesmo bloco vem escrito de jeitos diferentes em cada aba ("SHEKINHA-I",
# "SHEKINAH l", "MARCO-AMUY", "BR-154", "SCALON 7 PLACAS"...): um nome só para todos
SINONIMOS_BLOCO = {"SHEKINHA I": "SHEKINAH I", "SHEKINAH L": "SHEKINAH I", "BR 154": "154",
                   "SCALON 7 PLACAS": "SCALON SETE PLACAS"}


def bloco_norm(b):
    b = re.sub(r"\s+", " ", s(b).upper().replace("-", " ")).strip()
    return SINONIMOS_BLOCO.get(b, b)


def junta(dic, k, v):
    if v:
        dic[k] = round(dic.get(k, 0) + v, 2)


def quadro_bloco(ws, rotulo="BLOCO"):
    """Quadro 'BLOCO | Á REALIZAR' (ou REGIÃO) que a equipe monta na própria aba."""
    for i, r in enumerate(ws):
        if s(cel(r, 0)) == rotulo and s(cel(r, 1)).upper().startswith("Á REALIZAR"):
            out = {}
            for r2_ in ws[i + 1:]:
                k, v = s(cel(r2_, 0)), num(cel(r2_, 1))
                if not k or v is None:
                    break
                junta(out, bloco_norm(k) if rotulo == "BLOCO" else k, v)
            return out
    return {}


def soma(lin, j):
    return round(sum(num(cel(r, j)) or 0 for r in lin), 2)


def por_bloco(lin, jb, jv):
    out = {}
    for r in lin:
        b, v = s(cel(r, jb)), num(cel(r, jv)) or 0
        if b and not b.startswith("#"):
            junta(out, bloco_norm(b), v)
    return out


def programas(wb):
    P = []
    # pré-emergente cana soca: base = área colhida; saldo pelo quadro BLOCO da aba
    ws = wb["Pré-Emerg -Cana Soca SAFRA"]
    rot, lin, _ = tabela_fazendas(ws, ["Cód Faz", "Bloco", "Área Colhida", "Área Realizada", "Á Realizar"])
    blocos = quadro_bloco(ws)
    P.append({"id": "pre", "nome": "Pré-emergente cana soca", "base": soma(lin, rot["Área Colhida"]), "baseRot": "área colhida",
              "reforma": soma(lin, rot["Área Reforma"]), "realizado": soma(lin, rot["Área Realizada"]),
              "aRealizar": round(sum(blocos.values()), 2), "blocos": blocos, "regioes": quadro_bloco(ws, "REGIÃO")})
    # pré-emergente área de queima
    ws = wb["Pré-Emerg - Área Queima"]
    rot, lin, _ = tabela_fazendas(ws, ["Cód Faz", "Bloco", "Área Queimada", "Área Realizada", "Á Realizar"])
    P.append({"id": "queima", "nome": "Pré-emergente área de queima", "base": soma(lin, rot["Área Queimada"]), "baseRot": "área queimada",
              "realizado": soma(lin, rot["Área Realizada"]), "aRealizar": soma(lin, rot["Á Realizar"]),
              "blocos": por_bloco(lin, rot["Bloco"], rot["Á Realizar"])})
    # 3ª molécula (reforço de bordadura)
    ws = wb["Terceira Mol. - Soca SAFRA"]
    rot, lin, _ = tabela_fazendas(ws, ["Cód Faz", "Bloco", "Área Aberta", "Área Realizada", "Á Realizar"])
    P.append({"id": "mol3", "nome": "3ª molécula (bordadura)", "base": soma(lin, rot["Área Aberta"]), "baseRot": "área aberta",
              "realizado": soma(lin, rot["Área Realizada"]), "aRealizar": soma(lin, rot["Á Realizar"]),
              "blocos": por_bloco(lin, rot["Bloco"], rot["Á Realizar"])})
    # 1ª catação química: saldo dividido em mais e menos de 20 dias após a colheita
    ws = wb["1ª Catação Química SAFRA"]
    rot, lin, _ = tabela_fazendas(ws, ["Cód Faz", "Bloco", "Área Aberta", "Realizado Total"])
    j20 = next(j for k, j in rot.items() if k.startswith("Á Realizar") and ">" in k)
    j20m = next(j for k, j in rot.items() if k.startswith("Á Realizar") and "<" in k)
    mais, menos = soma(lin, j20), soma(lin, j20m)
    bl = por_bloco(lin, rot["Bloco"], j20)
    for b, v in por_bloco(lin, rot["Bloco"], j20m).items():
        bl[b] = round(bl.get(b, 0) + v, 2)
    P.append({"id": "cat1", "nome": "1ª catação química", "base": soma(lin, rot["Área Aberta"]), "baseRot": "área aberta",
              "reforma": soma(lin, rot["Área Reforma"]), "realizado": soma(lin, rot["Realizado Total"]),
              "mec": soma(lin, rot["Área Realizada MEC"]), "man": soma(lin, rot["Área Realizada MAN"]),
              "aRealizar": round(mais + menos, 2), "atrasado": mais, "blocos": bl,
              "blocosAtrasado": por_bloco(lin, rot["Bloco"], j20)})
    # 2ª catação química
    ws = wb["2ª Catação Química SAFRA"]
    rot, lin, _ = tabela_fazendas(ws, ["Cód Faz", "Bloco", "Programado", "Área Realizada", "Á Realizar"])
    P.append({"id": "cat2", "nome": "2ª catação química", "base": soma(lin, rot["Programado"]), "baseRot": "programado",
              "reforma": soma(lin, rot["Área Reforma"]), "realizado": soma(lin, rot["Área Realizada"]),
              "aRealizar": soma(lin, rot["Á Realizar"]), "blocos": por_bloco(lin, rot["Bloco"], rot["Á Realizar"])})
    # adubação de soca: resumo por bloco da própria planilha (localizada e cocho)
    ws = wb["ABERTO-ENCERRADO_BLOCO"]
    for titulo, pid, nome in (("LOCALIZADA", "loc", "Adubação localizada"), ("COCHO", "cocho", "Adubação com cocho")):
        ini = next(i for i, r in enumerate(ws) if s(cel(r, 5)) == titulo)
        cab = next(i for i in range(ini, len(ws)) if s(cel(ws[i], 0)) == "BLOCO")
        bl, tot = {}, None
        for r in ws[cab + 1:]:
            b = s(cel(r, 0))
            if not b:
                break
            vals = [num(cel(r, j)) or 0 for j in range(1, 7)]   # aberta, encerrada, total, reforma, realizado, a realizar
            if b.startswith("TOTAL"):
                tot = vals
                break
            junta(bl, bloco_norm(b), vals[5])
        P.append({"id": pid, "nome": nome, "base": round(tot[2], 2), "baseRot": "área com ordem",
                  "reforma": round(tot[3], 2), "realizado": round(tot[4], 2), "aRealizar": round(tot[5], 2),
                  "blocos": {b: v for b, v in bl.items() if v}})
    return P


# ------------------------------------------ fazendas e talhões da semana
# Detalhe por fazenda e talhão, sempre só do que foi feito nos dias da semana:
# - relatório de operações do sistema (RSC0819R, .xlsx) que estiver em dados/: todas
#   as operações, com OS, área executada e dose de cada produto;
# - aba "Realizado - 3ª Mol" da planilha SAFRA: OS da 3ª molécula (Uniport
#   bordadura), com a área de borda aplicada em cada talhão (coluna "Recom.").
# A operação do sistema vai para a frente pela descrição; a atividade só quando não
# há dúvida (o sistema não separa Uniport próprio de terceiro nem as equipes do cocho).
OPER_FRENTE = [("DRONE", "aerea", "Drone Inseticida"), ("UNIPORT", "uniport", None),
               ("COCHO", "cocho", None), ("FORMIGA", "manual", "Catação de Formiga"),
               ("MANUAL", "manual", None)]
# rótulo curto das operações mais comuns (o relatório traz a descrição longa e truncada)
ROTULOS = {
    24565: "Broca · drone · 1ª aplic.", 24566: "Broca · drone · 2ª aplic.",
    24593: "Cigarrinha · drone · 1ª aplic.", 14501: "Formiga · manual · 1ª aplic.",
    24003: "Adubação cocho · 1ª aplic.", 24004: "Adubação cocho · 2ª aplic.",
    14330: "Herb. Uniport barra total · 1ª aplic.", 24330: "Herb. Uniport barra total · 1ª aplic.",
    24331: "Herb. Uniport barra total · 2ª aplic.", 14327: "Herb. Uniport bordadura · 1ª aplic.",
    24327: "Herb. Uniport bordadura · 1ª aplic.", 14340: "Herb. manual · 1ª aplic.",
    14341: "Herb. manual · 2ª aplic.", 14342: "Herb. manual · 3ª aplic.",
}
UNIDADES = {"LT": "L", "KG": "kg", "TON": "t", "UN": "un", "HA": "ha"}
MINUSCULAS = {"de", "do", "da", "dos", "das", "e", "ou", "d'"}


def nome_faz(fundo):
    """'9006 - FAZ. CÓRREGO DAS FLORES II' -> 'Faz. Córrego das Flores II'"""
    t = s(fundo)
    t = t.split(" - ", 1)[-1] if " - " in t else t
    t = re.sub(r"^FAZ\.(?=\S)", "FAZ. ", t, flags=re.I)
    out = []
    for i, w in enumerate(t.split(" ")):
        lw = w.lower()
        if re.fullmatch(r"l{1,3}", w) and i > 0:         # "SUCURI ll" = Sucuri II
            out.append(w.upper().replace("L", "I"))
        elif re.fullmatch(r"[ivx]+", lw) and i > 0:
            out.append(w.upper())
        elif lw in MINUSCULAS and i > 0:
            out.append(lw)
        else:
            out.append(re.sub(r"^(\(?)(\w)", lambda m: m.group(1) + m.group(2).upper(), lw))
    return " ".join(out)


def talhao(v):
    n = num(v)
    return str(int(n)) if n is not None and n == int(n) else s(v)


def oper_frente(desc):
    d = desc.upper()
    for chave, fr, at in OPER_FRENTE:
        if chave in d:
            return fr, at
    return None, None


def rsc_semana(ws, dias):
    """Relatório de operações (RSC0819R): uma linha por OS × talhão, com os produtos."""
    out, nomes = [], {}
    faz = tal = ha_tal = op = cur = None
    for r in ws[2:]:
        c0 = s(cel(r, 0))
        if c0 == "PROPRIEDADE:":
            faz = int(num(cel(r, 1)))
            nomes[str(faz)] = nome_faz(cel(r, 5))
        elif c0 == "TALHÃO:":
            tal, ha_tal = talhao(cel(r, 1)), num(cel(r, 4))
        elif c0 == "Operação:":
            cod, desc = int(num(cel(r, 1))), s(cel(r, 3))
            op = (cod, desc, ROTULOS.get(cod) or re.sub(r"\(HA\)$", "", desc).strip(" (").capitalize())
        elif c0 == "OS" or cel(r, 0) is None:
            continue
        else:
            if c0:                                   # linha nova de OS; as seguintes sem OS são outros produtos dela
                d, cur = data_xl(cel(r, 2)), None
                fr, at = oper_frente(op[1])
                if d and iso(d) in dias and fr:
                    cur = {"f": fr, "a": at, "d": iso(d), "faz": faz, "t": tal, "ht": r2(ha_tal),
                           "ha": r2(num(cel(r, 5))), "os": int(num(c0)), "op": op[2], "prod": []}
                    out.append(cur)
            if cur is not None and s(cel(r, 7)):
                un = UNIDADES.get(s(cel(r, 8)), s(cel(r, 8)))
                dose, rec = num(cel(r, 9)), num(cel(r, 11))
                cur["prod"].append([s(cel(r, 7)), None if dose is None else round(dose, 4),
                                    None if rec is None else round(rec, 4), un])
    return out, nomes


def mol3_semana(wb, dias):
    """OS da 3ª molécula encerradas nos dias da semana (aba Realizado - 3ª Mol)."""
    nome = next((k for k in wb if k.startswith("Realizado - 3")), None)
    if not nome:
        return [], {}
    ws = wb[nome]
    rot = {s(v): j for j, v in enumerate(ws[0]) if isinstance(v, str)}
    out, nomes = [], {}
    for r in ws[1:]:
        d = data_xl(cel(r, rot["Data Op."]))
        faz = num(cel(r, rot["Propriedade"]))
        if not d or iso(d) not in dias or not faz:
            continue
        faz = int(faz)
        nomes[str(faz)] = nome_faz(cel(r, rot["Fundo Agrícola"]))
        out.append({"f": "uniport", "a": None, "d": iso(d), "faz": faz, "t": talhao(cel(r, rot["Talhão"])),
                    "ht": r2(num(cel(r, rot["Plantada"]))), "ha": r2(num(cel(r, rot["Recom."]))),
                    "os": int(num(cel(r, rot["OS"]))), "op": "3ª molécula · Uniport bordadura",
                    "corte": int(num(cel(r, rot["Corte"])) or 0) or None, "prod": []})
    return out, nomes


def locais_semana(wb_prog, rscs, dias):
    linhas, nomes, fontes = [], {}, []
    for arq, ws in rscs:
        lin, nm = rsc_semana(ws, dias)
        if lin:
            linhas += lin
            nomes.update(nm)
            fontes.append(arq)
    # a 3ª molécula entra quando a OS não veio no relatório do sistema
    ja = {(l["os"], l["faz"], l["t"]) for l in linhas}
    lin, nm = mol3_semana(wb_prog, dias)
    lin = [l for l in lin if (l["os"], l["faz"], l["t"]) not in ja]
    if lin:
        linhas += lin
        nomes.update(nm)
        fontes.append("Realizado - 3ª Mol")
    linhas.sort(key=lambda l: (l["f"], l["d"], l["faz"], num(l["t"]) or 0))
    return {"fazendas": nomes, "linhas": linhas, "fontes": fontes}


def data_atualizacao(wb):
    for r in wb["Pré-Emerg -Cana Soca SAFRA"][:6]:
        for j, v in enumerate(r):
            if s(v).startswith("DATA ATUALIZA"):
                d = data_xl(cel(r, j + 1))
                if d:
                    return iso(d)
    return None


# ------------------------------------------------------------------ main
def main():
    arquivos = sorted(glob.glob(os.path.join(AQUI, "dados", "*.xls*")))
    wb_sem = wb_prog = None
    rscs = []
    for a in arquivos:
        if os.path.basename(a).startswith("~$"):
            continue
        try:
            wb = le(a)
        except Exception as e:                       # arquivo de outro relatório na pasta
            print("ignorado:", os.path.basename(a), "-", e)
            continue
        if "Semanal-Acumulado" in wb:
            wb_sem, arq_sem = wb, os.path.basename(a)
        if "Pré-Emerg -Cana Soca SAFRA" in wb:
            wb_prog, arq_prog = wb, os.path.basename(a)
        primeira = next(iter(wb.values()), [])
        if primeira and s(cel(primeira[0], 3)).startswith("Período Acumulado"):   # RSC0819R do sistema
            rscs.append((os.path.basename(a), primeira))
    if not wb_sem or not wb_prog:
        sys.exit("Faltou planilha em dados/: preciso da que tem a aba Semanal-Acumulado (.xlsb) "
                 "e da que tem as abas dos programas (SAFRA 2026 .xlsx)")
    sem, frentes, avisos = semana(wb_sem["Semanal-Acumulado"])
    dados = {
        "fonte": {"semana": arq_sem, "programas": arq_prog, "atualizacaoProgramas": data_atualizacao(wb_prog)},
        "semana": sem, "frentes": frentes, "avisos": avisos, "programas": programas(wb_prog),
        "nomesMes": MESES, "locais": locais_semana(wb_prog, rscs, set(sem["dias"])),
    }
    js = "window.TRATOS = " + json.dumps(dados, ensure_ascii=False, separators=(",", ":"))
    html = open(HTML, encoding="utf-8").read()
    novo, n = re.subn(r"(/\*DADOS-INICIO\*/\n).*?(\n/\*DADOS-FIM\*/)",
                      lambda m: m.group(1) + js + ";" + m.group(2), html, flags=re.S)
    if n != 1:
        sys.exit("Marcadores DADOS-INICIO / DADOS-FIM não encontrados no index.html")
    open(HTML, "w", encoding="utf-8", newline="\n").write(novo)
    n_at = sum(len(f["atividades"]) for f in frentes)
    print(f"semana {sem['dias'][0]} a {sem['dias'][-1]} ({sem['qtdDias']} dias) · {len(frentes)} frentes · "
          f"{n_at} atividades · {len(dados['programas'])} programas · {len(avisos)} aviso(s)")
    loc = dados["locais"]
    print(f"fazendas e talhões da semana: {len(loc['linhas'])} linha(s) de "
          f"{len({l['faz'] for l in loc['linhas']})} fazenda(s) · fontes: {', '.join(loc['fontes']) or 'nenhuma'}")
    for a in avisos:
        print("  aviso:", a)


if __name__ == "__main__":
    main()
