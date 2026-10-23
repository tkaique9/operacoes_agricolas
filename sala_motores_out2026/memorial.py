"""Memorial técnico (PDF A4) e planilha de quadros (xlsx) do estudo da oficina de motores e transmissões.

python memorial.py  ->  Memorial_Tecnico_Sala_Motores_OUT2026.pdf, Quadros_Sala_Motores_OUT2026.xlsx
Rodar depois de pranchas.py: as figuras são recortes das pranchas (build/SM-0x.png).
Páginas em HTML, PDF impresso pelo Edge em modo headless.
"""
import base64
import html
import io
import json
import subprocess
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

import conteudo as C
import galpao as G
import layout as LY
import verificacao as V

AQUI = Path(__file__).parent
BUILD = AQUI / "build"
PDF = AQUI / "Memorial_Tecnico_Sala_Motores_OUT2026.pdf"
XLSX = AQUI / "Quadros_Sala_Motores_OUT2026.xlsx"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
DATA = "21/10/2026"
NAVY = "#1f3d63"


def n(x, d=0):
    s = f"{x:,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def e(s):
    return html.escape(str(s))


# ---------------------------------------------------------------- figuras (recortes das pranchas)
def recorte(folha, x0, y0, x1, y1):
    """Recorte em mm do A3 (origem embaixo à esquerda) -> data URI PNG."""
    im = Image.open(BUILD / f"{folha}.png")
    s = im.width / 420
    im = im.crop((int(x0 * s), int((297 - y1) * s), int(x1 * s), int((297 - y0) * s))).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def figura(uri, legenda, larg="100%"):
    return f'<figure><img src="{uri}" style="width:{larg}"><figcaption>{legenda}</figcaption></figure>'


# ---------------------------------------------------------------- blocos de html
def tabela(cab, linhas, cls="", larguras=None):
    col = "".join(f'<col style="width:{w}">' for w in larguras) if larguras else ""
    th = "".join(f"<th>{e(c)}</th>" for c in cab)
    tr = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in ln) + "</tr>" for ln in linhas)
    return f'<table class="{cls}"><colgroup>{col}</colgroup><thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>'


def conf(txt):
    """Marca em magenta o que depende de confirmação."""
    return f'<span class="conf">{e(txt)}</span>'


def lista(itens):
    return "<ul>" + "".join(f"<li>{i}</li>" for i in itens) + "</ul>"


# ---------------------------------------------------------------- dados derivados
Q1 = V.quadro_areas(1)
Q2 = V.quadro_areas(2)
A1 = {i: a for i, _, _, a, _ in Q1}
A2 = {i: a for i, _, _, a, _ in Q2}
PROD1 = sum(A1[f"S{i}"] for i in range(1, 11))
PROD2 = sum(A2[f"S{i}"] for i in range(1, 11))
QTD1 = {}
for t, _, _ in LY.A1_EQUIP:
    QTD1[t] = QTD1.get(t, 0) + 1
QTD1["E01"] = 1
SETOR_DE = {}
for sid, _, _, rets, _ in LY.A1_SETORES:
    for t, r, _ in LY.A1_EQUIP:
        if any(V.contido(r, q) for q in rets):
            SETOR_DE.setdefault(t, set()).add(sid)


def capa():
    return f"""
<section class="capa">
  <div class="faixa">ESTUDO CONCEITUAL — NÃO EXECUTIVO — NÃO LIBERADO PARA OBRA</div>
  <h1>Oficina de motores e transmissões<br>com lavagem de componentes</h1>
  <p class="sub">Memorial técnico do estudo de layout sobre o galpão do desenho CRV(MG)20102026 rev. 00
  (“Galpão para motor a diesel”, 7.850 x 20.000 mm) — CRV Açúcar &amp; Álcool, Capinópolis/MG</p>
  <table class="kv">
    <tr><td>Data</td><td>{DATA}</td></tr>
    <tr><td>Revisão</td><td>A — estudo preliminar</td></tr>
    <tr><td>Base</td><td>{e(G.DESENHO['numero'])} — {e(G.DESENHO['revisao'])}; {e(G.DESENHO['autoria'])}</td></tr>
    <tr><td>Responsável técnico</td><td>A definir — profissional legalmente habilitado, com ART/RRT</td></tr>
    <tr><td>Arquivos que acompanham</td><td>Pranchas SM-01 a SM-07 (PDF A3 e SVG); planta em CAD (DXF, mm 1:1);
        quadros em planilha (xlsx)</td></tr>
  </table>
  <p class="aviso">Este documento organiza o layout e os requisitos para validação interna, levantamento em campo,
  seleção de equipamentos e orçamento. Dimensões de equipamentos são envelopes de reserva, não especificação.
  Itens <span class="conf">em magenta</span> dependem de confirmação. Estruturas, instalações, içamento, prevenção
  contra incêndio e efluentes precisam de projeto e ART de profissionais habilitados antes de qualquer execução.</p>
</section>"""


def s1_resumo():
    return f"""
<h2>1. Resumo executivo</h2>
<p>O galpão tem 20,00 x 7,85 m entre eixos (152,85 m² internos), um escritório de 4,0 x 3,0 m no fundo, um portão de
4,0 x 5,0 m na frente e uma porta de 1,0 m nos fundos. O projeto já prevê um <b>pórtico móvel de 2 t com talha elétrica
sobre trilhos no piso</b>, com vão de 6.534 mm ao longo de todo o comprimento. Esse pórtico é o que mais condiciona o
layout: as pernas varrem duas faixas junto às paredes e, com a folga de segurança, sobram 5,01 m de largura útil
(Y 1.418 a 6.432).</p>
<p><b>Recomendação — Alternativa 1:</b> manter o pórtico do projeto, limitar o curso com batentes (pernas entre
X 5.300 e 19.250, longe do escritório) e organizar a oficina em <b>U</b>: lado A (porta) sujo — recebimento,
desmontagem e lavagem; lado B (escritório) limpo — inspeção, montagem, controle de qualidade e expedição; a
lavagem fica no fundo do lado A, onde o duto de vapor e a drenagem saem pela parede de fundos sem cruzar o
pórtico. A lavadora industrial fecha a transição sujo → limpo; os únicos contatos entre os lados (boca da lavadora e
portão) são tratados por procedimento.</p>
<p>A Alternativa 2 (ponte rolante, sem trilhos no piso) aumenta a área dos setores produtivos de
{n(PROD1, 1)} m² para {n(PROD2, 1)} m² (+{n((PROD2 / PROD1 - 1) * 100)}%) e permite uma segunda posição de montagem,
mas depende de viga de rolamento, pilares e fundações que o projeto atual não tem. Só se justifica se o galpão ainda
não foi executado e o estudo estrutural confirmar a viabilidade.</p>
<p><b>Decisão urgente:</b> trilhos, caixas de piso, dreno, canaleta e contenção da lavagem são embutidos no piso.
Precisam estar definidos <b>antes da concretagem</b> (pendência P03).</p>
<p><b>Pendências que impedem liberar para execução:</b> pesos e dimensões dos conjuntos (P01), geometria e alcance
do pórtico (P02), situação da obra (P03), origem do pórtico (P04), energia disponível (P07), destino dos efluentes
(P09), divergências estruturais (P16) e exigências do CBMMG (P18).</p>
{tabela(["Indicador", "Alternativa 1 (recomendada)", "Alternativa 2"], [
        ["Área dos setores produtivos S1–S10", f"{n(PROD1, 2)} m²", f"{n(PROD2, 2)} m²"],
        ["Faixas do pórtico (acesso restrito)", f"{n(A1['S13'], 2)} m²", "—"],
        ["Posições de motor (desmontagem + montagem)", "1 + 1", "1 + 2"],
        ["Corredor central", "1,20–1,25 m", "1,60 m"],
        ["Depende de obra estrutural nova", "Não", "Sim"]], larguras=["44%", "28%", "28%"])}
"""


def s2_escopo():
    return f"""
<h2>2. Objetivo e escopo</h2>
<p>Oficina de manutenção de motores diesel (máquinas agrícolas, caminhões e equipamentos de apoio), motores
industriais, transmissões, caixas de câmbio e redutores da usina, com volume médio de serviço. A oficina recebe o
conjunto, abre a ordem de serviço, desmonta, lava, inspeciona, monta, faz testes funcionais e libera. Usinagem e
serviços especializados ficam fora (premissa PR06).</p>
{tabela(["Atividade", "Onde", "Interno ou externo"], [[e(a), e(b), e(c)] for a, b, c in C.ATIVIDADES],
        larguras=["62%", "14%", "24%"])}
<p class="nota">* Substituição de camisas apenas nos modelos de camisa úmida e com ferramental do fabricante.</p>
"""


def s3_desenho():
    conf_rows = [[e(c["item"]), n(c["cota"]), n(c["medido"])]
                 for c in json.loads((AQUI / "desenho_original.json").read_text(encoding="utf-8"))["conferencia"]]
    alt = [[e(k), n(v), e(f)] for k, (v, f) in G.ALTURAS.items()]
    return f"""
<h2>3. Análise técnica do desenho original</h2>
<p><b>Documento:</b> {e(G.DESENHO['numero'])}, {e(G.DESENHO['titulo'])}, {e(G.DESENHO['formato'])}.
Vistas: {e(G.DESENHO['vistas'])}.</p>
<p><b>Método:</b> o PDF é vetorial (A1, 1:50), sem camadas nem cotas associativas. As linhas foram lidas
diretamente do arquivo e convertidas para milímetros (1 pt = 17,64 mm em 1:50). Cada medida usada foi conferida
contra a cota escrita quando ela existe; a diferença ficou em até 2 mm. Onde não há cota, a medida é marcada como
“medido” e entra como premissa.</p>
<h3>3.1 Dimensões confirmadas</h3>
{tabela(["Item", "Cota no desenho (mm)", "Medido nas linhas (mm)"], conf_rows, larguras=["60%", "20%", "20%"])}
{tabela(["Área", "m²", "Como"], [
        ["Entre eixos", n(G.AREA_EIXOS, 2), "20,00 x 7,85"],
        ["Construída (faces externas)", n(G.AREA_CONSTRUIDA, 2), "20,15 x 8,00"],
        ["Interna (faces internas — base do quadro de áreas)", n(G.AREA_INTERNA, 2), "19,85 x 7,70"],
        ["Escritório existente (com paredes)", n(G.AREA_ESCRITORIO, 2), "4,00 x 3,00"]], larguras=["60%", "15%", "25%"])}
<h3>3.2 Alturas acima do piso acabado</h3>
{tabela(["Elemento", "mm", "Fonte"], alt, larguras=["35%", "15%", "50%"])}
{figura(recorte("SM-01", 20, 114, 340, 287), "Figura 1 — Galpão existente redesenhado sobre as linhas do PDF original "
        "(azul-claro). Prancha SM-01, escala 1:75 no A3.")}
<h3>3.3 Interferências físicas</h3>
{tabela(["", "Interferência"], [[i, e(t)] for i, t in G.INTERFERENCIAS], larguras=["8%", "92%"])}
<h3>3.4 Divergências e informações ausentes</h3>
{tabela(["", "Divergência ou lacuna", "Quem resolve"], [[d, conf(t), e(q)] for d, t, q in G.DIVERGENCIAS],
        larguras=["7%", "73%", "20%"])}
<p><b>Acessos:</b> motores e transmissões só entram e saem pelo portão de 4,0 x 5,0 m (a empilhadeira deixa o conjunto
na posição de recebimento, dentro do alcance do gancho). A porta de 1,0 m serve a pessoas e é a segunda saída de
emergência. O vão de 2,25 m na lateral A (D08) não foi usado no layout até ser identificado.</p>
"""


def s4_premissas():
    pend = [[p, e(a), e(b), e(c), "<b>Sim</b>" if bl == "Sim" else "Não"] for p, a, b, c, bl in C.PENDENCIAS]
    return f"""
<h2>4. Premissas e pendências</h2>
<h3>4.1 Premissas adotadas</h3>
{tabela(["", "Premissa"], [[p, conf(t)] for p, t in C.PREMISSAS], larguras=["8%", "92%"])}
<h3>4.2 Pendências de levantamento e validação</h3>
{tabela(["", "Informação que falta", "Por que importa", "Quem levanta", "Bloqueia execução"], pend,
        larguras=["6%", "38%", "30%", "16%", "10%"])}
"""


def s5_alternativas():
    import pranchas as P
    comp = P.comparativo()
    return f"""
<h2>5. Alternativas de layout e recomendação</h2>
<h3>5.1 Alternativa 1 — layout funcional com o pórtico do projeto (recomendada)</h3>
<p>Mantém os trilhos e o pórtico do desenho, com curso limitado por batentes para que as pernas não passem ao lado do
escritório (hoje sobrariam ≈ 10 mm — D10). Toda a operação fica na faixa útil entre as pernas, em U: lado A sujo, lado
B limpo, corredor central de 1,20–1,25 m, lavagem no fundo do lado A e apoio (ferramentaria, químicos, quarentena,
lava-olhos e quadro elétrico) no canto dos fundos, fora do curso do pórtico. Não depende de obra estrutural nova.</p>
<h3>5.2 Alternativa 2 — layout otimizado com ponte rolante</h3>
<p>Troca o pórtico por ponte rolante univiga de 2 t apoiada em vigas de rolamento. Sem trilhos no piso, as paredes
ficam livres para bancadas leves, armários e pontos de utilidade; o corredor central passa a 1,60 m (paleteira) e cabe
uma segunda posição de montagem de motor. Exige viga de rolamento e pilares/fundações capazes das reações das rodas
— os pilares atuais (bloco 600 x 600 sobre uma estaca Ø 250 x 4 m) foram projetados para a cobertura.</p>
{figura(recorte("SM-07", 20, 126, 340, 287), "Figura 2 — Alternativa 2 (ponte rolante). Prancha SM-07, escala 1:75 no A3.")}
<h3>5.3 Comparação</h3>
{tabela(comp[0], [[e(c) for c in ln] for ln in comp[1:]], larguras=["34%", "33%", "33%"])}
<h3>5.4 Recomendação</h3>
<p>Implantar a <b>Alternativa 1</b>: é executável sobre o projeto aprovado, atende ao fluxo e à segregação
sujo/limpo com folgas de segurança verificadas e não depende de reforço estrutural. Reavaliar a Alternativa 2
somente se, antes da concretagem do piso, o estudo estrutural mostrar viabilidade a custo aceitável — ela recupera a
largura perdida para as pernas do pórtico e dá capacidade para uma segunda montagem simultânea.</p>
"""


def s6_layout():
    linhas = []
    for sid, nome, zona, rets, just in LY.A1_SETORES:
        r = rets[0]
        dims = " + ".join(f"{n(q[1] - q[0])} x {n(q[3] - q[2])}" for q in rets)
        eq = ", ".join(sorted({t for t, rr, _ in LY.A1_EQUIP if any(V.contido(rr, q) for q in rets)})) or "—"
        linhas.append([f"<b>{sid}</b>", e(nome), dims, n(A1.get(sid, sum(V.area(q) for q in rets)), 2), eq, e(just)])
    return f"""
<h2>6. Layout recomendado (Alternativa 1)</h2>
{figura(recorte("SM-02", 20, 126, 340, 287), "Figura 3 — Planta baixa proposta. Prancha SM-02 (Desenho 1), escala 1:75 no A3; "
        "cotas completas na prancha SM-03 (Desenho 2).")}
<p>Coordenadas em mm a partir do cruzamento dos eixos 1 (parede de fundos) e A (parede do lado da porta). Os setores
são retângulos sem sobreposição; a soma deles com a circulação dá exatamente a área interna (capítulo 14).</p>
{tabela(["", "Setor", "Dimensões (mm)", "m²", "Equipamentos", "Justificativa da posição"], linhas,
        larguras=["5%", "19%", "13%", "6%", "15%", "42%"], cls="pequena")}
{figura(recorte("SM-03", 20, 12, 137, 112), "Figura 4 — Corte transversal esquemático com as folgas verticais "
        "(prancha SM-03, 1:100 no A3).", larg="70%")}
"""


def s7_fluxo():
    return f"""
<h2>7. Fluxo operacional</h2>
{figura(recorte("SM-04", 20, 126, 340, 287), "Figura 5 — Fluxo de componentes, pessoas e resíduos. Prancha SM-04 (Desenho 3).")}
<p><b>Sequência:</b> (1) o conjunto chega pré-lavado por fora e a empilhadeira o deixa em E02; abre-se a OS com fotos;
(2) drenam-se óleo e líquido de arrefecimento (E03); (3) o pórtico leva o conjunto ao cavalete E04 ou à bancada E06;
as peças são separadas por OS nos carrinhos sujos E05; (4) peças pequenas vão ao tanque E08; (5) peças grandes e
cestos entram na lavadora E07 pelo carro extraível; (6) na descarga, as peças atravessam só o corredor até a mesa
E10 (escorrimento, secagem, inspeção visual, protetivo); (7) depois de frias, são medidas em E12 ou no escritório
(metrologia fina); aprovadas vão ao kit da OS (E14) com as peças novas; (8) montagem em E17/E15; (9) testes funcionais
e controle de qualidade em S10 (E31); (10) expedição com capa em E18.</p>
<p><b>Segregação:</b> lado A só recebe peças sujas; lado B só peças limpas. Carrinhos, cestos e etiquetas têm cores
diferentes por estado (capítulo 10). Os dois pontos de contato — a boca da lavadora e o portão — são controlados por
procedimento: limpar o carro da lavadora e o piso entre ciclos; não receber e expedir ao mesmo tempo; conjunto
liberado sai coberto. O pórtico é único para os dois lados: carga limpa não passa sobre posição suja aberta.</p>
<p><b>Retorno a etapa anterior:</b> peça reprovada vai à quarentena E20 (etiqueta vermelha) ou volta à lavadora pelo
corredor transversal; peça enviada à retífica sai pela expedição e volta pelo recebimento, passando de novo pela
lavagem antes de entrar no lado limpo. A etiqueta da OS acompanha cesto, carrinho e conjunto em todas as etapas.</p>
<p><b>Pessoas e veículos:</b> pessoas entram pela porta de fundos (lado do apoio e do escritório) ou pelo portão;
a empilhadeira só opera na zona do portão (S1/S10). <b>Rotas de fuga:</b> corredor central → portão; corredor
transversal (1,10 m) → rota de 1,25 m → porta de fundos. Distância máxima até uma saída ≈ 13 m.
{conf("A porta de fundos abre para dentro; confirmar com o projeto de PCI se deve abrir no sentido da fuga.")}</p>
{figura(recorte("SM-04", 20, 12, 245, 126), "Figura 6 — Fluxograma do processo (prancha SM-04).", larg="88%")}
"""


def s8_lavagem():
    return f"""
<h2>8. Área de lavagem de componentes</h2>
{figura(recorte("SM-05", 20, 120, 333, 287), "Figura 7 — Planta ampliada da lavagem. Prancha SM-05 (Desenho 4), escala 1:40 no A3.")}
<h3>8.1 Processo por tipo de componente</h3>
{tabela(["Componente", "Processo indicado", "Cuidados"], [
        ["Blocos, cabeçotes, cárteres e carcaças de ferro fundido/aço", "Aspersão aquosa alcalina quente na lavadora E07",
         "Proteger roscas e galerias; secar e aplicar protetivo logo após"],
        ["Carcaças, tampas e cárteres de alumínio", conf("E07 com detergente compatível com alumínio (pH e inibidor)"),
         conf("Química e temperatura conforme fabricante do componente e do detergente")],
        ["Engrenagens, eixos, rolamentos", "E07 em cesto ou E08", "Rolamentos a reaproveitar lavados à parte e protegidos"],
        ["Peças pequenas e fixadores", "Tanque E08 (solução aquosa, bomba e pincel)", "Cestos por OS; sem diesel ou solvente"],
        ["Vedações, sensores, chicotes, componentes de injeção", "Não lavar na lavadora", "Limpeza manual ou serviço especializado"],
        ["Carbonização pesada (pistões, válvulas)", conf("Avaliar processo específico (imersão, ultrassom) se a demanda justificar"),
         "Hoje: serviço externo"]], larguras=["32%", "34%", "34%"])}
<h3>8.2 Lavadora industrial (E07) — critérios de seleção</h3>
{lista([conf("Maior peça (dimensões e peso) define diâmetro, altura útil e carga da mesa giratória (P01)."),
        conf("Química, temperatura e pressão aprovadas para ferro, aço e alumínio; compatível com vedações remanescentes."),
        "Carro extraível para carga pelo pórtico; porta intertravada; apreciação de riscos NR-12 pelo fornecedor.",
        "Separador de óleo (skimmer), filtragem do banho e remoção de lodo pela frente da máquina.",
        conf("Utilidades: tensão, potência de aquecimento, água de reposição; condensador de vapor ou exaustão (EX-1)."),
        "Consumo de detergente e frequência de troca do banho → volume mensal de resíduo e custo operacional.",
        "Não escolher só por preço ou área: o envelope reservado (2.500 x 2.200 x 2.600 mm + carro de 1.600 mm) "
        "é o limite do layout; máquina maior obriga a revisar S5."])}
<h3>8.3 Organização física</h3>
<p>Peças sujas chegam pelo lado A (carrinhos E05) → preparação (retirada de tampões, separação em cestos) → tanque E08
ou carro da lavadora → ciclo fechado → descarga direta para a mesa E10 no lado limpo → inspeção visual → estabilização
térmica → inspeção dimensional (E12). Atrás da lavadora fica o corredor transversal de 1,10 m (acesso de serviço);
ao lado, o apoio S12 com armário de químicos com bacia (E21), lava-olhos e chuveiro (E22) e o quadro elétrico (E23)
fora da área úmida.</p>
<h3>8.4 Piso, contenção e drenagem</h3>
<p>Área S5 (4,30 x 2,28 m) com piso impermeável resistente a álcali e calor ({conf("ex.: uretano-cimento")}), juntas
seladas e rodapé sanitário; contenção perimetral por lombada transponível de 50 mm ou piso rebaixado 50 mm;
caimento ≥ 1% para a canaleta CN-1 com grelha; dreno enterrado com caixa de inspeção até a caixa de contenção
estanque CX-1 fora da parede de fundos. A CX-1 não tem saída: é esvaziada por caminhão de sucção licenciado
{conf("ou ligada à estação de tratamento da usina, se licenciada para esse efluente (P09)")}. O restante do piso é
limpo a seco (absorvente e lavadora de piso), com bandejas de gotejamento em E02, E04 e E17.</p>
{figura(recorte("SM-05", 20, 28, 190, 115), "Figura 8 — Corte esquemático A-A da lavagem (prancha SM-05).", larg="80%")}
<h3>8.5 Efluentes, resíduos e produtos químicos</h3>
{lista(["<b>Banho esgotado da lavadora:</b> bombeado para IBC/tambor e enviado a destinador licenciado. "
        "A solução alcalina emulsiona o óleo: <b>não</b> enviar para caixa separadora de água e óleo, que só funciona com "
        "óleo livre.",
        "<b>Óleo do separador da lavadora e óleo drenado dos conjuntos:</b> OLUC, coleta por coletor autorizado com "
        "certificado (CONAMA 362/2005).",
        "<b>Lodo, filtros, estopas e absorventes contaminados:</b> resíduos perigosos (classificar pela NBR 10004), "
        "recipientes fechados e identificados, abrigo externo (NBR 12235), transporte com MTR.",
        "<b>Líquido de arrefecimento usado:</b> recipiente próprio, destinação licenciada.",
        "<b>Nenhum descarte</b> na rede pluvial, na rede sanitária ou no solo sem tratamento e autorização.",
        "<b>Produtos químicos:</b> armário com bacia de contenção (E21), FDS no local, rotulagem GHS, incompatíveis "
        "separados; inflamáveis só em armário próprio e dentro dos limites (NR-20/NBR 17505, se aplicável)."])}
"""


def s9_icamento():
    m = C.MATRIZ_ICAMENTO
    return f"""
<h2>9. Movimentação de cargas</h2>
<p><b>Equipamento previsto no projeto:</b> {e(G.PORTICO)}. Vão entre trilhos 6.534 mm; viga com topo a 4.460 mm e
base a ≈ 4.150 mm do piso; pernas de ≈ 520 mm de largura (medido no corte).</p>
<h3>9.1 Requisitos para confirmar a capacidade</h3>
{lista([conf("Peso máximo de cada motor, transmissão e redutor atendido (manual de serviço), com acessórios que "
             "permanecem no conjunto (volante, carcaça, coletores)."),
        conf("Peso dos dispositivos de içamento: balancim, cintas, correntes, manilhas e adaptadores de cavalete."),
        "Capacidade nominal do pórtico e da talha ≥ soma acima, com os critérios do fabricante e da NBR 8400; "
        "acessórios com carga de trabalho identificada.",
        conf("Altura útil do gancho com o maior conjunto suspenso e folga sobre cavaletes, bancadas e lavadora (P02)."),
        "Cobertura: todas as posições que exigem içamento (E02, E04, E06, carro da E07, E10, E12, E15, E17, E18) estão "
        f"dentro do alcance adotado (X {n(LY.GANCHO_X[0])}–{n(LY.GANCHO_X[1])}, Y {n(LY.GANCHO_Y[0])}–{n(LY.GANCHO_Y[1])}).",
        "Interferências: dentro do curso, nada suspenso abaixo de ~4,6 m; luminárias acima da viga e abaixo do forro.",
        "Curso limitado por batentes (pernas X 5.300–19.250), alarme na translação e faixas S13 livres e sinalizadas.",
        "Trilhos ancorados e nivelados; base dos trilhos e laje verificadas para as cargas de roda."])}
<h3>9.2 Matriz comparativa</h3>
{tabela(["Critério", "Ponte rolante univiga 2 t", "Pórtico 2 t sobre trilhos (projeto)", "Monovia central + talha",
         "Guindaste de oficina / pórtico com rodízios"], [[e(c) for c in ln] for ln in m], cls="pequena",
        larguras=["14%", "22%", "24%", "20%", "20%"])}
<p><b>Critérios de escolha:</b> (1) peso do maior conjunto com dispositivos; (2) área a cobrir; (3) altura útil;
(4) capacidade da estrutura e das fundações; (5) interferências com piso, iluminação e dutos; (6) custo e prazo.
Neste galpão, com o projeto mantido, o pórtico atende; a ponte rolante é melhor para o layout, mas só com estrutura
própria verificada. A monovia não serve ao fluxo em U.</p>
<p class="nota">O dimensionamento final do equipamento de içamento, da estrutura de apoio, dos trilhos, dos acessórios e
dos dispositivos de segurança é de profissional habilitado, com ART (NR-11, NR-12, ABNT NBR 8400).</p>
"""


def s10_bancadas():
    import pranchas as P  # noqa: F401
    ident = [
        ("Aguardando manutenção", "S1 (E02) e carrinhos E05", "Etiqueta da OS laranja; estado de chegada fotografado"),
        ("Sujas aguardando lavagem", "S4 e zona de carga de S5", "Cestos laranja com etiqueta da OS"),
        ("Limpas aguardando inspeção", "S6 (E10)", "Cestos azuis; protetivo nas ferrosas"),
        ("Aprovadas aguardando montagem", "Kit da OS (E14) em S8/S9", "Etiqueta verde “aprovado” + nº da OS; capa"),
        ("Reprovadas ou em avaliação", "Quarentena E20 (S12)", "Etiqueta vermelha com o motivo; rack vermelho"),
        ("Ferramentas e dispositivos", "E19 (S12) e carros por mecânico", "Painel com silhueta; torquímetros com selo de calibração"),
        ("Instrumentos de medição", "Escritório/metrologia (E24)", "Armário fechado; certificado de calibração"),
        ("Reparados aguardando retirada", "Expedição E18 (S10)", "Etiqueta de liberação do CQ; capa de proteção"),
        ("Resíduos", "E03 e abrigo externo", "Recipientes identificados por tipo"),
    ]
    return f"""
<h2>10. Bancadas, ferramentas e armazenamento</h2>
{lista(["<b>Desmontagem e montagem de motores:</b> cavalete rotativo (E04/E17) — o motor gira e o mecânico trabalha em pé, "
        "com 800 mm livres nas laterais de trabalho; bandeja de gotejamento sob o motor.",
        "<b>Transmissões e redutores:</b> bancada reforçada de 2.000 x 900 mm com morsa (E06 suja, E15 limpa); altura da "
        "tarefa ajustada à peça (NR-17); peças acima do limite manual com auxílio do pórtico.",
        "<b>Prensa hidráulica (E16)</b> no lado limpo, para rolamentos e buchas novos; capacidade a definir.",
        "<b>Peças por OS:</b> carrinhos (E05 sujos, E14 limpos) em vez de estantes fixas — o kit da OS anda com o serviço "
        "e evita busca e mistura. Parafusos e peças pequenas em cestos identificados dentro do carrinho.",
        "<b>Ferramentas especiais e dispositivos:</b> armários E19 no apoio S12; ferramentas de uso diário em carros por "
        "mecânico. <b>Instrumentos de precisão</b> no escritório/metrologia, longe de água, óleo e sopro.",
        "<b>Quarentena</b> (E20) separada, identificada em vermelho."])}
{tabela(["Categoria", "Onde fica", "Identificação visual"], [[e(a), e(b), e(c)] for a, b, c in ident],
        larguras=["30%", "30%", "40%"])}
"""


def s11_equipamentos():
    linhas = []
    for t, (nome, func, qb, dim, esp, util, req, rest, prio, levantar) in LY.EQUIP.items():
        q = QTD1.get(t, qb)
        linhas.append([f"<b>{t}</b>", e(nome), e(func), str(q), e(dim), e(esp), e(util), e(req), e(rest), e(prio),
                       conf(levantar) if levantar != "—" else "—"])
    return f"""
<section class="paisagem">
<h2>11. Lista preliminar de equipamentos e requisitos de instalação</h2>
<p class="nota">Quantidades da Alternativa 1. Dimensões são envelopes de reserva até a escolha de fabricante e modelo;
o espaço operacional inclui abertura de portas, acesso para manutenção e movimentação de carga.</p>
{tabela(["Tag", "Equipamento", "Função", "Qtd", "Dimensão de referência (mm)", "Espaço operacional", "Utilidades",
         "Requisitos de instalação", "Restrições de segurança", "Prioridade", "Dados a levantar"], linhas,
        cls="mini", larguras=["3%", "11%", "11%", "3%", "8%", "9%", "9%", "11%", "12%", "6%", "17%"])}
</section>"""


def s12_infra():
    return f"""
<h2>12. Infraestrutura</h2>
{figura(recorte("SM-06", 20, 126, 340, 287), "Figura 9 — Infraestrutura preliminar. Prancha SM-06 (Desenho 5).")}
<p><b>Princípio:</b> dentro do curso do pórtico nada pode descer do teto abaixo de ~4,6 m e nada pode cruzar as
faixas das pernas. Por isso as tomadas das estações ficam em <b>caixas de piso</b> (CP-01 a CP-07) alimentadas por
eletrodutos sob o corredor central; o ar comprimido chega por <b>enroladores nas pontas</b> (R1–R4, 15 m cobrem a
faixa útil); a lavadora recebe energia, água e exaustão <b>por cima do apoio S12</b>, fora do curso, saindo pela parede
de fundos.</p>
{tabela(["Sistema", "Solução preliminar", "A levantar / projeto"], [
        ["Energia", "QDO (E23) na parede de fundos; circuitos C1 lavadora (trifásico dedicado), C2 pórtico (festão no trilho B), "
                    "C3 caixas de piso 220/380 V, C4/C5 iluminação geral e de emergência, C6 escritório; DR nas tomadas; aterramento",
         conf("Tensão, potência disponível e ponto de entrega (P07); projeto NBR 5410; SPDA NBR 5419")],
        ["Iluminação", "LED industrial acima da viga do pórtico e abaixo do forro (≈ 11 pontos); reforço localizado em E06, E12, "
                       "E15 e na metrologia; emergência nas saídas e corredores",
         conf("Cálculo luminotécnico pela NBR ISO/CIE 8995-1")],
        ["Ar comprimido", "Compressor em abrigo externo ou rede existente; rede interna pelos fundos; rede externa pela fachada A "
                          "até a frente; enroladores R1–R4; tratamento do ar", conf("Demanda de ar e local do compressor (P08)")],
        ["Água", "Entrada pelos fundos para lava-olhos/lavatório (E22) e enchimento da lavadora", "Pressão e vazão disponíveis"],
        ["Drenagem industrial", "Só a área S5 tem ralo: CN-1 → dreno → CI → CX-1 estanque", conf("Destino do efluente (P09)")],
        ["Ventilação e exaustão", "Ventilação natural pelos vãos superiores das laterais; duto EX-1 da lavadora a ~3 m até os "
                                  "fundos; coletor de névoa no sopro (E10); ar-condicionado no escritório/metrologia",
         conf("Fechamento dos vãos (D12); vazão de exaustão do fabricante")],
        ["Prevenção contra incêndio", "Extintores (posições indicativas), iluminação de emergência, sinalização, rotas livres",
         conf("Projeto de PCI conforme CBMMG (P18)")]], larguras=["16%", "56%", "28%"])}
"""


def s13_seguranca():
    riscos = [[r, e(p), e(o), str(pp), str(ss), f"<b>{e(nv)}</b>", e(c), e(nr)] for r, p, o, pp, ss, nv, c, nr in C.RISCOS]
    return f"""
<h2>13. Segurança, meio ambiente e conformidade</h2>
<p>As normas abaixo foram consideradas pela pertinência à operação; nem todas se aplicam automaticamente. As versões
vigentes devem ser verificadas na data do projeto executivo.</p>
{tabela(["Referência", "Tema", "Como se aplica aqui"], [[e(a), e(b), e(c)] for a, b, c in C.NORMAS], cls="pequena",
        larguras=["17%", "31%", "52%"])}
<section class="paisagem">
<h3>13.1 Matriz de riscos e medidas preventivas</h3>
<p class="nota">{e(C.ESCALA_RISCO)}</p>
{tabela(["", "Perigo", "Onde", "P", "S", "Nível", "Medidas de controle", "Referência"], riscos, cls="mini",
        larguras=["4%", "17%", "11%", "3%", "3%", "6%", "45%", "11%"])}
</section>
<h3>13.2 Responsabilidades técnicas</h3>
{tabela(["Item", "Profissional / responsável", "Documento"], [
        ["Base dos trilhos, laje, contenção, fundações (e estrutura da ponte rolante, se Alt. 2)", "Engenheiro civil/estrutural", "Projeto + ART"],
        ["Pórtico, talha e acessórios de içamento", "Fabricante + engenheiro mecânico", "Projeto/certificado, manual, ART, inspeções"],
        ["Instalações elétricas e SPDA", "Engenheiro eletricista", "Projeto + ART; prontuário NR-10 quando aplicável"],
        ["Drenagem, efluentes e resíduos", "Engenheiro (sanitário/ambiental) + meio ambiente da usina", "Projeto + ART; licença/autorização"],
        ["Exaustão e ventilação", "Engenheiro mecânico", "Projeto + ART"],
        ["Prevenção contra incêndio", "Profissional habilitado", "Projeto aprovado no CBMMG"],
        ["Apreciação de riscos de máquinas, PGR, ergonomia", "SESMT (engenharia de segurança)", "NR-12, NR-01, NR-17 (AEP/AET)"]],
        larguras=["44%", "30%", "26%"])}
"""


def s14_areas():
    l1 = [[e(i), e(nm), n(a, 2), n(p, 1) + "%"] for i, nm, z, a, p in Q1]
    import pranchas as P
    cat = [[e(nm), n(a, 2), n(a / G.AREA_INTERNA * 100, 1) + "%", e(obs)] for nm, a, obs in P.linhas_categorias(1)]
    cat.append(["<b>Total (área interna)</b>", f"<b>{n(G.AREA_INTERNA, 2)}</b>", "<b>100%</b>", "Sem sobreposição"])
    return f"""
<h2>14. Quadro de áreas (Alternativa 1)</h2>
<p>Base: área interna entre as faces das paredes, {n(G.AREA_INTERNA, 2)} m² (área construída {n(G.AREA_CONSTRUIDA, 2)} m²;
entre eixos {n(G.AREA_EIXOS, 2)} m²). Cada setor é um retângulo sem sobreposição; a circulação é a diferença entre a área
interna e a soma dos setores, por isso não há dupla contagem. O escritório entra pela área interna da sala; suas
paredes aparecem em linha própria.</p>
{tabela(["Categoria", "m²", "% da área interna", "Observações"], cat, larguras=["32%", "12%", "14%", "42%"])}
{tabela(["", "Setor", "m²", "%"], l1, cls="pequena", larguras=["7%", "69%", "12%", "12%"])}
<p>Área produtiva (S1–S10): {n(PROD1, 2)} m² na Alternativa 1 e {n(PROD2, 2)} m² na Alternativa 2.</p>
"""


def s15_verificacao():
    linhas = []
    for nalt in (1, 2):
        for item, ok, det in V.verificar(nalt):
            linhas.append([f"Alt. {nalt}", e(item), "<b>OK</b>" if ok else '<b class="falha">FALHA</b>', e(det)])
    return f"""
<h2>15. Verificação de consistência</h2>
<p>As verificações abaixo são calculadas a partir das mesmas coordenadas usadas nas pranchas e no CAD
(<code>verificacao.py</code>):</p>
{tabela(["", "Verificação", "Resultado", "Detalhe"], linhas, cls="pequena", larguras=["7%", "40%", "9%", "44%"])}
<p><b>O que ainda não pode ser verificado</b> (depende das pendências): capacidade de carga (P01, P04), alcance real do
gancho e largura do truque (P02), compatibilidade da lavadora escolhida com o envelope (P11), viabilidade dos
embutidos (P03), cargas elétricas (P07) e destino dos efluentes (P09). Enquanto essas pendências estiverem abertas,
o layout não está liberado para execução.</p>
"""


def s16_plano():
    return f"""
<h2>16. Plano de execução por etapas</h2>
<p class="nota">Sem prazos: dependem das respostas às pendências e da contratação.</p>
{tabela(["Etapa", "Entregáveis", "Responsáveis sugeridos", "Depende de", "Riscos", "Critério de aprovação"],
        [[f"<b>{e(a)}</b>", e(b), e(c), e(d), e(r), e(k)] for a, b, c, d, r, k in C.PLANO], cls="pequena",
        larguras=["12%", "30%", "16%", "9%", "15%", "18%"])}
"""


def s17_orcamento():
    return f"""
<h2>17. Orçamento preliminar</h2>
<p>{e(C.ORCAMENTO_BASE)}</p>
{tabela(["Categoria", "Itens e quantidades do layout", "Valor"], [[f"<b>{e(a)}</b>", e(b), "a cotar"] for a, b in C.ORCAMENTO],
        larguras=["24%", "64%", "12%"])}
"""


def s18_proxima():
    return f"""
<h2>18. Recomendações para a próxima fase</h2>
{lista(["Responder as pendências bloqueantes (P01–P04, P07, P09, P16, P18) e conferir as cotas em campo.",
        "Decidir entre Alternativa 1 e 2 <b>antes de concretar o piso</b>; em seguida, congelar os embutidos.",
        "Pedir ao fornecedor do pórtico a geometria do truque, a aproximação do gancho e a altura útil; recalcular a faixa útil.",
        "Especificar a lavadora com a maior peça real e pedir propostas com envelope, utilidades e destino do banho.",
        "Contratar os projetos complementares (civil/piso, elétrico, drenagem/efluentes, exaustão, PCI) com ART.",
        "Incluir a oficina no PGR, fazer a apreciação de riscos NR-12 e a AEP (NR-17).",
        "Definir procedimentos: recebimento e OS, içamento, lavagem, segregação por cor, quarentena, expedição, resíduos."])}
<h3>Verificações necessárias antes da execução</h3>
{lista([f"☐ {e(c)}" for c in C.CHECKLIST])}
<h3>Arquivos do estudo</h3>
{tabela(["Arquivo", "Conteúdo"], [
        ["Pranchas_Sala_Motores_OUT2026.pdf", "SM-01 existente e interferências; SM-02 planta proposta (Desenho 1); SM-03 cotas e corte "
         "(Desenho 2); SM-04 fluxo (Desenho 3); SM-05 lavagem (Desenho 4); SM-06 infraestrutura (Desenho 5); SM-07 Alternativa 2"],
        ["build/SM-0x.svg", "As mesmas pranchas em vetor editável (Inkscape, Illustrator)"],
        ["Planta_Sala_Motores_OUT2026.dxf", "CAD em mm (1:1), por camadas: existente, Alt. 1 com cotas, fluxos e infraestrutura, Alt. 2"],
        ["Quadros_Sala_Motores_OUT2026.xlsx", "Áreas, equipamentos, riscos, pendências, premissas, normas, plano, orçamento, verificação"],
        ["Memorial_Tecnico_Sala_Motores_OUT2026.pdf", "Este documento"]], larguras=["38%", "62%"])}
"""


CSS = """
@page { size: A4; margin: 15mm 14mm 16mm 14mm;
  @bottom-left { content: "Oficina de motores e transmissões — memorial técnico — estudo conceitual, não executivo";
                 font: 7pt 'Segoe UI', Arial; color: #777; }
  @bottom-right { content: "Página " counter(page) " de " counter(pages); font: 7pt 'Segoe UI', Arial; color: #777; } }
@page paisagem { size: A4 landscape; margin: 12mm 12mm 14mm 12mm; }
.paisagem { page: paisagem; }
body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 9.2pt; color: #1d2733; line-height: 1.38; }
h1 { font-size: 26pt; color: """ + NAVY + """; margin: 30mm 0 6mm; line-height: 1.15; }
h2 { font-size: 13.5pt; color: """ + NAVY + """; border-bottom: 2px solid """ + NAVY + """; padding-bottom: 2px;
     margin: 9mm 0 3mm; break-after: avoid; }
h2:not(:first-child) { break-before: page; }
.paisagem h2, .paisagem h3 { break-before: auto; }
h3 { font-size: 10.5pt; color: #2c5282; margin: 5mm 0 2mm; break-after: avoid; }
p { margin: 0 0 2.4mm; text-align: justify; }
ul { margin: 0 0 3mm 5mm; padding-left: 4mm; } li { margin-bottom: 1.2mm; }
table { width: 100%; border-collapse: collapse; margin: 1.5mm 0 4mm; font-size: 8.2pt; table-layout: fixed; }
th { background: """ + NAVY + """; color: #fff; text-align: left; padding: 3px 4px; font-weight: 600; }
td { border-bottom: 1px solid #d7dde5; padding: 3px 4px; vertical-align: top; overflow-wrap: anywhere; }
tr { break-inside: avoid; } thead { display: table-header-group; }
table.pequena { font-size: 7.6pt; } table.mini { font-size: 6.8pt; } table.mini td, table.mini th { padding: 2px 3px; }
table.kv td:first-child { width: 28%; font-weight: 600; color: #4a5568; }
figure { margin: 2mm 0 5mm; text-align: center; break-inside: avoid; }
figure img { border: 1px solid #d7dde5; }
figcaption { font-size: 7.8pt; color: #4a5568; margin-top: 1.5mm; }
.conf { color: #c2185b; } .falha { color: #b3302d; }
.nota { font-size: 8pt; color: #4a5568; }
.capa .faixa { background: #b3302d; color: #fff; font-weight: 700; padding: 3mm 4mm; letter-spacing: .5px; }
.capa .sub { font-size: 11.5pt; color: #4a5568; margin-bottom: 8mm; text-align: left; }
.capa .aviso { margin-top: 10mm; padding: 4mm; background: #f4f6f9; border-left: 4px solid #b3302d; }
code { font-size: 8pt; background: #f1f3f6; padding: 0 2px; }
"""


def montar():
    corpo = "".join(f() for f in (capa, s1_resumo, s2_escopo, s3_desenho, s4_premissas, s5_alternativas, s6_layout,
                                  s7_fluxo, s8_lavagem, s9_icamento, s10_bancadas, s11_equipamentos, s12_infra,
                                  s13_seguranca, s14_areas, s15_verificacao, s16_plano, s17_orcamento, s18_proxima))
    return f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<title>Memorial técnico — Oficina de motores e transmissões</title><style>{CSS}</style></head>
<body>{corpo}</body></html>"""


# ---------------------------------------------------------------- planilha
def planilha():
    import pranchas as P
    wb = openpyxl.Workbook()
    cab_fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))

    def aba(nome, cab, linhas, larguras):
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for ln in linhas:
            ws.append(list(ln))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cab_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for i, w in enumerate(larguras, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    wb.remove(wb.active)
    linhas = []
    for nalt, q in ((1, Q1), (2, Q2)):
        for i, nm, z, a, p in q:
            linhas.append((f"Alt. {nalt}", i, nm, z, round(a, 2), round(p / 100, 4)))
    ws = aba("Quadro de áreas", ["Alternativa", "Setor", "Nome", "Zona", "Área (m²)", "% da área interna"], linhas,
             [12, 8, 60, 12, 12, 16])
    for c in ws["F"][1:]:
        c.number_format = "0.0%"
    aba("Áreas por categoria", ["Categoria", "Área (m²)", "Observações"],
        [(nm, round(a, 2), obs) for nm, a, obs in P.linhas_categorias(1)], [36, 12, 60])
    eq = []
    for t, (nome, func, qb, dim, esp, util, req, rest, prio, lev) in LY.EQUIP.items():
        eq.append((t, nome, func, QTD1.get(t, qb), sum(1 for x in LY.A2_EQUIP if x[0] == t) or (1 if t == "E01" else 0),
                   ", ".join(sorted(SETOR_DE.get(t, []))), dim, esp, util, req, rest, prio, lev))
    aba("Equipamentos", ["Tag", "Equipamento", "Função", "Qtd Alt. 1", "Qtd Alt. 2", "Setor (Alt. 1)",
                         "Dimensão de referência (mm)", "Espaço operacional", "Utilidades", "Requisitos de instalação",
                         "Restrições de segurança", "Prioridade", "Dados a levantar"], eq,
        [6, 34, 34, 9, 9, 10, 26, 28, 28, 30, 34, 12, 40])
    pos = [(t, alt, r[0], r[1], r[2], r[3], r[1] - r[0], r[3] - r[2])
           for alt, lst in (("Alt. 1", LY.A1_EQUIP), ("Alt. 2", LY.A2_EQUIP)) for t, r, _ in lst]
    aba("Posições (mm)", ["Tag", "Alternativa", "X de", "X até", "Y de", "Y até", "Comprimento", "Largura"], pos,
        [8, 12, 10, 10, 10, 10, 12, 10])
    aba("Riscos", ["ID", "Perigo", "Onde", "P", "S", "Nível", "Medidas de controle", "Referência"], C.RISCOS,
        [6, 34, 18, 5, 5, 8, 80, 16])
    aba("Pendências", ["ID", "Informação que falta", "Por que importa", "Quem levanta", "Bloqueia execução"],
        C.PENDENCIAS, [6, 60, 46, 24, 12])
    aba("Premissas", ["ID", "Premissa"], C.PREMISSAS, [6, 120])
    aba("Interferências e divergências", ["ID", "Descrição", "Quem resolve"],
        [(i, t, "") for i, t in G.INTERFERENCIAS] + [(d, t, q) for d, t, q in G.DIVERGENCIAS], [6, 110, 28])
    aba("Dimensões do galpão", ["Item", "Valor (mm)", "Fonte"],
        [(k, v, f) for k, (v, f) in G.ALTURAS.items()] +
        [("Comprimento entre eixos", G.L, "cota"), ("Largura entre eixos", G.B, "cota"), ("Parede", G.PAREDE, "cota"),
         ("Vão dos trilhos", G.TRILHOS_Y[1] - G.TRILHOS_Y[0], "cota"), ("Bloco de fundação (planta)", G.BLOCO, "medido"),
         ("Largura da perna do pórtico", G.PERNA_LARGURA, "medido (premissa)")], [44, 14, 60])
    aba("Normas", ["Referência", "Tema", "Como se aplica"], C.NORMAS, [24, 50, 90])
    aba("Plano de execução", ["Etapa", "Entregáveis", "Responsáveis", "Depende de", "Riscos", "Critério de aprovação"],
        C.PLANO, [26, 60, 30, 14, 30, 40])
    aba("Orçamento", ["Categoria", "Itens e quantidades do layout", "Valor (R$)", "Data-base", "Impostos", "Frete",
                      "Instalação", "Contingência", "Fonte da cotação"],
        [(a, b, "a cotar", "", "", "", "", "", "") for a, b in C.ORCAMENTO], [30, 90, 12, 12, 10, 10, 10, 12, 20])
    aba("Verificação", ["Alternativa", "Verificação", "OK", "Detalhe"],
        [(f"Alt. {nalt}", i, "OK" if ok else "FALHA", d) for nalt in (1, 2) for i, ok, d in V.verificar(nalt)],
        [12, 60, 8, 70])
    aba("Checklist pré-execução", ["Item", "Verificado (S/N)", "Responsável", "Data"],
        [(c, "", "", "") for c in C.CHECKLIST], [100, 16, 20, 12])
    wb.save(XLSX)


def main():
    htm = BUILD / "memorial.html"
    htm.write_text(montar(), encoding="utf-8")
    tmp = BUILD / "memorial.pdf"
    tmp.unlink(missing_ok=True)
    subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={tmp}", htm.resolve().as_uri()], check=True, timeout=180)
    tmp.replace(PDF)  # falha se o PDF estiver aberto em outro programa
    planilha()
    print(PDF, XLSX)


if __name__ == "__main__":
    main()
