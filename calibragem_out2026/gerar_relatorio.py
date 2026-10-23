"""Gera o relatório da diretoria (PDF, A4 paisagem) e a base por pneu (xlsx).

python gerar_relatorio.py  ->  Relatorio_Calibragem_Reboques_OUT2026.pdf
                               Base_Auditoria_Calibragem_OUT2026.xlsx
Gráficos em matplotlib, páginas em HTML, PDF impresso pelo Edge em modo headless.
"""
import base64
import collections
import statistics as st
import subprocess
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import analise  # noqa: E402

AQUI = Path(__file__).parent
BUILD = AQUI / "build"
BUILD.mkdir(exist_ok=True)
PDF = AQUI / "Relatorio_Calibragem_Reboques_OUT2026.pdf"
XLSX = AQUI / "Base_Auditoria_Calibragem_OUT2026.xlsx"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

NAVY, RED, ORANGE, GREEN, GRAY = "#1f3d63", "#b3302d", "#d9902f", "#1e6b4a", "#7d8ea3"
CLASSES = ["< 105", "105 a 114", "115 a 125", "> 125", "sem leitura"]
COR = {"< 105": RED, "105 a 114": ORANGE, "115 a 125": GREEN, "> 125": "#5b6fa8", "sem leitura": "#b8c0ca"}
ROTULO = {"< 105": "Abaixo de 105 psi", "105 a 114": "105 a 114 psi", "115 a 125": "115 a 125 psi (padrão)",
          "> 125": "Acima de 125 psi", "sem leitura": "Sem leitura"}

plt.rcParams.update({"font.size": 8, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.edgecolor": "#c9d0d8", "axes.grid": True, "grid.color": "#e8ecf0",
                     "axes.axisbelow": True, "legend.frameon": False})


# ---------------------------------------------------------------- formatação
def n(x, d=0):
    s = f"{x:,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(a, b, d=1):
    return n(a / b * 100, d) + "%"


def img(fig, nome):
    caminho = BUILD / f"{nome}.png"
    fig.savefig(caminho, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return "data:image/png;base64," + base64.b64encode(caminho.read_bytes()).decode()


# ---------------------------------------------------------------- dados
P, D, OBS = analise.main()
MED = [p for p in P if p["psi"] is not None]
CFGS = ["9 eixos", "11 eixos"]
TOTAL = len(P)
FROTAS = list(dict.fromkeys(p["frota"] for p in P))
CONJ = list(analise.CONJUNTOS)


def sel(cfg=None, medidos=False):
    base = MED if medidos else P
    return [p for p in base if cfg is None or p["config"] == cfg]


cont = {cfg: collections.Counter(p["classe"] for p in sel(cfg)) for cfg in CFGS + [None]}
conf = cont[None]["115 a 125"]
media = st.mean(p["psi"] for p in MED)
sem = cont[None]["sem leitura"]
exato = sum(1 for p in MED if p["psi"] == analise.PADRAO)
abaixo115 = sum(1 for p in MED if p["psi"] < 115)
sem_int = sum(1 for p in P if p["psi"] is None and p["interno"])

DD = [d for d in D if "dif" in d]
dup10 = sum(d["dif"] >= 10 for d in DD)
dup20 = sum(d["dif"] >= 20 for d in DD)
dup_meio = sum(1 for d in D if (d["ext"] is None) != (d["int"] is None))

# por conjunto
CJ = []
for c, (fab, ano, cfg, viag, kmed, kmax, suc, est, taxa) in analise.CONJUNTOS.items():
    s = [p for p in P if p["conjunto"] == c]
    m = [p for p in s if p["psi"] is not None]
    CJ.append({
        "conj": c, "fab": fab, "ano": ano, "cfg": cfg, "cfg2": "9 eixos" if "9" in cfg else "11 eixos",
        "kmed": kmed, "kmax": kmax, "suc": suc, "est": est, "taxa": taxa, "pneus": len(s),
        "aud": "parcial (só 63076)" if c == "63075 + 63076" else "completo",
        "media": st.mean(p["psi"] for p in m), "sem": len(s) - len(m),
        "conf": sum(1 for p in m if 115 <= p["psi"] <= 125),
        "psi_med": analise.pressao_para(kmed), "psi_max": analise.pressao_para(kmax),
        "sob_med": sum(p["sobrecarga_med"] for p in m), "sob_max": sum(p["sobrecarga_max"] for p in m),
    })
for c in CJ:
    c["sob_med_pct"] = c["sob_med"] / c["pneus"] * 100
    c["sob_max_pct"] = c["sob_max"] / c["pneus"] * 100
corr_media = np.corrcoef([c["media"] for c in CJ], [c["suc"] for c in CJ])[0, 1]
corr_sob = np.corrcoef([c["sob_med_pct"] for c in CJ], [c["suc"] for c in CJ])[0, 1]
nove = [c for c in CJ if c["cfg2"] == "9 eixos"]
corr_sob9 = np.corrcoef([c["sob_med_pct"] for c in nove], [c["suc"] for c in nove])[0, 1]

sob_med9 = sum(p["sobrecarga_med"] for p in sel("9 eixos", True))
sob_max9 = sum(p["sobrecarga_max"] for p in sel("9 eixos", True))
sob_med11 = sum(p["sobrecarga_med"] for p in sel("11 eixos", True))
uso9 = st.mean(p["uso_med"] for p in sel("9 eixos", True))
uso11 = st.mean(p["uso_med"] for p in sel("11 eixos", True))
G9, G11 = analise.GERAL["9 eixos"], analise.GERAL["11 eixos"]
req = {k: analise.pressao_para(v) for k, v in
       {"9med": G9["kg_med"], "9max": G9["kg_max"], "11med": G11["kg_med"], "11max": G11["kg_max"]}.items()}
cap = analise.capacidade


# ---------------------------------------------------------------- gráficos
def g_distribuicao():
    fig, ax = plt.subplots(figsize=(7.4, 3.2))
    valores = sorted(set(p["psi"] for p in MED))
    base = np.zeros(len(valores))
    for cfg, cor in (("11 eixos", NAVY), ("9 eixos", RED)):
        q = [sum(1 for p in sel(cfg, True) if p["psi"] == v) for v in valores]
        ax.bar(range(len(valores)), q, bottom=base, color=cor, width=0.75,
               label=f"{cfg} ({len(sel(cfg, True))} medidos)")
        base += q
    for i, t in enumerate(base):
        ax.text(i, t + 1, int(t), ha="center", fontsize=7, fontweight="bold", color="#333")
    i0, i1 = valores.index(115), valores.index(125)
    ax.axvspan(i0 - 0.45, i1 + 0.45, color=GREEN, alpha=0.08, zorder=0)
    ax.text((i0 + i1) / 2, max(base) * 1.02, "faixa 115–125 psi", ha="center", color=GREEN, fontsize=7.5)
    ax.set_xticks(range(len(valores)), [n(v) for v in valores])
    ax.set_xlabel("Pressão medida (psi)")
    ax.set_ylabel("Pneus")
    ax.set_ylim(0, max(base) * 1.12)
    ax.legend(loc="upper left")
    return img(fig, "01_distribuicao")


def g_mapa():
    ordem = []
    for c in CJ:
        for f in c["conj"].split(" + "):
            if f in FROTAS:
                ordem.append((f, c["conj"], c["cfg2"]))
    ordem.append(("63105", "fora da base de pesagem", "11 eixos"))
    cols = [(e, pos) for e in range(1, 6) for pos in ("EE", "EI", "DI", "DE")]
    fig, ax = plt.subplots(figsize=(11.6, 5.0))
    ax.grid(False)
    idx = {(p["frota"], p["eixo"], p["pos"]): p for p in P}
    for r, (f, cj, cfg) in enumerate(ordem):
        y = len(ordem) - 1 - r
        for c, (e, pos) in enumerate(cols):
            p = idx.get((f, e, pos))
            x = c + (e - 1) * 0.35
            if p is None:
                continue
            ax.add_patch(plt.Rectangle((x, y), 0.94, 0.86, color=COR[p["classe"]], lw=0))
            txt = n(p["psi"]) if p["psi"] is not None else ("S/P" if p["leitura"] == "S/PINO" else
                                                           "S/F" if p["leitura"] == "S/FUNCIONAR" else "–")
            ax.text(x + 0.47, y + 0.42, txt, ha="center", va="center", fontsize=6.6,
                    color="#333" if p["classe"] == "sem leitura" else "white", fontweight="bold")
        ax.text(-0.25, y + 0.42, f"{f}", ha="right", va="center", fontsize=7.5, fontweight="bold")
        ax.text(-2.3, y + 0.42, cfg, ha="right", va="center", fontsize=6.8, color="#666")
    for e in range(1, 6):
        x0 = (e - 1) * 4 + (e - 1) * 0.35
        ax.text(x0 + 2, len(ordem) + 0.75, f"{e}º eixo", ha="center", fontsize=7.5, fontweight="bold")
        for k, pos in enumerate(("EE", "EI", "DI", "DE")):
            ax.text(x0 + k + 0.47, len(ordem) + 0.1, pos, ha="center", fontsize=6.3, color="#666")
    ax.set_xlim(-4.4, 21.6)
    ax.set_ylim(-0.3, len(ordem) + 1.35)
    ax.axis("off")
    hs = [plt.Rectangle((0, 0), 1, 1, color=COR[k]) for k in CLASSES]
    ax.legend(hs, [ROTULO[k] for k in CLASSES], loc="lower center", ncol=5, bbox_to_anchor=(0.5, -0.07),
              fontsize=7)
    return img(fig, "02_mapa")


def g_duplos():
    fig, ax = plt.subplots(figsize=(5.6, 3.0))
    faixas = [("0", lambda d: d == 0), ("5", lambda d: d == 5), ("10 a 15", lambda d: 10 <= d <= 15),
              ("20 a 25", lambda d: 20 <= d <= 25), ("30 ou +", lambda d: d >= 30)]
    q = [sum(1 for d in DD if f(d["dif"])) for _, f in faixas]
    cores = [GREEN, GREEN, ORANGE, RED, RED]
    ax.bar([r for r, _ in faixas], q, color=cores, width=0.65)
    for i, v in enumerate(q):
        ax.text(i, v + 0.6, v, ha="center", fontweight="bold", fontsize=8)
    ax.set_xlabel("Diferença de pressão entre os dois pneus do rodado duplo (psi)")
    ax.set_ylabel("Rodados duplos")
    ax.set_ylim(0, max(q) * 1.18)
    return img(fig, "03_duplos")


def g_curva():
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    xs = np.linspace(65, 123.3, 200)
    ax.plot(xs, [cap(x) for x in xs], color=NAVY, lw=2.2)
    ax.text(73, 2470, "capacidade do pneu (ETRTO)", color=NAVY, fontsize=7, fontweight="bold")
    ax.scatter(analise.TAB_PSI, analise.TAB_KG, color=NAVY, s=14, zorder=3)
    linhas = [(G9["kg_max"], RED, "--", f"9 eixos máx. {n(G9['kg_max'])} kg", req["9max"]),
              (G9["kg_med"], RED, ":", f"9 eixos méd. {n(G9['kg_med'])} kg", req["9med"]),
              (G11["kg_max"], GRAY, "--", f"11 eixos máx. {n(G11['kg_max'])} kg", req["11max"]),
              (G11["kg_med"], GRAY, ":", f"11 eixos méd. {n(G11['kg_med'])} kg", req["11med"])]
    for kg, cor, ls, txt, psi in linhas:
        ax.axhline(kg, color=cor, ls=ls, lw=1.1)
        ax.plot([psi, psi], [1700, kg], color=cor, lw=0.8, ls=ls)
        ax.scatter([psi], [kg], color=cor, s=28, zorder=4)
        x_txt, ha = (66, "left") if cor == RED else (99.5, "left")
        ax.text(x_txt, kg + 22, txt, color=cor, fontsize=7.2, fontweight="bold", ha=ha)
        ax.text(psi + 0.6, 1720, f"{n(psi)} psi", color=cor, fontsize=7, rotation=90, va="bottom")
    ax.axvline(120, color=GREEN, lw=1.4)
    ax.text(120.4, 3200, "padrão 120 psi", color=GREEN, fontsize=7.5, fontweight="bold")
    ax.set_xlim(64, 127)
    ax.set_ylim(1700, 3300)
    ax.set_xlabel("Pressão a frio (psi)")
    ax.set_ylabel("Carga por pneu (kg)")
    return img(fig, "04_curva")


def g_uso():
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    bins = np.arange(50, 135, 5)
    for cfg, cor in (("11 eixos", NAVY), ("9 eixos", RED)):
        v = [p["uso_med"] for p in sel(cfg, True)]
        ax.hist(v, bins=bins, color=cor, alpha=0.85, label=f"{cfg} — média {n(st.mean(v))}%", edgecolor="white")
    ax.axvline(100, color="#333", ls="--", lw=1)
    ax.text(101, 42, "acima de 100% =\ncarga maior que a\ncapacidade na pressão medida",
            fontsize=7, color="#333", va="top")
    ax.set_xlabel("Carga média do conjunto ÷ capacidade do pneu na pressão medida (%)")
    ax.set_ylabel("Pneus")
    ax.set_ylim(0, 88)
    ax.legend(loc="upper right")
    return img(fig, "05_uso")


def g_sucata():
    fig, axs = plt.subplots(1, 2, figsize=(11.4, 3.4))
    for ax, chave, titulo, r in ((axs[0], "media", "Pressão média medida (psi)", corr_media),
                                 (axs[1], "sob_med_pct", "% de pneus com pressão insuficiente para a carga média", corr_sob)):
        for c in CJ:
            cor = RED if c["cfg2"] == "9 eixos" else NAVY
            ax.scatter(c[chave], c["suc"], s=60, color=cor, edgecolor="white", zorder=3)
            if c["suc"] >= 4 or c["conj"] in ("63075 + 63076", "63037 + 63025", "63021 + 63020"):
                ax.annotate(c["conj"], (c[chave], c["suc"]), xytext=(5, 3), textcoords="offset points", fontsize=6.5)
        ax.set_xlabel(titulo)
        ax.set_ylabel("Pneus perdidos (15/07 a 25/09)")
        ax.set_title(f"Correlação: {n(r, 2)}", fontsize=8.5, color="#333")
    axs[0].scatter([], [], color=RED, label="9 eixos")
    axs[0].scatter([], [], color=NAVY, label="11 eixos")
    axs[0].legend(loc="upper left")
    fig.tight_layout(w_pad=3)
    return img(fig, "06_sucata")


# ---------------------------------------------------------------- html
CSS = """
@page { size: A4 landscape; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; font-family: Arial, Helvetica, sans-serif; color: #1d2733; }
.page { width: 297mm; height: 210mm; position: relative; page-break-after: always; overflow: hidden; }
.page:last-child { page-break-after: auto; }
.hdr { background: #142a47; color: #fff; height: 21mm; padding: 5.5mm 11mm 0 11mm; position: relative; }
.hdr h1 { margin: 0; font-size: 17.5pt; letter-spacing: .2px; }
.hdr .sub { font-size: 10pt; color: #c9d3e0; margin-top: 2.3mm; }
.hdr .r1 { position: absolute; right: 11mm; top: 7mm; font-size: 9pt; color: #c9d3e0; }
.hdr .r2 { position: absolute; right: 11mm; top: 14.6mm; font-size: 9pt; color: #c9d3e0; }
.body { padding: 5.5mm 11mm 0 11mm; }
.ftr { position: absolute; bottom: 5mm; left: 11mm; right: 11mm; border-top: 1px solid #dde2e8;
       padding-top: 1.5mm; font-size: 7.4pt; color: #6b7785; }
.ftr span { float: right; }
.kpis { display: flex; gap: 4mm; margin-bottom: 5mm; }
.kpi { flex: 1; border: 1px solid #d5dbe2; border-top: 5px solid #142a47; border-radius: 6px;
       padding: 4mm 5mm 3.5mm 5mm; background: #fff; }
.kpi.red { border-color: #b3302d; border-top-color: #b3302d; background: #fbe9e8; }
.kpi.green { border-color: #1e6b4a; border-top-color: #1e6b4a; background: #e6f2ec; }
.kpi.orange { border-color: #d9902f; border-top-color: #d9902f; background: #fbf1e3; }
.kpi .t { font-size: 8.3pt; font-weight: bold; color: #6b7785; text-transform: uppercase; }
.kpi .v { font-size: 22pt; font-weight: bold; margin: 2.5mm 0 1.5mm 0; color: #142a47; }
.kpi.red .v { color: #b3302d; } .kpi.green .v { color: #1e6b4a; } .kpi.orange .v { color: #b8741c; }
.kpi .s { font-size: 8.2pt; color: #6b7785; display: flex; justify-content: space-between; align-items: center; }
.badge { white-space: nowrap; background: #142a47; color: #fff; font-weight: bold; border-radius: 4px; padding: 1.3mm 3mm; font-size: 8.8pt; }
.red .badge { background: #b3302d; } .green .badge { background: #1e6b4a; } .orange .badge { background: #d9902f; }
h2 { font-size: 12.5pt; margin: 0 0 2.5mm 0; border-left: 4.5mm solid #142a47; padding-left: 3.5mm; line-height: 6.5mm; }
h2.o { border-left-color: #d9902f; } h2.g { border-left-color: #1e6b4a; } h2.r { border-left-color: #b3302d; }
.row { display: flex; gap: 6mm; }
.col { flex: 1; min-width: 0; }
img.ch { width: 100%; display: block; }
table { border-collapse: collapse; width: 100%; font-size: 8.4pt; }
th { background: #142a47; color: #fff; padding: 1.6mm 2mm; text-align: right; font-weight: bold; border: 1px solid #fff; }
th:first-child, td:first-child { text-align: left; }
td { padding: 1.35mm 2mm; text-align: right; border: 1px solid #e1e6ec; }
tr:nth-child(even) td { background: #f5f7f9; }
tr.tot td { background: #e9edf1; font-weight: bold; }
td.l, th.l { text-align: left; } td.c, th.c { text-align: center; }
.small td, .small th { font-size: 7.6pt; padding: 1.05mm 1.6mm; }
.small td { white-space: nowrap; }
.wrap td:last-child, .wrap th:last-child { white-space: normal; text-align: left; }
.box { background: #eef1f4; border: 1px solid #d5dbe2; border-radius: 6px; padding: 3.5mm 5mm; font-size: 9pt; line-height: 1.5; }
.box.o { background: #fbf1e3; border-color: #d9902f; }
.box.r { background: #fbe9e8; border-color: #b3302d; }
.box.g { background: #e6f2ec; border-color: #1e6b4a; }
.box h3 { margin: 0 0 1.5mm 0; font-size: 10.5pt; color: #142a47; }
.box.o h3 { color: #c47a1b; } .box.r h3 { color: #b3302d; } .box.g h3 { color: #1e6b4a; }
.box p { margin: 0; } .box ol, .box ul { margin: 0; padding-left: 5mm; }
.note { font-size: 7.4pt; color: #6b7785; margin-top: 1.5mm; }
.dot { display: inline-block; width: 2.6mm; height: 2.6mm; border-radius: 50%; margin-right: 1.5mm; vertical-align: middle; }
b.r { color: #b3302d; } b.g { color: #1e6b4a; }
"""

TITULO = "REBOQUES CANAVIEIROS — AUDITORIA DE CALIBRAGEM"
FONTE = ("Fonte: CONFERENCIA DE CALIBRAGEM.xlsx (padrão 120 psi) · Relatório Peso por Pneu x Sucateamento "
         "(15/07 a 25/09/2026) · Relatório Pneu 295/80R22,5 · tabela ETRTO 295/80R22,5 148 (duplo).")
PAGINAS = []


def pagina(sub, aba, corpo):
    PAGINAS.append((sub, aba, corpo))


def montar():
    tot = len(PAGINAS)
    out = []
    for i, (sub, aba, corpo) in enumerate(PAGINAS, 1):
        out.append(f"""<div class="page"><div class="hdr"><h1>{TITULO}</h1><div class="sub">{sub}</div>
<div class="r1">CRV Industrial · Unidade Capinópolis/MG</div><div class="r2">{aba}</div></div>
<div class="body">{corpo}</div><div class="ftr">{FONTE}<span>Página {i} de {tot}</span></div></div>""")
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{''.join(out)}</body></html>"


def kpi(t, v, s, b, cls=""):
    return (f'<div class="kpi {cls}"><div class="t">{t}</div><div class="v">{v}</div>'
            f'<div class="s"><span>{s}</span><span class="badge">{b}</span></div></div>')


def tabela(cab, linhas, cls="", tot=None):
    h = "".join(f"<th>{c}</th>" for c in cab)
    corpo = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in ln) + "</tr>" for ln in linhas)
    if tot:
        corpo += '<tr class="tot">' + "".join(f"<td>{c}</td>" for c in tot) + "</tr>"
    return f'<table class="{cls}"><tr>{h}</tr>{corpo}</table>'


# ---------------------------------------------------------------- página 1: visão geral
def p1():
    linhas = []
    for k in CLASSES:
        a, b, t = cont["9 eixos"][k], cont["11 eixos"][k], cont[None][k]
        linhas.append([f'<span class="dot" style="background:{COR[k]}"></span>{ROTULO[k]}', a, b, t, pct(t, TOTAL)])
    tb = tabela(["Faixa", "9 eixos", "11 eixos", "Total", "%"], linhas,
                tot=["Total de posições", len(sel("9 eixos")), len(sel("11 eixos")), TOTAL, "100%"])
    m9 = st.mean(p["psi"] for p in sel("9 eixos", True))
    m11 = st.mean(p["psi"] for p in sel("11 eixos", True))
    corpo = f"""<div class="kpis">
{kpi("Posições auditadas", f"{TOTAL} pneus", "20 frotas · 10 conjuntos", f"{len(MED)} medidos")}
{kpi("Dentro do padrão (115–125 psi)", pct(conf, TOTAL), f"{conf} de {TOTAL} · em 120: {exato}", "meta 100%", "red")}
{kpi("Pressão média medida", f"{n(media, 1)} psi", f"Padrão: {analise.PADRAO} psi", f"{n(media - analise.PADRAO, 1)} psi", "red")}
{kpi("Sem leitura", f"{sem} pneus", f"{sem_int} deles em posição interna", pct(sem, TOTAL), "orange")}
</div>
<div class="row"><div class="col" style="flex:1.45"><h2>Pressão medida — todos os pneus (psi)</h2><img class="ch" src="{g_distribuicao()}"></div>
<div class="col"><h2>Conformidade por faixa</h2>{tb}
<p class="note">Faixa aceita: padrão 120 psi ± 5 psi. "Sem leitura" = célula com "-", "S/PINO" ou "S/FUNCIONAR".</p>
<div class="box r" style="margin-top:3mm"><h3>Leitura para a diretoria</h3>
<p>{conf} de {TOTAL} pneus ({pct(conf, TOTAL)}) estão na faixa; só {exato} ({pct(exato, TOTAL)}) estão em 120 psi.</p>
<p>O desvio é para baixo: {abaixo115} dos {len(MED)} pneus medidos ({pct(abaixo115, len(MED))}) estão abaixo de 115 psi; só {cont[None]['> 125']} acima de 125.</p>
<p>Média 9 eixos {n(m9, 1)} psi · 11 eixos {n(m11, 1)} psi — as duas configurações estão igualmente baixas.</p></div>
</div></div>"""
    pagina("Pressão medida x padrão de 120 psi — 280 posições de pneu em 20 frotas", "Visão Geral", corpo)


# ---------------------------------------------------------------- página 2: mapa
def p2():
    m_int = st.mean(p["psi"] for p in MED if p["interno"])
    m_ext = st.mean(p["psi"] for p in MED if not p["interno"])
    pior = min(FROTAS, key=lambda f: st.mean(p["psi"] for p in MED if p["frota"] == f))
    vp = [p["psi"] for p in MED if p["frota"] == pior]
    inteiras = [f for f in FROTAS if all(p["classe"] == "115 a 125" for p in P if p["frota"] == f)]
    minimo = min(MED, key=lambda p: p["psi"])
    obs = "; ".join(f"{f}: " + " / ".join(o) for f, o in OBS.items())
    corpo = f"""<h2>Mapa da calibragem — cada célula é um pneu (EE/DE = externos, EI/DI = internos do rodado duplo)</h2>
<img class="ch" src="{g_mapa()}" style="width:93%;margin:0 auto">
<div class="box" style="margin-top:2mm"><h3>Leitura para a diretoria</h3>
<p>{"Nenhuma frota tem" if not inteiras else str(len(inteiras)) + " frota(s) com"} todos os pneus na faixa. A de menor média é a {pior} ({n(st.mean(vp), 1)} psi, com pneus de {n(min(vp))} a {n(max(vp))} psi); o pneu mais baixo é o {minimo['eixo']}º {minimo['pos']} da {minimo['frota']}, com {n(minimo['psi'])} psi.
{sem} pneus sem leitura: {sem_int} em posição interna (EI/DI), onde o acesso à válvula é mais difícil — pneu que não é medido não é calibrado.
Média dos internos {n(m_int, 1)} psi e dos externos {n(m_ext, 1)} psi. Observação da planilha — {obs}.</p></div>"""
    pagina("Mapa por frota, eixo e posição — onde estão os pneus fora do padrão", "Mapa da Auditoria", corpo)


# ---------------------------------------------------------------- página 3: por frota
def p3():
    linhas = []
    for f in FROTAS:
        s = [p for p in P if p["frota"] == f]
        m = [p["psi"] for p in s if p["psi"] is not None]
        c = sum(1 for v in m if 115 <= v <= 125)
        dd = [d for d in DD if d["frota"] == f]
        linhas.append((c / len(s), [f, s[0]["conjunto"].replace(" (fora da base de pesagem)", " *"), s[0]["config"],
                                    max(p["eixo"] for p in s), len(s), n(st.mean(m), 1), n(min(m)), n(max(m)), c,
                                    sum(1 for v in m if v < 105), len(s) - len(m),
                                    f"{sum(d['dif'] >= 10 for d in dd)} de {len(dd)}", pct(c, len(s), 0)]))
    linhas = [ln for _, ln in sorted(linhas, key=lambda t: t[0])]
    tb = tabela(["Frota", "Conjunto", "Config. do conjunto", "Eixos do reboque", "Pneus", "Média", "Mín.", "Máx.", "115–125",
                 "&lt; 105", "Sem leit.", "Duplos ≥ 10 psi", "% conforme"], linhas, "small",
                tot=["Total", "", "", "", TOTAL, n(media, 1), n(min(p["psi"] for p in MED)),
                     n(max(p["psi"] for p in MED)), conf, cont[None]["< 105"], sem, f"{dup10} de {len(DD)}",
                     pct(conf, TOTAL, 0)])
    corpo = f"""<h2>Ranking por frota — ordenado pelo % de pneus dentro de 115–125 psi (menor para maior)</h2>{tb}
<p class="note">"Config. do conjunto" conta os eixos do cavalo mecânico + os dois reboques; "Eixos do reboque" conta só a frota da linha. Ex.: 63037 (2 eixos) + 63025 (4 eixos) = 6 eixos nos reboques + 3 do cavalo = 9 eixos; 63071 (3) + 63074 (5) = 8 + 3 = 11 eixos. Pneus = eixos do reboque × 4.<br>* 63105 não está entre os 65 conjuntos pesados de 15/07 a 25/09; tem 3 eixos como os semirreboques Randon e foi tratado como 11 eixos.
"Duplos ≥ 10 psi" = rodados duplos com diferença de 10 psi ou mais entre o pneu externo e o interno, entre os pares com as duas leituras.</p>"""
    pagina("Ranking por frota — conformidade, mínimo, sem leitura e rodado duplo", "Por Frota", corpo)


# ---------------------------------------------------------------- página 4: rodado duplo
def p4():
    piores = sorted(DD, key=lambda d: -d["dif"])[:10]
    tb = tabela(["Frota", "Eixo", "Lado", "Externo", "Interno", "Diferença"],
                [[d["frota"], f"{d['eixo']}º", "Esquerdo" if d["lado"] == "E" else "Direito", n(d["ext"]),
                  n(d["int"]), f"{n(d['dif'])} psi"] for d in piores], "small")
    corpo = f"""<div class="kpis">
{kpi("Rodados duplos auditados", f"{len(D)}", f"{len(DD)} com as duas leituras", "pares")}
{kpi("Diferença ≥ 10 psi no par", f"{dup10} pares", f"de {len(DD)} com as duas leituras", pct(dup10, len(DD)), "red")}
{kpi("Diferença ≥ 20 psi no par", f"{dup20} pares", "um pneu carrega bem mais que o outro", pct(dup20, len(DD)), "red")}
{kpi("Par com um pneu sem leitura", f"{dup_meio} pares", "um lado do duplo sem controle", pct(dup_meio, len(D)), "orange")}
</div>
<div class="row"><div class="col"><h2>Diferença de pressão dentro do rodado duplo</h2><img class="ch" src="{g_duplos()}"></div>
<div class="col"><h2>Top 10 · maiores diferenças</h2>{tb}</div></div>
<div class="box o" style="margin-top:4mm"><h3>Por que importa</h3>
<p>Os dois pneus do duplo dividem a carga do lado do eixo. Com pressões diferentes, o pneu mais murcho deflete mais e passa a carga para o vizinho:
um fica sobrecarregado e o outro trabalha deformado e aquecido. A boa prática é diferença máxima de 5 psi entre os dois — aqui {pct(dup10, len(DD), 0)} dos pares passam de 10 psi.</p></div>"""
    pagina("Rodado duplo — diferença de pressão entre o pneu externo e o interno do mesmo lado", "Rodado Duplo", corpo)


# ---------------------------------------------------------------- página 5: calibragem ideal
def p5():
    linhas = [
        ["9 eixos", n(G9["kg_med"]), n(G9["kg_max"]), f"{n(req['9med'])} psi", f"{n(req['9max'])} psi",
         "<b>120 psi</b>", "115 psi", f"{n(cap(120) / G9['kg_max'] * 100 - 100)}%"],
        ["11 eixos", n(G11["kg_med"]), n(G11["kg_max"]), f"{n(req['11med'])} psi", f"{n(req['11max'])} psi",
         "<b>110 psi</b>", "100 psi", f"{n(cap(110) / G11['kg_max'] * 100 - 100)}%"],
    ]
    tb = tabela(["Config.", "Kg/pn méd.", "Kg/pn máx.", "psi p/ méd.", "psi p/ máx.",
                 "Recomendado", "Mínimo", "Folga"], linhas, "small")
    cap_tab = tabela(["Pressão (psi)"] + [n(v) for v in (90, 100, 105, 110, 115, 120)],
                     [["Capacidade (kg/pneu)"] + [n(cap(v)) for v in (90, 100, 105, 110, 115, 120)]], "small")
    corpo = f"""<div class="row"><div class="col" style="flex:1.25"><h2>Capacidade do pneu x pressão e a carga real</h2>
<img class="ch" src="{g_curva()}">{cap_tab}</div>
<div class="col"><h2 class="g">Calibragem ideal pelo peso</h2>{tb}
<p class="note">Mín. = pressão que a tabela ETRTO exige para a carga (rodado duplo). Folga = capacidade na pressão recomendada ÷ carga máxima por pneu registrada na pesagem.</p>
<div class="box g" style="margin-top:3mm"><h3>Leitura para a diretoria</h3>
<p><b>9 eixos:</b> a viagem mais pesada ({n(G9['kg_max'])} kg/pneu) exige {n(req['9max'])} psi. O padrão de 120 psi está certo e <b class="r">não tem folga</b>:
a 110 psi o pneu sustenta {n(cap(110))} kg e a 100 psi, {n(cap(100))} kg — abaixo da carga média de muitos conjuntos.</p>
<p><b>11 eixos:</b> a carga máxima ({n(G11['kg_max'])} kg/pneu) exige {n(req['11max'])} psi. 110 psi já deixa {n(cap(110) / G11['kg_max'] * 100 - 100)}% de folga;
manter 120 psi como padrão único também atende.</p>
<p>Em resumo: o peso confirma o padrão de 120 psi para os 9 eixos; a baixa calibragem é crítica neles e tolerável nos 11 eixos.</p></div>
</div></div>"""
    pagina("Calibragem ideal pelo peso por pneu — tabela de carga x pressão do 295/80R22,5", "Calibragem Ideal", corpo)


# ---------------------------------------------------------------- página 6: pressão x carga por conjunto
def p6():
    linhas = []
    for c in sorted(CJ, key=lambda c: -c["kmed"]):
        linhas.append([c["conj"], c["cfg"], n(c["kmed"]), n(c["kmax"]), n(c["psi_med"]), n(c["psi_max"]),
                       n(c["media"], 1), f"{c['sob_med']} ({n(c['sob_med_pct'])}%)", f"{c['sob_max']} ({n(c['sob_max_pct'])}%)",
                       c["suc"]])
    tb = tabela(["Conjunto", "Config.", "Kg/pn méd.", "Kg/pn máx.", "psi p/ méd.", "psi p/ máx.", "psi medido",
                 "Insuf. méd.", "Insuf. máx.", "Pneus perdidos"], linhas, "small")
    corpo = f"""<div class="row"><div class="col" style="flex:0.9"><h2>Uso da capacidade na pressão medida</h2><img class="ch" src="{g_uso()}">
<div class="box" style="margin-top:2mm"><h3>Leitura para a diretoria</h3>
<p>Na pressão medida, o pneu dos 9 eixos usa em média {n(uso9)}% da capacidade só com a carga média; o dos 11 eixos, {n(uso11)}%.</p>
<p>9 eixos: {sob_med9} pneus já ficam sobrecarregados na carga média do conjunto e {sob_max9} na carga máxima ({pct(sob_max9, len(sel('9 eixos', True)), 0)} dos medidos).
11 eixos: {sob_med11} (o pneu de 50 psi da 63097).</p></div></div>
<div class="col" style="flex:1.25"><h2>Pressão necessária x medida, por conjunto</h2>{tb}
<p class="note">psi p/ méd. e p/ máx. = pressão mínima da tabela ETRTO para a carga por pneu média e máxima do conjunto.
Insuf. = pneus medidos com pressão abaixo dessa mínima (entre parênteses, % das posições do conjunto). Pneus perdidos = pneus sucateados no conjunto de 15/07 a 25/09/2026. 63075 + 63076: auditado só o 63076.</p></div></div>"""
    pagina("Pressão medida x carga real de cada conjunto — pneus rodando acima da capacidade", "Pressão x Carga", corpo)


# ---------------------------------------------------------------- página 7: calibragem x sucateamento
def p7():
    causas = [["Estouro por impacto", 281, "33%", "Direta: pneu murcho e sobrecarregado deflete mais e rompe a carcaça no impacto"],
              ["Aquecimento de roda", 117, "14%", "Indireta: a causa principal é freio/cubo; pneu murcho soma calor"],
              ["Furou e rodou vazio", 99, "12%", "Direta: falta de conferência diária — pneu vazio não detectado"],
              ["Roda travada - freio", 72, "9%", "Não relacionada à calibragem"],
              ["Desloc. no ombro", 41, "5%", "Direta: calor por baixa pressão/sobrecarga separa o ombro"],
              ["Saliência lateral", 30, "4%", "Parcial: flanco deformado fica exposto a cortes e impactos"]]
    tb = tabela(["Motivo (Abr–Set/26)", "Qtd", "%", "Relação com a calibragem"],
                [[a, b, c, d] for a, b, c, d in causas], "small wrap")
    pior = max(CJ, key=lambda c: c["suc"])
    corpo = f"""<h2>Calibragem x sucateamento — 10 conjuntos auditados</h2><img class="ch" src="{g_sucata()}" style="width:96%">
<div class="row" style="margin-top:2mm"><div class="col" style="flex:1.15">{tb}
<p class="note">Causas do sucateamento total do 295/80R22,5 (844 un, todos os equipamentos); reboque canavieiro = 462 un (55%).</p></div>
<div class="col"><div class="box r"><h3>A calibragem é o motivo do estouro?</h3>
<p>Sozinha, a pressão não explica: a pressão média do conjunto não tem relação com a quantidade de pneus perdidos (correlação {n(corr_media, 2)}).</p>
<p>O que se relaciona é <b>pressão abaixo do que a carga exige</b>: correlação {n(corr_sob, 2)} entre % de pneus com pressão insuficiente e pneus perdidos
({n(corr_sob9, 2)} só nos 9 eixos). O conjunto que mais perdeu, {pior['conj']} ({pior['suc']} pneus), tem {pior['sob_med']} pneus insuficientes já na carga média.</p>
<p>Conclusão: a calibragem baixa é <b>causa contribuinte</b> do estouro nos 9 eixos, que rodam no limite da carga; nos 11 eixos a mesma pressão tem folga.</p></div></div></div>"""
    pagina("A calibragem explica o estouro de pneu? — cruzamento com os pneus perdidos por conjunto", "Calibragem x Estouro", corpo)


# ---------------------------------------------------------------- página 8: plano de ação, metodologia, síntese
def p8():
    corpo = f"""<div class="row"><div class="col"><h2 class="g">Plano de ação</h2><div class="box g" style="font-size:8.6pt">
<ol>
<li><b>Imediato:</b> recalibrar as {TOTAL - conf} posições fora da faixa, começando pelos {cont[None]['< 105']} pneus abaixo de 105 psi e pelos 9 eixos.</li>
<li><b>Sem leitura:</b> instalar extensor de válvula nos pneus internos e repor válvula/bico nos "S/PINO" — {sem} pneus hoje sem controle.</li>
<li><b>Padrão por configuração:</b> 9 eixos 120 psi (mínimo 115); 11 eixos 110 a 120 psi (mínimo 100). Sempre a frio, antes do primeiro turno.</li>
<li><b>Rodado duplo:</b> diferença máxima de 5 psi entre externo e interno; casar pneus de mesmo diâmetro e vida.</li>
<li><b>Rotina:</b> conferência diária com manômetro nos 9 eixos e semanal nos 11 eixos; manômetros aferidos.</li>
<li><b>Indicador:</b> % de pneus em 115–125 psi por frota, semanal, meta ≥ 90%.</li>
<li><b>Próxima auditoria (30 dias):</b> registrar data, hora, pneu frio/quente e nº de fogo — para ligar pressão ao pneu sucateado.</li>
<li><b>Carga dos 9 eixos:</b> mesmo a 120 psi a viagem mais pesada usa {n(G9['kg_max'] / cap(120) * 100)}% da capacidade; avaliar teto de carga líquida nos conjuntos Verde/Rodofort.</li>
</ol></div></div>
<div class="col"><h2 class="o">Metodologia, premissas e limites</h2><div class="box o" style="font-size:8.1pt;line-height:1.42">
<ul>
<li>Faixa conforme: 120 psi ± 5 (115 a 125). Valores da planilha usados como lidos.</li>
<li>Capacidade: tabela ETRTO do 295/80R22,5 índice 148 em rodado duplo (3.150 kg a 8,5 bar/123 psi), interpolada; abaixo de 65 psi, extrapolada.</li>
<li>Carga por pneu de cada conjunto = média e máxima do relatório de peso por pneu (rateio uniforme, inclui a parte apoiada no cavalo — limite superior).</li>
<li>A data da auditoria e se a medição foi a frio ou a quente não constam na planilha. Pneu quente marca ~10% a mais: se medido quente, a pressão a frio é ainda menor.</li>
<li>A auditoria é uma foto; os pneus perdidos são de 15/07 a 25/09. Com 10 conjuntos, as correlações são indicativas, não prova.</li>
<li>Pneus perdidos em quantidade, sem ajuste pelo nº de viagens (a maioria fez 120 a 230; 63021 + 63020 fez só 29).</li>
<li>9 eixos e 11 eixos também diferem em idade e fabricante (2009–2023 x 2023–2024).</li>
<li>63105 fora da base de pesagem (tratado como 11 eixos); do conjunto 63075 + 63076 só o 63076 foi auditado.</li>
</ul></div></div></div>
<h2 style="margin-top:4mm">Síntese</h2><div class="box">
<p>Só {pct(conf, TOTAL)} dos {TOTAL} pneus estão em 115–125 psi; média {n(media, 1)} psi ({n(media - 120, 1)} psi do padrão); {sem} sem leitura; {pct(dup10, len(DD), 0)} dos duplos com diferença ≥ 10 psi.</p>
<p>Pelo peso: 9 eixos precisam de {n(req['9max'])} psi na carga máxima — padrão 120 psi correto e sem folga. 11 eixos precisam de {n(req['11max'])} psi — 110 psi bastam.</p>
<p>Calibragem baixa + carga alta: nos 9 eixos {sob_max9} pneus medidos rodam abaixo da pressão exigida pela carga máxima; a quantidade de pneus perdidos acompanha esse % (correlação {n(corr_sob, 2)}), não a pressão média ({n(corr_media, 2)}).</p>
<p>Conclusão: a calibragem fora do padrão contribui para o estouro, sobretudo nos 9 eixos; corrigir pressão e rodado duplo é a ação de menor custo — cada pneu poupado vale ~R$ 1.961 (preço médio Abr–Set/26).</p></div>"""
    pagina("Plano de ação, metodologia e síntese", "Ação & Síntese", corpo)


# ---------------------------------------------------------------- página 9: resumo do rodado duplo
def par(frota, eixo, lado):
    d = next(d for d in D if d["frota"] == frota and d["eixo"] == eixo and d["lado"] == lado)
    c = next(c for c in CJ if frota in c["conj"].split(" + "))
    alto, baixo = max(d["ext"], d["int"]), min(d["ext"], d["int"])
    carga = 2 * c["kmed"] * alto / (alto + baixo)  # divisão da carga do par na proporção da pressão
    return d, c, alto, baixo, carga


def p9():
    linhas = []
    for frota, eixo, lado, efeito in (
            ("63044", 3, "D", "os dois pneus acima do limite"),
            ("63097", 2, "E", "se o furado esvaziar, o outro leva o par inteiro"),
            ("63071", 2, "D", "peso dentro do limite; arrasto e desgaste irregular")):
        d, c, alto, baixo, carga = par(frota, eixo, lado)
        uso = carga / cap(alto) * 100
        if frota == "63097":
            efeito = f"se o furado esvaziar: {n(2 * c['kmed'])} kg num pneu de {n(cap(alto))} kg ({n(2 * c['kmed'] / cap(alto) * 100 - 100)}% acima)"
        linhas.append([f"{frota} · {eixo}º eixo {'esq.' if lado == 'E' else 'dir.'}", c["cfg2"], f"{n(alto)} / {n(baixo)} psi",
                       f"{n(alto / (alto + baixo) * 100)}%", f"{n(carga)} kg", f"{n(cap(alto))} kg", f"{n(uso)}%", efeito])
    tb = tabela(["Par (rodado duplo)", "Config.", "Pressões", "Peso no mais cheio", "Carga nele", "Aguenta",
                 "Uso", "Efeito"], linhas, "small wrap")
    danos = [("Sobrecarga e estouro do mais cheio", "Ele segura o peso que seria do vizinho e estoura no primeiro buraco ou toco.", "Estouro por impacto · 281"),
             ("Calor no pneu murcho", "Flexiona demais a cada volta, esquenta e a borracha descola da lona.", "Desloc. no ombro · 41"),
             ("Efeito dominó", "Se o murcho esvazia, o vizinho leva o peso do par inteiro e estoura também.", "Furou e rodou vazio · 99"),
             ("Desgaste por arrasto", "Com diâmetro menor, o murcho é arrastado no chão; perde vida e recapagem.", "Carcaça sem reforma"),
             ("Flancos encostando", "O murcho abaula para o lado e esfrega no vizinho: calor e corte.", "Saliência lateral · 30")]
    cards = "".join(f'<div class="kpi red" style="padding:3mm 3.5mm;min-width:0"><div class="t" style="color:#b3302d;height:8mm">{i}. {t}</div>'
                    f'<p style="font-size:8.4pt;margin:1.5mm 0 2mm 0;line-height:1.4;height:16mm">{x}</p>'
                    f'<div class="s"><span class="badge" style="font-size:7.8pt">{b}</span></div></div>'
                    for i, (t, x, b) in enumerate(danos, 1))
    corpo = f"""<div class="kpis">
{kpi("Pares com diferença ≥ 10 psi", f"{dup10} de {len(DD)}", "pares com as duas leituras", pct(dup10, len(DD), 0), "red")}
{kpi("Pares com diferença ≥ 20 psi", f"{dup20} pares", "um pneu leva bem mais peso", pct(dup20, len(DD), 0), "red")}
{kpi("Par com um lado sem leitura", f"{dup_meio} pares", "metade do par sem controle", pct(dup_meio, len(D), 0), "orange")}
{kpi("Diferença máxima aceita", "5 psi", "entre externo e interno", "padrão", "green")}
</div>
<h2 class="r">Por que o par com pressões diferentes perde os dois pneus</h2>
<div class="kpis" style="gap:3mm">{cards}</div>
<div class="row"><div class="col" style="flex:1.55"><h2>Exemplos da auditoria</h2>{tb}
<p class="note">Estimativa: a carga média do conjunto no par dividida na proporção da pressão de cada pneu; "Aguenta" = capacidade ETRTO na pressão do pneu mais cheio. Conta aproximada, para ordem de grandeza.</p></div>
<div class="col"><div class="box g"><h3>Como evitar</h3><ul>
<li>Diferença máxima de <b>5 psi</b> entre os dois pneus do par.</li>
<li>Calibrar os dois juntos, a frio — extensor de válvula no pneu interno.</li>
<li>Montar pares iguais: mesmo diâmetro, desgaste e vida (nova ou recapada).</li>
<li>Pneu furado ou murcho: trocar ou consertar na hora — não deixar rodar.</li>
</ul></div></div></div>"""
    pagina("Resumo — o prejuízo de rodar com o rodado duplo descalibrado", "Resumo Rodado Duplo", corpo)


# ---------------------------------------------------------------- base xlsx
def base_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pneus"
    cab = ["Frota", "Conjunto", "Config.", "Eixo", "Posição", "Leitura", "Pressão (psi)", "Faixa",
           "Kg/pneu méd. conjunto", "Kg/pneu máx. conjunto", "psi mín. p/ carga máx.",
           "Capacidade na pressão medida (kg)", "Insuficiente p/ carga méd.", "Insuficiente p/ carga máx.",
           "Célula na planilha", "Ação"]
    ws.append(cab)
    for p in P:
        if p["psi"] is None:
            acao = "Medir: instalar extensor/válvula" if p["interno"] else "Medir e calibrar"
        elif p["psi"] < 115:
            acao = f"Calibrar para {120} psi"
        elif p["psi"] > 125:
            acao = "Reduzir para 120 psi"
        else:
            acao = "OK"
        ws.append([int(p["frota"]), p["conjunto"], p["config"], p["eixo"], p["pos"], p["leitura"], p["psi"],
                   ROTULO[p["classe"]], p["kg_med"], p["kg_max"], p["psi_min_max"], p.get("capacidade"),
                   "SIM" if p.get("sobrecarga_med") else ("" if p["psi"] is None else "não"),
                   "SIM" if p.get("sobrecarga_max") else ("" if p["psi"] is None else "não"), p["celula"], acao])
    ws2 = wb.create_sheet("Rodado duplo")
    ws2.append(["Frota", "Conjunto", "Eixo", "Lado", "Externo (psi)", "Interno (psi)", "Diferença (psi)"])
    for d in D:
        ws2.append([int(d["frota"]), d["conjunto"], d["eixo"], "Esquerdo" if d["lado"] == "E" else "Direito",
                    d["ext"], d["int"], d.get("dif")])
    ws3 = wb.create_sheet("Conjuntos")
    ws3.append(["Conjunto", "Fabricante", "Ano", "Config.", "Posições auditadas", "Pressão média", "115–125 psi",
                "Sem leitura", "Kg/pn méd.", "Kg/pn máx.", "psi p/ méd.", "psi p/ máx.", "Insuf. p/ méd.",
                "Insuf. p/ máx.", "Sucateados", "Estouro", "Sucat./1.000 pn-v"])
    for c in CJ:
        ws3.append([c["conj"], c["fab"], c["ano"], c["cfg"], c["pneus"], round(c["media"], 1), c["conf"], c["sem"],
                    c["kmed"], c["kmax"], round(c["psi_med"], 1), round(c["psi_max"], 1), c["sob_med"], c["sob_max"],
                    c["suc"], c["est"], c["taxa"]])
    ws4 = wb.create_sheet("Tabela ETRTO")
    ws4.append(["bar", "psi", "kg por pneu (duplo, índice 148)", "kg por eixo (4 pneus)"])
    for b, ps, kg, ke in zip(analise.BAR, analise.TAB_PSI, analise.TAB_KG, analise.EIXO_KG):
        ws4.append([b, round(ps, 1), kg, ke])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="142A47")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(10, min(34, max(len(str(c.value or "")) for c in col) + 2))
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions
    wb.save(XLSX)


def main():
    for f in (p1, p2, p3, p4, p5, p6, p7, p8, p9):
        f()
    html = BUILD / "relatorio.html"
    html.write_text(montar(), encoding="utf-8")
    tmp = BUILD / "relatorio.pdf"
    tmp.unlink(missing_ok=True)
    subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={tmp}", html.resolve().as_uri()], check=True, timeout=120)
    tmp.replace(PDF)  # falha se o PDF estiver aberto em outro programa
    base_xlsx()
    print(PDF, XLSX)


if __name__ == "__main__":
    main()
