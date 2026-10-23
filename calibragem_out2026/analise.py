"""Auditoria de calibragem dos reboques canavieiros x peso por pneu x sucateamento.

Lê dados/CONFERENCIA DE CALIBRAGEM.xlsx (aba Planilha1), cruza com o peso por pneu e o
sucateamento de cada conjunto (Relatorio_Peso_por_Pneu_Reboques_Canavieiros.pdf) e com a
tabela ETRTO de carga x pressão do 295/80R22,5 148 (rodado duplo). Grava analise.json.
"""
import json
import re
from pathlib import Path

import openpyxl

AQUI = Path(__file__).parent
PLANILHA = AQUI / "dados" / "CONFERENCIA DE CALIBRAGEM.xlsx"

PADRAO = 120          # psi, padrão da unidade (OBS da planilha: "CALIBRAGEM PADRÃO 120 LBS")
TOL = 5               # psi, faixa aceita em torno do padrão (115 a 125)
DIF_DUPLO = 5         # psi, diferença máxima entre os dois pneus do mesmo rodado duplo

# ETRTO, 295/80R22,5 índice 148 em rodado duplo (Hankook, load inflation poster EU):
# carga por eixo de 4 pneus em 4,5 ... 8,5 bar -> 7.575 ... 12.600 kg. Aqui por pneu.
BAR = [4.5, 5.0, 5.5, 6.0, 6.5, 7.0, 7.5, 8.0, 8.5]
EIXO_KG = [7575, 8240, 8890, 9535, 10165, 10785, 11395, 12000, 12600]
TAB_PSI = [b * 14.5038 for b in BAR]
TAB_KG = [k / 4 for k in EIXO_KG]


def capacidade(psi):
    """kg por pneu (duplo) que a pressão sustenta; interpolação linear, extrapola abaixo de 65 psi."""
    p, k = TAB_PSI, TAB_KG
    if psi <= p[0]:
        return k[0] + (psi - p[0]) * (k[1] - k[0]) / (p[1] - p[0])
    for i in range(1, len(p)):
        if psi <= p[i]:
            return k[i - 1] + (psi - p[i - 1]) * (k[i] - k[i - 1]) / (p[i] - p[i - 1])
    return k[-1]


def pressao_para(kg):
    """psi mínimo para o pneu (duplo) suportar kg."""
    p, k = TAB_PSI, TAB_KG
    if kg <= k[0]:
        return p[0]
    for i in range(1, len(k)):
        if kg <= k[i]:
            return p[i - 1] + (kg - k[i - 1]) * (p[i] - p[i - 1]) / (k[i] - k[i - 1])
    return None  # acima da capacidade máxima do pneu


# Conjuntos do relatório de peso por pneu (pesagens 15/07 a 25/09/2026)
# frota auditada -> conjunto; dados: fabricante, ano, config, viagens, kg/pn méd, kg/pn máx,
# sucateados, estouro, sucateados por 1.000 pneu-viagem
CONJUNTOS = {
    "63031 + 63030": ("Santa Izabel", 2020, "9 eixos", 125, 2496, 2714, 2, 1, 0.67),
    "63127 + 63128": ("Sergomel", 2020, "Verde · 9 eixos", 147, 2681, 2918, 15, 4, 4.25),
    "63037 + 63025": ("Santa Izabel", 2020, "9 eixos", 152, 2478, 2718, 2, 2, 0.55),
    "63021 + 63020": ("Santa Izabel", 2020, "9 eixos", 29, 2528, 2717, 1, 1, 1.44),
    "63044 + 63050": ("Santa Izabel", 2020, "9 eixos", 123, 2459, 2649, 4, 0, 1.36),
    "63077 + 63078": ("Randon", 2023, "11 eixos", 210, 1955, 2085, 2, 0, 0.30),
    "63097 + 63098": ("Randon", 2024, "11 eixos", 214, 1968, 2163, 0, 0, 0.00),
    "63071 + 63074": ("Randon", 2023, "11 eixos", 185, 2019, 2231, 3, 1, 0.51),
    "63110 + 63106": ("Randon", 2024, "11 eixos", 225, 1951, 2126, 1, 0, 0.14),
    "63075 + 63076": ("Randon", 2023, "11 eixos", 197, 1954, 2112, 3, 0, 0.48),
}
FROTA_CONJ = {f: c for c in CONJUNTOS for f in c.split(" + ")}
# 63105 não está entre os 65 conjuntos pesados; 3 eixos como os semirreboques Randon de 11 eixos
FORA_BASE = {"63105": ("11 eixos", 1965, 2353)}  # média e máximo dos 11 eixos

GERAL = {  # médias das configurações no relatório de peso
    "9 eixos": {"kg_med": 2588, "kg_max": 3049, "taxa": 1.37, "conj": 41},
    "11 eixos": {"kg_med": 1965, "kg_max": 2353, "taxa": 0.24, "conj": 24},
}


def ler_planilha():
    ws = openpyxl.load_workbook(PLANILHA, data_only=True)["Planilha1"]
    pneus, obs, frota = [], {}, None
    linhas = list(ws.iter_rows(min_row=1, max_row=ws.max_row))
    for i, row in enumerate(linhas):
        a = row[0].value
        if isinstance(a, (int, float)):
            frota = str(int(a))
        g = row[6].value if len(row) > 6 else None
        if frota and isinstance(g, str) and g.strip() and g.strip() != "OBS":
            obs.setdefault(frota, []).append(g.strip())
        rotulos = [c.value for c in row[1:5]]
        if not all(isinstance(r, str) and re.search(r"(EE|EI|DI|DE)", r) for r in rotulos):
            continue
        valores = [c.value for c in linhas[i + 1][1:5]]
        for rot, val in zip(rotulos, valores):
            m = re.match(r"\s*(\d+)\D*(EE|EI|DI|DE)", rot)
            eixo, pos = int(m.group(1)), m.group(2)
            if isinstance(val, (int, float)):
                psi, leitura = float(val), "medido"
            else:
                psi, leitura = None, (str(val).strip() if val is not None else "-")
            pneus.append({"frota": frota, "eixo": eixo, "pos": pos, "psi": psi,
                          "leitura": leitura, "celula": linhas[i + 1][1 + rotulos.index(rot)].coordinate})
    return pneus, obs


def classe(psi):
    if psi is None:
        return "sem leitura"
    if psi < PADRAO - 15:
        return "< 105"
    if psi < PADRAO - TOL:
        return "105 a 114"
    if psi <= PADRAO + TOL:
        return "115 a 125"
    return "> 125"


def main():
    pneus, obs = ler_planilha()
    for p in pneus:
        f = p["frota"]
        if f in FROTA_CONJ:
            c = FROTA_CONJ[f]
            fab, ano, cfg, viag, kmed, kmax, suc, est, taxa = CONJUNTOS[c]
            p.update(conjunto=c, config="9 eixos" if "9" in cfg else "11 eixos", kg_med=kmed, kg_max=kmax)
        else:
            cfg, kmed, kmax = FORA_BASE[f]
            p.update(conjunto=f + " (fora da base de pesagem)", config=cfg, kg_med=kmed, kg_max=kmax)
        p["classe"] = classe(p["psi"])
        p["interno"] = p["pos"] in ("EI", "DI")
        p["psi_min_max"] = round(pressao_para(p["kg_max"]), 1)
        p["psi_min_med"] = round(pressao_para(p["kg_med"]), 1)
        if p["psi"] is not None:
            cap = capacidade(p["psi"])
            p["capacidade"] = round(cap)
            p["sobrecarga_med"] = cap < p["kg_med"]
            p["sobrecarga_max"] = cap < p["kg_max"]
            p["uso_med"] = round(p["kg_med"] / cap * 100, 1)

    # rodados duplos: (frota, eixo, lado)
    duplos = []
    chave = {(p["frota"], p["eixo"], p["pos"]): p for p in pneus}
    for (f, e, pos), p in chave.items():
        if pos in ("EE", "DE"):
            lado = "E" if pos == "EE" else "D"
            par = chave.get((f, e, lado + "I"))
            if par is None:
                continue
            a, b = p["psi"], par["psi"]
            d = {"frota": f, "eixo": e, "lado": lado, "ext": a, "int": b, "config": p["config"],
                 "conjunto": p["conjunto"]}
            if a is not None and b is not None:
                d["dif"] = abs(a - b)
            duplos.append(d)

    json.dump({"pneus": pneus, "duplos": duplos, "obs": obs,
               "tabela": {"psi": TAB_PSI, "kg": TAB_KG}}, open(AQUI / "analise.json", "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    return pneus, duplos, obs


if __name__ == "__main__":
    main()
