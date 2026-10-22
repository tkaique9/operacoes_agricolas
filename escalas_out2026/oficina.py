"""Oficina (102.xxx): escala de hoje por departamento e simulação das escalas pela regra de cada um.

python oficina.py  ->  oficina.json
Planilhas: dados/oficina/ativos_oficina.xlsx (abas Planilha2 com as anotações por departamento, Sheet1 com o
cadastro e Planilha1 com sexo e CLT/rural) e dados/escalas.xlsx (Relatório de Revezamento).
"""
import collections
import json
import math
import sys
from pathlib import Path

import openpyxl

import analise

AQUI = Path(__file__).parent
ARQ = AQUI / "dados" / "oficina" / "ativos_oficina.xlsx"

# Regra de cada departamento, tirada das anotações da aba Planilha2 (coluna F e linhas em azul).
#   nao      = não precisa simular        5x2  = horário de dia, segunda a sexta
#   24h      = roda 24 h (3 turnos)       1turno = 1 pessoa por turno 24 h, o resto 5x2
#   2turnos  = turno A e B (sem noite)    azul = só as linhas em azul rodam 24 h, o resto 5x2
#   sem_nota = sem anotação na planilha (tratado como 5x2 até confirmar)
REGRAS = {
    "102.001.001": "nao", "102.001.011": "azul", "102.001.012": "24h", "102.001.013": "5x2",
    "102.002.002": "24h", "102.002.003": "24h", "102.002.004": "1turno", "102.002.005": "5x2",
    "102.002.006": "24h", "102.002.007": "24h", "102.002.009": "24h", "102.002.012": "sem_nota",
    "102.003.001": "24h", "102.004.001": "24h", "102.005.001": "24h", "102.005.003": "2turnos",
    "102.006.001": "24h", "102.006.002": "5x2", "102.006.003": "24h",
}
TURNO_DE = {"05": "Manhã", "06": "Manhã", "07": "Manhã", "13": "Tarde", "14": "Tarde", "15": "Tarde",
            "21": "Noite", "22": "Noite", "23": "Noite"}
TURNOS = ["Manhã", "Tarde", "Noite"]


def dep_fmt(cod):
    s = str(cod)
    return f"{s[:3]}.{s[3:6]}.{s[6:]}"


def ler():
    wb = openpyxl.load_workbook(ARQ)  # com formatação, para achar as linhas em azul
    notas, azul, dep = {}, set(), None
    for row in wb["Planilha2"].iter_rows(min_row=4):
        a = row[0].value
        if a and "Total" not in str(a):
            dep = str(a).split(" - ")[0]
        for c in row[5:7]:
            if c.value:
                notas.setdefault(dep, []).append(str(c.value).strip())
        f = row[2].fill
        if row[1].value and f is not None and f.fill_type and f.fgColor.type == "theme":
            azul.add(row[1].value)
    val = openpyxl.load_workbook(ARQ, read_only=True, data_only=True)  # valores (a aba Sheet1 tem PROCV)
    cad = [r for r in val["Sheet1"].iter_rows(values_only=True) if isinstance(r[1], int)]
    p1 = list(val["Planilha1"].iter_rows(values_only=True))
    idx = {c: i for i, c in enumerate(p1[0])}
    extra = {r[idx["matricula"]]: (r[idx["sexo"]], r[idx["clt_rural"]]) for r in p1[1:]}
    return cad, notas, azul, extra


CEN = {c["id"]: c for c in analise.CENARIOS}
SUGS = ["5x2", "4x2", "12x36", "3x3", "4x2_12"]
# na oficina todos são CLT urbano: a noite de 12 h com 10 h de trabalho passa de 10 h contadas (hora reduzida)
RISCO_OF = {"5x2": "Baixo", "4x2": "Baixo", "12x36": "Médio", "3x3": "Médio", "4x2_12": "Não serve"}
ORDEM_RISCO = {"Baixo": 0, "Médio": 1, "Não serve": 9}


def arred(x):
    """Inteiro mais próximo, empate para baixo: 2,5 pessoas por dia hoje = garantir 2."""
    return math.ceil(x - 0.5)


def postos_padrao(d, ps):
    """Pessoas que o turno precisa garantir todo dia (manhã, tarde, noite). A diretoria pode mudar no HTML."""
    regra = d["regra"]
    if regra in ("azul", "1turno"):     # PCM em azul e caldeiraria: 1 por turno nas 24 h
        return {"Manhã": 1, "Tarde": 1, "Noite": 1}
    if regra == "2turnos":              # turno A e B, sem noite
        return {"Manhã": 1, "Tarde": 1, "Noite": 0}
    if regra != "24h":
        return None
    pres = {t: sum(p.get("pres", 0) for p in ps if p.get("turno") == t) for t in TURNOS}
    if not any(pres.values()):          # laboratório: hoje só de dia em 12x36
        return {t: 1 for t in TURNOS}
    return {t: max(1, arred(v)) for t, v in pres.items()}


def precisa(c, P, volante=False):
    """Pessoas para garantir P[turno] todo dia. Turma fixa por turno, ou com folguista volante entre turnos."""
    f = c["ciclo"] / c["trab"]
    f5 = CEN["5x2"]["ciclo"] / CEN["5x2"]["trab"]
    v = [P[t] for t in TURNOS]
    if c["turnos"] == 3:
        return math.ceil(sum(v) * f - 1e-9) if volante else sum(math.ceil(x * f - 1e-9) for x in v)
    p24 = min(v)
    if p24 == 0:                        # não roda à noite: escala de 12 h não se aplica
        return None
    ex = [x - p24 for x in v]           # o que passa da cobertura 24 h fica em 5x2 de 8 h
    if volante:
        return math.ceil(2 * p24 * f + sum(ex) * f5 - 1e-9)
    return 2 * math.ceil(p24 * f - 1e-9) + sum(math.ceil(e * f5 - 1e-9) for e in ex)


def simular(d, ps):
    regra = d["regra"]
    P = postos_padrao(d, ps)
    equipe = [p for p in ps if p["tipo"] in ("5x1", "6x1", "12x36", "Sem escala de origem")
              and (regra != "azul" or p["azul"])]
    adm = sum(1 for p in ps if p["tipo"] == "ADM") if regra == "24h" else 0
    res = {"postos": P, "equipe": len(equipe), "adm_usavel": adm, "sim": {}}
    if P is None:
        return res
    for cid in SUGS:
        fixo, vol = precisa(CEN[cid], P), precisa(CEN[cid], P, True)
        if fixo is None:
            continue
        cf, cv = fixo - len(equipe), vol - len(equipe)
        if regra == "1turno":           # o 24 h sai da própria equipe; quem sobra vai para 5x2 de dia
            cf, cv = max(0, cf), max(0, cv)
        res["sim"][cid] = {"precisa": fixo, "precisa_volante": vol, "contratar": cf, "contratar_volante": cv}
    ok = [k for k in res["sim"] if RISCO_OF[k] != "Não serve"]
    res["sugerida"] = min(ok, key=lambda k: (max(res["sim"][k]["contratar"], 0), ORDEM_RISCO[RISCO_OF[k]], SUGS.index(k)))
    return res


def main():
    cad, notas, azul, extra = ler()
    esc = [r for r in analise.ler_planilha("escalas.xlsx") if isinstance(r[0], int) and isinstance(r[3], int)]
    codigos, hist = {}, collections.defaultdict(list)
    for r in esc:
        codigos.setdefault(r[3], analise.classificar_codigo(r[3], r[4].strip()))
    for r in sorted(esc, key=lambda r: r[2], reverse=True):
        hist[r[0]].append(r)

    pessoas = []
    for r in cad:
        mat, sit = r[1], r[3]
        dep = dep_fmt(r[0])
        h = hist.get(mat, [])
        atual = h[0] if h else None
        origem = next((x for x in h if x[3] not in analise.AUSENCIA), None)
        sexo, cr = extra.get(mat, (None, None))
        p = {"matricula": mat, "nome": r[2].strip(), "situacao": sit, "funcao": r[4].strip(), "dep": dep,
             "departamento": str(r[6]).partition(" - ")[2].strip() or str(r[6]).strip(), "sexo": sexo, "clt_rural": cr, "azul": mat in azul,
             "cod_atual": atual[3] if atual else None, "escala_atual": atual[4].strip() if atual else "sem registro"}
        cod = origem[3] if origem else None
        if atual and atual[3] in analise.AUSENCIA:
            p["obs"] = analise.AUSENCIA[atual[3]] + (f"; volta para {cod}" if cod else "; sem escala de origem")
        c = codigos.get(cod) if cod else None
        p["cod_escala"] = cod
        if sit == "AF":
            p["tipo"] = "Afastado"
        elif c is None:
            p["tipo"] = "Sem registro" if atual is None else "Sem escala de origem"
        elif c["grupo"] == "5x1 turnos 24 h":
            p.update(tipo="5x1", turno=TURNO_DE[c["inicio"][:2]], inicio=c["inicio"], h_dia=c["h_dia"], pres=5 / 6)
        elif cod in (919, 552, 924):
            p.update(tipo="6x1", turno=TURNO_DE[c["inicio"][:2]], inicio=c["inicio"], h_dia=c["h_dia"], pres=6 / 7)
        elif c["regime"] == "12x36":
            p.update(tipo="12x36", turno="Dia 12 h", pres=0.5)
        else:
            p.update(tipo="ADM", turno="ADM")
        p["regra"] = REGRAS.get(dep, "sem_nota")
        pessoas.append(p)

    deps = []
    for dep in sorted({p["dep"] for p in pessoas}):
        ps = [p for p in pessoas if p["dep"] == dep]
        d = {"dep": dep, "nome": ps[0]["departamento"], "regra": REGRAS.get(dep, "sem_nota"), "notas": notas.get(dep, []),
             "n": len(ps), "situacao": dict(collections.Counter(p["situacao"] for p in ps)),
             "tipos": dict(collections.Counter(p["tipo"] for p in ps)),
             "pessoas_turno": {t: sum(1 for p in ps if p.get("turno") == t) for t in TURNOS + ["Dia 12 h", "ADM"]},
             "presenca": {t: round(sum(p.get("pres", 0) for p in ps if p.get("turno") == t), 2) for t in TURNOS + ["Dia 12 h"]},
             "funcoes": collections.Counter(p["funcao"] for p in ps).most_common(4),
             "mulheres": sum(1 for p in ps if p["sexo"] == "F"), "azul": sum(1 for p in ps if p["azul"])}
        deps.append(d)
    for d in deps:
        d.update(simular(d, [p for p in pessoas if p["dep"] == d["dep"]]))
    out = {"pessoas": pessoas, "departamentos": deps, "codigos_turnos": sorted({p["cod_escala"] for p in pessoas if p.get("turno")}),
           "clt_rural": dict(collections.Counter(p["clt_rural"] for p in pessoas))}
    (AQUI / "oficina.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return out


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    o = main()
    print("CLT/rural:", o["clt_rural"], "| pessoas:", len(o["pessoas"]))
    for d in o["departamentos"]:
        pt, pr = d["pessoas_turno"], d["presenca"]
        print(f"{d['dep']} {d['nome'][:34]:34s} {d['regra']:8s} n={d['n']:2d} {d['tipos']}")
        print(f"      pessoas M/T/N/12h/ADM = {pt['Manhã']}/{pt['Tarde']}/{pt['Noite']}/{pt['Dia 12 h']}/{pt['ADM']}"
              f"  presença/dia M/T/N/12h = {pr['Manhã']}/{pr['Tarde']}/{pr['Noite']}/{pr['Dia 12 h']}  azul={d['azul']}  mulheres={d['mulheres']}")
    print()
    tot = collections.Counter()
    for d in o["departamentos"]:
        if not d["sim"]:
            continue
        P, sm = d["postos"], d["sim"]
        linha = " ".join(f"{k}:{v['contratar']:+d}({v['contratar_volante']:+d})" for k, v in sm.items())
        print(f"{d['nome'][:30]:30s} {d['regra']:7s} P={P['Manhã']}/{P['Tarde']}/{P['Noite']} equipe={d['equipe']:2d} "
              f"adm={d['adm_usavel']} {linha}  => {d['sugerida']}")
        for k in sm:
            tot[k] += sm[k]["contratar"]
        tot["sugerida"] += sm[d["sugerida"]]["contratar"]
    print("contratar com turma fixa:", dict(tot))
