"""Gera Base_Consolidada_Fertilizantes_OUT2026.xlsx a partir de dados_consolidados.json.

Uso:  python gerar_planilha_auxiliar.py

Todas as contas ficam em fórmulas do Excel, ligadas às premissas da aba Premissas (copiadas
da planilha de origem, com a célula de onde vieram). As colunas "Planilha" da aba Balanço
trazem os valores da origem para conferência lado a lado. Depois de gerar, rode
recalcular_planilha.ps1 para o Excel calcular e gravar os resultados.
"""

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).parent
D = json.loads((AQUI / "dados_consolidados.json").read_text(encoding="utf-8"))
A, B = D["cenarios"]["A"], D["cenarios"]["B"]
SAIDA = AQUI / "Base_Consolidada_Fertilizantes_OUT2026.xlsx"

FONTE = "Arial"
AZUL = Font(name=FONTE, size=10, color="0000FF")          # dado copiado da origem
PRETO = Font(name=FONTE, size=10, color="000000")         # fórmula
VERDE = Font(name=FONTE, size=10, color="008000")         # ligação simples a outra aba
NEGRITO = Font(name=FONTE, size=10, bold=True)
CABEC = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
TITULO = Font(name=FONTE, size=14, bold=True, color="0E3B2C")
SUB = Font(name=FONTE, size=9, italic=True, color="66736D")
F_CABEC = PatternFill("solid", fgColor="0E3B2C")
F_EDITAR = PatternFill("solid", fgColor="FFFF00")
F_TOTAL = PatternFill("solid", fgColor="EEF3EA")
LINHA = Border(bottom=Side(style="thin", color="D5DED2"))

FMT = {"ha": "#,##0.00", "t": "#,##0.000", "dias": "#,##0.0", "data": "dd/mm/yyyy", "%": "0.0%",
       "ha/dia": "#,##0.00", "t/dia": "#,##0.00", "t/ha": "#,##0.00", "kg/ha": "#,##0", "fração": "0.00",
       "R$": '"R$" #,##0;-"R$" #,##0;-', "R$/ha": '"R$" #,##0.00', "R$/t": '"R$" #,##0.00', "": "General"}

wb = Workbook()


def cabecalho(ws, linha, titulos, larguras=None):
    for j, t in enumerate(titulos, start=1):
        c = ws.cell(linha, j, t)
        c.font, c.fill = CABEC, F_CABEC
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if larguras:
        for j, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def titulo(ws, texto, sub):
    ws["A1"], ws["A2"] = texto, sub
    ws["A1"].font, ws["A2"].font = TITULO, SUB


def escreve(ws, ref, valor, fonte=PRETO, fmt=None):
    c = ws[ref]
    c.value = valor
    c.font = fonte
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    return c


# ============================================================ Leia-me
ws = wb.active
ws.title = "Leia-me"
titulo(ws, "Base consolidada — necessidade de fertilizantes (posição 16/10/2026)",
       f"Origem: {D['fonte']['arquivo']} — abas {A['aba']} (cenário A) e {B['aba']} (cenário B). Nenhum valor da origem foi alterado.")
linhas = [
    ("Abas", ""),
    ("Premissas", "Dados copiados da planilha de origem, com a célula de onde vieram (texto azul)."),
    ("Balanço", "Toda a conta da necessidade, do saldo e da compra, em fórmulas; colunas F:I conferem com a origem."),
    ("Estoque", "Estoque por fórmula em 16/10/2026 e área equivalente (estoque ÷ dose implícita da origem)."),
    ("Curva_16-00-20", "Necessidade acumulada × disponível, com a premissa ilustrativa de aplicação uniforme até o término da safra."),
    ("Sensibilidade", "Dias de vinhaça com K, variação de demanda e preços: células amarelas são editáveis."),
    ("Indicadores", "Painel de acompanhamento: o que já é mensurável e o que depende de dado."),
    ("Inconsistências", "Pontos da planilha de origem que precisam de correção ou complemento."),
    ("Conferência", f"{len(D['checks'])} conferências: fórmulas e totais da origem refeitos em Python e comparados com o valor gravado pelo Excel, mais a checagem da sensibilidade."),
    ("", ""),
    ("Legenda", ""),
    ("Texto azul", "dado copiado da planilha de origem (ver coluna Origem)"),
    ("Texto preto", "fórmula"),
    ("Texto verde", "ligação simples a outra aba"),
    ("Fundo amarelo", "premissa que o usuário pode alterar (aba Sensibilidade)"),
    ("", ""),
    ("Premissas desta base", ""),
    ("Cenários", "Os cenários A e B são as duas abas da planilha de origem; só muda o número de dias de vinhaça com K adequado (C13)."),
    ("Localizada CRV", "Calculada como MÁXIMO(0; área de localizada estimada − Cargill a adubar). Reproduz a aba A (=C14−H36) e o 0 digitado em H38 da aba B."),
    ("Nutrientes", "Pela garantia expressa no nome da fórmula (N-P2O5-K2O, em %). O teor de K da vinhaça não está na origem e não foi estimado."),
    ("Curva", "Distribuição uniforme da necessidade entre a data-base e o término da safra — ilustrativa; a origem não traz demanda mensal nem datas de entrega."),
    ("Preços", "A origem não traz preços; nenhum preço foi assumido. Informe-os na aba Sensibilidade."),
]
for i, (a, b) in enumerate(linhas, start=4):
    ws.cell(i, 1, a).font = NEGRITO
    ws.cell(i, 2, b).font = PRETO
ws["A15"].font = Font(name=FONTE, size=10, color="0000FF")
ws["A16"].font = PRETO
ws["A17"].font = VERDE
ws["A18"].fill = F_EDITAR
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 120

# ============================================================ Premissas
ws = wb.create_sheet("Premissas")
titulo(ws, "Premissas copiadas da planilha de origem", "Texto azul = valor da origem. Altere aqui só se a origem mudar.")
cabecalho(ws, 4, ["Item", "Unidade", "Cenário A (LOC. 1.000ha)", "Cenário B (LOC. 450ha)", "Origem (célula nas duas abas)"], [44, 10, 22, 22, 46])
cr = A["consumo_realizado"]
premissas = [
    ("data_base", "Data-base", "data", date.fromisoformat(A["data_base"]), date.fromisoformat(B["data_base"]), "A10"),
    ("data_limite", "Data limite da vinhaça localizada", "data", date.fromisoformat(A["data_limite_localizada"]), date.fromisoformat(B["data_limite_localizada"]), "B10"),
    ("termino", "Término da safra", "data", date.fromisoformat(A["termino_safra"]), date.fromisoformat(B["termino_safra"]), "C10"),
    ("ritmo", "Ritmo da vinhaça localizada (média de setembro)", "ha/dia", A["ritmo_localizada_ha_dia"], B["ritmo_localizada_ha_dia"], "C12"),
    ("dias_k", "Dias com vinhaça de K adequado", "dias", A["dias_k_vinhaca"], B["dias_k_vinhaca"], "C13"),
    ("tratos_cargill", "Área de tratos da Cargill", "ha", A["cargill"]["tratos_ha"], B["cargill"]["tratos_ha"], "E36 (H36 usa 1.200 digitado)"),
    ("real_cargill", "Realizado Cargill (vinhaça localizada)", "ha", A["cargill"]["realizado_ha"], B["cargill"]["realizado_ha"], "F36 = B20"),
    ("colher1", "Área à colher CRV — parcela 1", "ha", A["crv"]["a_colher_partes"][0], B["crv"]["a_colher_partes"][0], "D38 (=11593+9354+8779)"),
    ("colher2", "Área à colher CRV — parcela 2", "ha", A["crv"]["a_colher_partes"][1], B["crv"]["a_colher_partes"][1], "D38"),
    ("colher3", "Área à colher CRV — parcela 3", "ha", A["crv"]["a_colher_partes"][2], B["crv"]["a_colher_partes"][2], "D38"),
    ("plantio", "Plantio (renovação) excluído dos tratos", "ha", A["crv"]["plantio_ha"], B["crv"]["plantio_ha"], "C38"),
    ("real_loc", "Realizado CRV — vinhaça localizada", "ha", A["realizado"]["crv_localizada_ha"], B["realizado"]["crv_localizada_ha"], "B22 = F38"),
    ("real_min", "Realizado CRV — fertilizante mineral", "ha", A["realizado"]["crv_mineral_ha"], B["realizado"]["crv_mineral_ha"], "B23 = F39"),
    ("bisada", "Cana bisada — área", "ha", A["crv"]["bisada_ha"], B["crv"]["bisada_ha"], "A45 = G39"),
    ("tch", "Cana bisada — TCH estimado", "t/ha", A["crv"]["tch_bisada"], B["crv"]["tch_bisada"], "B45"),
    ("dose46", "Dose média 46-00-00", "kg/ha", A["doses_kg_ha"]["46-00-00"], B["doses_kg_ha"]["46-00-00"], "J38"),
    ("dose16", "Dose média 16-00-20", "kg/ha", A["doses_kg_ha"]["16-00-20"], B["doses_kg_ha"]["16-00-20"], "J39"),
    ("est46", "Estoque 46-00-00 (CRV-MG)", "t", A["balanco"]["46-00-00"]["estoque_t"], B["balanco"]["46-00-00"]["estoque_t"], "C50 = J50"),
    ("est16", "Estoque 16-00-20 (CRV-MG)", "t", A["balanco"]["16-00-20"]["estoque_t"], B["balanco"]["16-00-20"]["estoque_t"], "C51 = J49"),
    ("rec46", "Comprado a receber 46-00-00", "t", A["balanco"]["46-00-00"]["a_receber_t"], B["balanco"]["46-00-00"]["a_receber_t"], "D50"),
    ("rec16", "Comprado a receber 16-00-20", "t", A["balanco"]["16-00-20"]["a_receber_t"], B["balanco"]["16-00-20"]["a_receber_t"], "D51"),
    ("n16", "Teor de N do 16-00-20", "fração", 0.16, 0.16, "garantia no nome da fórmula"),
    ("k16", "Teor de K2O do 16-00-20", "fração", 0.20, 0.20, "garantia no nome da fórmula"),
    ("n46", "Teor de N do 46-00-00", "fração", 0.46, 0.46, "garantia no nome da fórmula"),
    ("fundacao", "Plantio projetado — fundação", "ha", D["plantio_projetado_ha"]["fundacao"], D["plantio_projetado_ha"]["fundacao"], "nota A42"),
    ("renovacao", "Plantio projetado — renovação", "ha", D["plantio_projetado_ha"]["renovacao"], D["plantio_projetado_ha"]["renovacao"], "nota A42"),
]
P = {}
for i, (k, item, un, va, vb, orig) in enumerate(premissas, start=5):
    P[k] = i
    ws.cell(i, 1, item).font = PRETO
    ws.cell(i, 2, un).font = PRETO
    for col, v in ((3, va), (4, vb)):
        c = ws.cell(i, col, v)
        c.font = AZUL
        c.number_format = FMT[un]
    ws.cell(i, 5, orig).font = PRETO
    for j in range(1, 6):
        ws.cell(i, j).border = LINHA
ws.cell(P["dias_k"], 5).comment = Comment("Única premissa que difere entre os cenários A e B.", "Planejamento")


def pr(k, col):
    return f"Premissas!${col}${P[k]}"


# ============================================================ Balanço
ws = wb.create_sheet("Balanço")
titulo(ws, "Balanço de necessidade, disponibilidade e compra", "Colunas C:D em fórmulas; F:G trazem o valor gravado na planilha de origem; H:I devem ser zero.")
cabecalho(ws, 4, ["Item", "Unidade", "Cenário A", "Cenário B", "Conta", "Planilha A", "Planilha B", "Dif. A", "Dif. B", "Célula na origem"],
          [46, 9, 16, 16, 52, 14, 14, 10, 10, 18])
bA, bB = A["balanco"], B["balanco"]
L = {}
linhas_bal = [
    # chave, item, unidade, fórmula com {c} = coluna da premissa e {r} = coluna do balanço, valores na origem, célula
    ("dias_lim", "Dias até o limite da localizada", "dias", "={P_data_limite}-{P_data_base}", A["dias_limite_localizada"], B["dias_limite_localizada"], "B11"),
    ("dias_saf", "Dias até o término da safra", "dias", "={P_termino}-{P_data_base}", A["dias_safra"], B["dias_safra"], "C11"),
    ("loc_est", "Área de localizada estimada", "ha", "={P_dias_k}*{P_ritmo}", A["area_localizada_estimada_ha"], B["area_localizada_estimada_ha"], "C14"),
    ("cargill", "Cargill a adubar", "ha", "={P_tratos_cargill}-{P_real_cargill}", A["cargill"]["a_tratar_ha"], B["cargill"]["a_tratar_ha"], "H36"),
    ("colher", "Área à colher CRV", "ha", "={P_colher1}+{P_colher2}+{P_colher3}", A["crv"]["a_colher_ha"], B["crv"]["a_colher_ha"], "D38"),
    ("tratos", "Área de tratos CRV", "ha", "={B_colher}-{P_plantio}", A["crv"]["tratos_ha"], B["crv"]["tratos_ha"], "E38"),
    ("real_crv", "Realizado CRV", "ha", "={P_real_loc}+{P_real_min}", A["crv"]["realizado_ha"], B["crv"]["realizado_ha"], "F40"),
    ("loc_crv", "Localizada CRV a adubar", "ha", "=MAX(0,{B_loc_est}-{B_cargill})", A["crv"]["a_tratar_localizada_ha"], B["crv"]["a_tratar_localizada_ha"], "H38"),
    ("mineral", "Fertilizante mineral CRV a adubar", "ha", "={B_tratos}-{B_real_crv}-{B_loc_crv}-{P_bisada}", A["crv"]["a_tratar_mineral_ha"], B["crv"]["a_tratar_mineral_ha"], "H39"),
    ("total", "Total a adubar", "ha", "={B_loc_crv}+{B_mineral}+{B_cargill}", A["total_a_tratar_ha"], B["total_a_tratar_ha"], "H40"),
    ("nec46", "Necessidade de 46-00-00", "t", "={P_dose46}*{B_loc_crv}/1000", bA["46-00-00"]["necessidade_t"], bB["46-00-00"]["necessidade_t"], "K38 = B50"),
    ("nec16", "Necessidade de 16-00-20", "t", "={P_dose16}*{B_mineral}/1000", bA["16-00-20"]["necessidade_t"], bB["16-00-20"]["necessidade_t"], "K39 = B51"),
    ("nec", "Necessidade total", "t", "={B_nec46}+{B_nec16}", None, None, ""),
    ("disp46", "Disponível 46-00-00 (estoque + a receber)", "t", "={P_est46}+{P_rec46}", bA["46-00-00"]["disponivel_t"], bB["46-00-00"]["disponivel_t"], "E50"),
    ("disp16", "Disponível 16-00-20 (estoque + a receber)", "t", "={P_est16}+{P_rec16}", bA["16-00-20"]["disponivel_t"], bB["16-00-20"]["disponivel_t"], "E51"),
    ("saldo46", "Saldo 46-00-00 (disponível − necessidade)", "t", "={B_disp46}-{B_nec46}", bA["46-00-00"]["saldo_t"], bB["46-00-00"]["saldo_t"], "F50"),
    ("saldo16", "Saldo 16-00-20 (disponível − necessidade)", "t", "={B_disp16}-{B_nec16}", bA["16-00-20"]["saldo_t"], bB["16-00-20"]["saldo_t"], "F51"),
    ("compra16", "Compra necessária de 16-00-20", "t", "=MAX(0,-{B_saldo16})", A["compra_16_t"], B["compra_16_t"], "B56"),
    ("cob_est", "Cobertura do 16-00-20 só com estoque", "%", "={P_est16}/{B_nec16}", None, None, ""),
    ("cob_disp", "Cobertura do 16-00-20 com estoque + a receber", "%", "={B_disp16}/{B_nec16}", None, None, ""),
    ("part16", "Participação do 16-00-20 na necessidade", "%", "={B_nec16}/{B_nec}", None, None, ""),
    ("n_nec", "N contido na necessidade", "t", "={B_nec16}*{P_n16}+{B_nec46}*{P_n46}", None, None, ""),
    ("k_nec", "K2O contido na necessidade (só 16-00-20)", "t", "={B_nec16}*{P_k16}", None, None, ""),
    ("n_exc", "N contido no excedente de 46-00-00", "t", "=MAX(0,{B_saldo46})*{P_n46}", None, None, ""),
    ("n_def", "N do déficit de 16-00-20", "t", "={B_compra16}*{P_n16}", None, None, ""),
    ("k_def", "K2O do déficit de 16-00-20", "t", "={B_compra16}*{P_k16}", None, None, ""),
    ("tratos_tot", "Área de tratos total (CRV + Cargill)", "ha", "={B_tratos}+{P_tratos_cargill}", None, None, "E40"),
    ("real_tot", "Área já adubada (CRV + Cargill)", "ha", "={B_real_crv}+{P_real_cargill}", None, None, ""),
    ("pct_real", "Área já adubada / área de tratos", "%", "={B_real_tot}/{B_tratos_tot}", None, None, ""),
    ("bisada_t", "Cana bisada (t)", "t", "={P_bisada}*{P_tch}", A["crv"]["bisada_t"], B["crv"]["bisada_t"], "C45"),
    ("ritmo_ref", "Ritmo médio para adubar o restante até o término (referência)", "ha/dia", "={B_total}/{B_dias_saf}", None, None, ""),
    ("diario16", "Consumo diário de 16-00-20 (premissa linear)", "t/dia", "={B_nec16}/{B_dias_saf}", None, None, ""),
    ("dias_est", "Dias de cobertura do estoque de 16-00-20", "dias", "={P_est16}/{B_diario16}", None, None, ""),
    ("dias_disp", "Dias de cobertura do disponível de 16-00-20", "dias", "={B_disp16}/{B_diario16}", None, None, ""),
    ("data_est", "Esgotamento do estoque de 16-00-20 (premissa linear)", "data", "={P_data_base}+{B_dias_est}", None, None, ""),
    ("data_disp", "Esgotamento do disponível de 16-00-20 (premissa linear)", "data", "={P_data_base}+{B_dias_disp}", None, None, ""),
]
for i, (k, *_r) in enumerate(linhas_bal, start=5):
    L[k] = i


ROTULO_P = {k: item for k, item, *_ in premissas}
ROTULO_B = {k: item for k, item, *_ in linhas_bal}


def legivel(formula):
    out = formula.lstrip("=")
    for k, t in ROTULO_P.items():
        out = out.replace("{P_" + k + "}", t)
    for k, t in ROTULO_B.items():
        out = out.replace("{B_" + k + "}", t)
    return out.replace("*", " × ").replace("/", " ÷ ").replace("-", " − ").replace("+", " + ").replace("MAX(0,", "máximo(0; ")


def resolve(formula, col):
    out = formula
    for k, r in P.items():
        out = out.replace("{P_" + k + "}", f"Premissas!{'C' if col == 'C' else 'D'}{r}")
    for k, r in L.items():
        out = out.replace("{B_" + k + "}", f"{col}{r}")
    return out


for k, item, un, f, va, vb, cel in linhas_bal:
    i = L[k]
    ws.cell(i, 1, item).font = PRETO
    ws.cell(i, 2, un).font = PRETO
    for col in ("C", "D"):
        escreve(ws, f"{col}{i}", resolve(f, col), PRETO, un)
    ws.cell(i, 5, legivel(f)).font = SUB
    if va is not None:
        escreve(ws, f"F{i}", va, AZUL, un)
        escreve(ws, f"G{i}", vb, AZUL, un)
        escreve(ws, f"H{i}", f"=ROUND(C{i}-F{i},6)", PRETO, "0.000000")
        escreve(ws, f"I{i}", f"=ROUND(D{i}-G{i},6)", PRETO, "0.000000")
    ws.cell(i, 10, cel).font = PRETO
    if k in ("total", "nec", "compra16"):
        for j in range(1, 11):
            ws.cell(i, j).fill = F_TOTAL
        for j in (1, 3, 4):
            ws.cell(i, j).font = NEGRITO
    for j in range(1, 11):
        ws.cell(i, j).border = LINHA
fim_bal = 5 + len(linhas_bal)
ws.cell(fim_bal + 1, 1, "Conferência geral (soma das diferenças absolutas)").font = NEGRITO
escreve(ws, f"H{fim_bal + 1}", f"=SUMPRODUCT(ABS(H5:I{fim_bal - 1}))", NEGRITO, "0.000000")
ws.freeze_panes = "C5"

# ============================================================ Estoque
ws = wb.create_sheet("Estoque")
titulo(ws, "Estoque de fertilizantes em 16/10/2026", "Área equivalente = estoque ÷ dose implícita usada na própria planilha (quadro I47:K61).")
cabecalho(ws, 4, ["Fórmula", "Grupo", "Estoque (t)", "Dose implícita (t/ha)", "Área equivalente (ha)", "Origem"], [16, 18, 14, 20, 20, 18])
est = A["estoque"]
for i, e in enumerate(est, start=5):
    ws.cell(i, 1, e["produto"]).font = PRETO
    ws.cell(i, 2, e["grupo"]).font = PRETO
    escreve(ws, f"C{i}", e["estoque_t"], AZUL, "t")
    escreve(ws, f"D{i}", e["dose_t_ha"], AZUL, "t/ha")
    escreve(ws, f"E{i}", f"=C{i}/D{i}", PRETO, "ha")
    ws.cell(i, 6, f"I{e['linha']}:K{e['linha']}").font = PRETO
ult = 4 + len(est)
for j, g in enumerate(("Socaria", "Plantio", "Não classificado"), start=ult + 2):
    ws.cell(j, 1, f"Subtotal {g}").font = NEGRITO
    escreve(ws, f"C{j}", f'=SUMIF($B$5:$B${ult},"{g}",C$5:C${ult})', NEGRITO, "t")
    escreve(ws, f"E{j}", f'=SUMIF($B$5:$B${ult},"{g}",E$5:E${ult})', NEGRITO, "ha")
tot = ult + 5
ws.cell(tot, 1, "Total").font = NEGRITO
escreve(ws, f"C{tot}", f"=SUM(C5:C{ult})", NEGRITO, "t")
ws.cell(tot + 1, 1, "A origem soma K49:K53 em SOCARIA e K54:K57 em PLANTIO; o Oxyfertil (linha 58) fica fora dos dois.").font = SUB
ws.cell(tot + 2, 1, "As áreas não se somam entre produtos: o 46-00-00 só substitui o 16-00-20 onde há vinhaça localizada (K).").font = SUB

# ============================================================ Curva
ws = wb.create_sheet("Curva_16-00-20")
titulo(ws, "Curva de abastecimento do 16-00-20", "Premissa ilustrativa: necessidade distribuída por igual entre a data-base e o término da safra.")
cabecalho(ws, 4, ["Dia", "Data", "Necessidade acumulada A (t)", "Estoque A (t)", "Estoque + a receber A (t)",
                  "Necessidade acumulada B (t)", "Estoque B (t)", "Estoque + a receber B (t)"], [8, 12, 24, 14, 22, 24, 14, 22])
for i, d in enumerate(range(0, A["dias_safra"] + 1, 5), start=5):
    escreve(ws, f"A{i}", d, AZUL, "0")
    escreve(ws, f"B{i}", f"=Premissas!$C${P['data_base']}+A{i}", PRETO, "data")
    escreve(ws, f"C{i}", f"='Balanço'!$C${L['diario16']}*A{i}", PRETO, "t")
    escreve(ws, f"D{i}", f"=Premissas!$C${P['est16']}", VERDE, "t")
    escreve(ws, f"E{i}", f"='Balanço'!$C${L['disp16']}", VERDE, "t")
    escreve(ws, f"F{i}", f"='Balanço'!$D${L['diario16']}*A{i}", PRETO, "t")
    escreve(ws, f"G{i}", f"=Premissas!$D${P['est16']}", VERDE, "t")
    escreve(ws, f"H{i}", f"='Balanço'!$D${L['disp16']}", VERDE, "t")

# ============================================================ Sensibilidade
ws = wb.create_sheet("Sensibilidade")
titulo(ws, "Sensibilidade e cenários definidos pelo usuário", "Células amarelas são editáveis. Nada aqui é previsão oficial da usina.")
for col, w in zip("ABCDEFGH", [40, 16, 18, 18, 18, 18, 20, 4]):
    ws.column_dimensions[col].width = w

ws["A4"] = "1. Dias de vinhaça localizada com K adequado (mesma lógica da aba A da origem)"
ws["A4"].font = NEGRITO
cabecalho_s = ["Dias (editável)", "Localizada CRV (ha)", "Mineral CRV (ha)", "Necessidade 16-00-20 (t)", "Necessidade 46-00-00 (t)", "Compra 16-00-20 (t)", "Excedente 46-00-00 (t)"]
for j, t in enumerate(cabecalho_s, start=1):
    c = ws.cell(5, j, t)
    c.font, c.fill = CABEC, F_CABEC
    c.alignment = Alignment(wrap_text=True, vertical="center")
for i, s in enumerate(D["sensibilidade_dias"], start=6):
    c = escreve(ws, f"A{i}", s["dias"], AZUL, "0")
    c.fill = F_EDITAR
    escreve(ws, f"B{i}", f"=MAX(0,A{i}*Premissas!$C${P['ritmo']}-'Balanço'!$C${L['cargill']})", PRETO, "ha")
    escreve(ws, f"C{i}", f"='Balanço'!$C${L['tratos']}-'Balanço'!$C${L['real_crv']}-Premissas!$C${P['bisada']}-B{i}", PRETO, "ha")
    escreve(ws, f"D{i}", f"=C{i}*Premissas!$C${P['dose16']}/1000", PRETO, "t")
    escreve(ws, f"E{i}", f"=B{i}*Premissas!$C${P['dose46']}/1000", PRETO, "t")
    escreve(ws, f"F{i}", f"=MAX(0,D{i}-'Balanço'!$C${L['disp16']})", PRETO, "t")
    escreve(ws, f"G{i}", f"='Balanço'!$C${L['disp46']}-E{i}", PRETO, "t")
ult_s = 5 + len(D["sensibilidade_dias"])
ws.cell(ult_s + 1, 1, "20 dias reproduz o cenário A e 0 dia o cenário B. Acima de 20 dias, só com vinhaça de K adequado e até a data limite (45 dias).").font = SUB

r = ult_s + 3
ws.cell(r, 1, "2. Variação de demanda definida pelo usuário").font = NEGRITO
ws.cell(r + 1, 1, "Variação sobre a necessidade de 16-00-20 (ex.: 10% ou −5%)").font = PRETO
c = escreve(ws, f"B{r + 1}", 0, AZUL, "0.0%")
c.fill = F_EDITAR
c.comment = Comment("Premissa do usuário. 0% = necessidade da planilha.", "Planejamento")
for j, t in enumerate(["Item", "Cenário A", "Cenário B"], start=1):
    cc = ws.cell(r + 2, j, t)
    cc.font, cc.fill = CABEC, F_CABEC
VAR = f"$B${r + 1}"
escreve(ws, f"A{r + 3}", "Necessidade de 16-00-20 ajustada (t)")
escreve(ws, f"B{r + 3}", f"='Balanço'!C{L['nec16']}*(1+{VAR})", PRETO, "t")
escreve(ws, f"C{r + 3}", f"='Balanço'!D{L['nec16']}*(1+{VAR})", PRETO, "t")
escreve(ws, f"A{r + 4}", "Compra de 16-00-20 ajustada (t)")
escreve(ws, f"B{r + 4}", f"=MAX(0,B{r + 3}-'Balanço'!C{L['disp16']})", PRETO, "t")
escreve(ws, f"C{r + 4}", f"=MAX(0,C{r + 3}-'Balanço'!D{L['disp16']})", PRETO, "t")

r2 = r + 7
ws.cell(r2, 1, "3. Preços definidos pelo usuário (a origem não traz preços)").font = NEGRITO
ws.cell(r2 + 1, 1, "Preço do 16-00-20 (R$/t)").font = PRETO
ws.cell(r2 + 2, 1, "Preço do 46-00-00 (R$/t)").font = PRETO
for rr_ in (r2 + 1, r2 + 2):
    c = escreve(ws, f"B{rr_}", None, AZUL, "R$/t")
    c.fill = F_EDITAR
ws.cell(r2 + 1, 3, "← informe o preço; enquanto vazio, o orçamento mostra \"informar preço\"").font = SUB
for j, t in enumerate(["Item", "Cenário A", "Cenário B"], start=1):
    cc = ws.cell(r2 + 3, j, t)
    cc.font, cc.fill = CABEC, F_CABEC
P16, P46 = f"$B${r2 + 1}", f"$B${r2 + 2}"
orc = [
    ("Orçamento da compra de 16-00-20 (R$)", lambda col: f"=IF({P16}=\"\",\"informar preço\",'Balanço'!{col}{L['compra16']}*{P16})", "R$"),
    ("Orçamento da compra ajustada pela variação de demanda (R$)", lambda col: f"=IF({P16}=\"\",\"informar preço\",{'B' if col == 'C' else 'C'}{r + 4}*{P16})", "R$"),
    ("Custo do 16-00-20 por hectare (R$/ha)", lambda col: f"=IF({P16}=\"\",\"informar preço\",{P16}*Premissas!{col}{P['dose16']}/1000)", "R$/ha"),
    ("Custo do 46-00-00 por hectare (R$/ha)", lambda col: f"=IF({P46}=\"\",\"informar preço\",{P46}*Premissas!{col}{P['dose46']}/1000)", "R$/ha"),
    ("Valor do excedente de 46-00-00 (R$)", lambda col: f"=IF({P46}=\"\",\"informar preço\",MAX(0,'Balanço'!{col}{L['saldo46']})*{P46})", "R$"),
    ("Efeito de cada R$ 100/t no preço do 16-00-20 sobre a compra (R$)", lambda col: f"='Balanço'!{col}{L['compra16']}*100", "R$"),
    ("Efeito de cada R$ 100/t no preço do 46-00-00 sobre o excedente (R$)", lambda col: f"=MAX(0,'Balanço'!{col}{L['saldo46']})*100", "R$"),
]
for i, (item, fn, un) in enumerate(orc, start=r2 + 4):
    ws.cell(i, 1, item).font = PRETO
    escreve(ws, f"B{i}", fn("C"), PRETO, un)
    escreve(ws, f"C{i}", fn("D"), PRETO, un)

# ============================================================ Indicadores
ws = wb.create_sheet("Indicadores")
titulo(ws, "Painel de indicadores de acompanhamento", "\"pendente\" = depende de dado que a planilha de origem não traz.")
cabecalho(ws, 4, ["Indicador", "Unidade", "Cenário A", "Cenário B", "Situação", "Conta ou dado que falta"], [36, 9, 16, 16, 18, 70])
ind = [
    ("Necessidade total prevista", "t", f"='Balanço'!C{L['nec']}", f"='Balanço'!D{L['nec']}", "mensurável", "Balanço: necessidade de 46-00-00 + 16-00-20"),
    ("Volume comprado a receber", "t", f"=Premissas!C{P['rec16']}+Premissas!C{P['rec46']}", f"=Premissas!D{P['rec16']}+Premissas!D{P['rec46']}", "mensurável", "D50 + D51 da origem (sem datas de entrega)"),
    ("Volume entregue", "t", "pendente", "pendente", "depende de dado", "registro de recebimentos por data"),
    ("Área já adubada", "ha", f"='Balanço'!C{L['real_tot']}", f"='Balanço'!D{L['real_tot']}", "mensurável", "toneladas aplicadas só existem para a vinhaça localizada (linhas 27–28)"),
    ("Saldo a comprar de 16-00-20", "t", f"='Balanço'!C{L['compra16']}", f"='Balanço'!D{L['compra16']}", "mensurável", "B56 da origem"),
    ("Aderência ao cronograma", "%", "pendente", "pendente", "depende de dado", "cronograma de entrega e de aplicação"),
    ("Custo previsto × realizado", "R$", "pendente", "pendente", "depende de dado", "preços contratados e notas fiscais"),
    ("Estoque projetado de 46-00-00 no fim da safra", "t", f"='Balanço'!C{L['saldo46']}", f"='Balanço'!D{L['saldo46']}", "mensurável", "o 16-00-20 zera depois da compra"),
]
for i, (nome, un, fa, fb, sit, obs) in enumerate(ind, start=5):
    ws.cell(i, 1, nome).font = PRETO
    ws.cell(i, 2, un).font = PRETO
    escreve(ws, f"C{i}", fa, PRETO if fa.startswith("=") else SUB, un)
    escreve(ws, f"D{i}", fb, PRETO if fb.startswith("=") else SUB, un)
    ws.cell(i, 5, sit).font = Font(name=FONTE, size=10, bold=True, color="1F6B3A" if sit == "mensurável" else "B07F0E")
    ws.cell(i, 6, obs).font = PRETO

# ============================================================ Inconsistências
ws = wb.create_sheet("Inconsistências")
titulo(ws, "Inconsistências e dados faltantes na planilha de origem", "Nenhuma célula da origem foi alterada; os pontos abaixo pedem validação.")
cabecalho(ws, 4, ["#", "Gravidade", "Onde", "Achado", "Efeito"], [4, 10, 30, 90, 60])
for i, inc in enumerate(D["inconsistencias"], start=5):
    for j, k in enumerate(("id", "gravidade", "onde", "achado", "efeito"), start=1):
        c = ws.cell(i, j, inc[k])
        c.font = PRETO
        c.alignment = Alignment(wrap_text=True, vertical="top")

# ============================================================ Conferência
ws = wb.create_sheet("Conferência")
titulo(ws, "Conferência das fórmulas da planilha de origem", "Cada fórmula foi refeita em Python (extrair_dados.py) e comparada com o valor gravado pelo Excel.")
cabecalho(ws, 4, ["Aba", "Célula", "Item", "Valor na planilha", "Recalculado", "Diferença", "OK"], [14, 10, 44, 18, 18, 12, 6])
for i, ck in enumerate(D["checks"], start=5):
    ws.cell(i, 1, ck["aba"]).font = PRETO
    ws.cell(i, 2, ck["celula"]).font = PRETO
    ws.cell(i, 3, ck["item"]).font = PRETO
    escreve(ws, f"D{i}", ck["planilha"], AZUL, "#,##0.000")
    escreve(ws, f"E{i}", ck["recalculado"], AZUL, "#,##0.000")
    escreve(ws, f"F{i}", f"=ROUND(E{i}-D{i},6)", PRETO, "0.000000")
    escreve(ws, f"G{i}", f'=IF(ABS(F{i})<0.000001,"sim","NÃO")', PRETO)
fim_c = 4 + len(D["checks"])
ws.cell(fim_c + 2, 3, "Conferências sem divergência").font = NEGRITO
escreve(ws, f"D{fim_c + 2}", f'=COUNTIF(G5:G{fim_c},"sim")', NEGRITO, "0")
escreve(ws, f"E{fim_c + 2}", f"=COUNTA(G5:G{fim_c})", NEGRITO, "0")

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save(SAIDA)
print(f"{SAIDA.name}: {len(wb.worksheets)} abas")
