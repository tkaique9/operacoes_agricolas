"""Lê a planilha "Necessidade Fertilizante OUT 2026.xlsx", refaz cada conta e grava
dados_consolidados.json, que alimenta a apresentação e a planilha auxiliar.

Uso:  python extrair_dados.py "<caminho da planilha>"

Nada é escrito na planilha de origem. Cada fórmula da origem é recalculada aqui e
comparada com o valor gravado pelo Excel; divergência interrompe o script.
"""

import json
import sys
from datetime import timedelta
from pathlib import Path

import openpyxl

ABAS = {"A": "LOC. 1.000ha", "B": "LOC. 450ha"}
NOME_CENARIO = {
    "A": "Cenário A · Localizada 1.000 ha",
    "B": "Cenário B · Localizada 450 ha",
}
TOL = 1e-6

# Teor de nutrientes lido do próprio nome da fórmula (N-P2O5-K2O, em %).
TEOR = {"16-00-20": (0.16, 0.00, 0.20), "46-00-00": (0.46, 0.00, 0.00)}

checks = []


def confere(rotulo, calculado, planilha, aba, celula):
    ok = abs(float(calculado) - float(planilha)) <= TOL
    checks.append({
        "aba": aba, "celula": celula, "item": rotulo,
        "planilha": round(float(planilha), 6), "recalculado": round(float(calculado), 6),
        "ok": ok,
    })
    if not ok:
        raise SystemExit(f"Divergência em {aba}!{celula} ({rotulo}): planilha {planilha} × recalculado {calculado}")


def produto(txt):
    """'16-00-20.' e '16-00-20' são a mesma fórmula na planilha."""
    return str(txt).strip().rstrip(".").strip()


def le_aba(ws, wf, chave):
    aba = ws.title
    v = lambda c: ws[c].value
    f = lambda c: wf[c].value

    data_base, data_limite, termino = v("A10"), v("B10"), v("C10")
    ritmo, dias_k = v("C12"), v("C13")
    confere("Dias até limite da localizada", (data_limite - data_base).days, v("B11"), aba, "B11")
    confere("Dias até término da safra", (termino - data_base).days, v("C11"), aba, "C11")
    area_loc_est = dias_k * ritmo
    confere("Área estimada de localizada", area_loc_est, v("C14"), aba, "C14")

    # 2. Áreas adubadas (realizado)
    real = {
        "cargill_localizada_ha": v("B20"),
        "crv_localizada_ha": v("B22"),
        "crv_mineral_ha": v("B23"),
    }
    confere("Total CRV adubado", real["crv_localizada_ha"] + real["crv_mineral_ha"], v("B24"), aba, "B24")
    consumo_real = []
    for lin in (27, 28):
        area, prod, dose = v(f"B{lin}"), produto(v(f"C{lin}")), v(f"D{lin}")
        confere(f"Consumo realizado {prod}", dose * area / 1000, v(f"E{lin}"), aba, f"E{lin}")
        consumo_real.append({"produto": prod, "area_ha": area, "dose_kg_ha": dose,
                             "volume_t": dose * area / 1000, "fazendas": v(f"F{lin}"), "linha": lin})
    quadro_localizada = {"total_ha": v("B26"), "partes_ha": [v("B27"), v("B28"), v("B29")]}

    # 3. Áreas remanescentes
    cargill = {"colhida_ha": v("B36"), "a_colher_ha": v("D36"), "tratos_ha": v("E36"),
               "realizado_ha": v("F36")}
    cargill["a_tratar_ha"] = 1200 - cargill["realizado_ha"]  # a planilha digita 1200 (= E36)
    confere("Cargill a tratar", cargill["a_tratar_ha"], v("H36"), aba, "H36")

    plantio = v("C38")
    a_colher = 11593 + 9354 + 8779
    confere("CRV áreas à colher", a_colher, v("D38"), aba, "D38")
    tratos = a_colher - plantio
    confere("CRV área de tratos", tratos, v("E38"), aba, "E38")
    confere("Realizado localizada = B22", real["crv_localizada_ha"], v("F38"), aba, "F38")
    confere("Realizado mineral = B23", real["crv_mineral_ha"], v("F39"), aba, "F39")
    realizado_crv = real["crv_localizada_ha"] + real["crv_mineral_ha"]
    confere("Total realizado CRV", realizado_crv, v("F40"), aba, "F40")

    bisada_ha, tch = v("A45"), v("B45")
    confere("Cana bisada (t)", bisada_ha * tch, v("C45"), aba, "C45")
    confere("Área bisada = A45", bisada_ha, v("G39"), aba, "G39")

    h38_formula = f("H38")
    if isinstance(h38_formula, str) and h38_formula.startswith("="):
        loc_crv = area_loc_est - cargill["a_tratar_ha"]
        confere("CRV a tratar com localizada", loc_crv, v("H38"), aba, "H38")
        h38_digitado = False
    else:
        loc_crv = v("H38")
        h38_digitado = True
    mineral = tratos - realizado_crv - loc_crv - bisada_ha
    confere("CRV a tratar com mineral", mineral, v("H39"), aba, "H39")
    total_tratar = mineral + loc_crv + cargill["a_tratar_ha"]
    confere("Total a tratar", total_tratar, v("H40"), aba, "H40")

    dose46, dose16 = v("J38"), v("J39")
    nec46 = dose46 * loc_crv / 1000
    nec16 = dose16 * mineral / 1000
    confere("Necessidade 46-00-00", nec46, v("K38"), aba, "K38")
    confere("Necessidade 16-00-20", nec16, v("K39"), aba, "K39")
    confere("Área colhida total", v("B38") + v("B36"), v("B40"), aba, "B40")
    confere("Área à colher total", a_colher + cargill["a_colher_ha"], v("D40"), aba, "D40")
    confere("Área de tratos total", tratos + cargill["tratos_ha"], v("E40"), aba, "E40")

    # 4. Resumo (balanço por fórmula)
    balanco = {}
    for lin, prod, nec in ((50, "46-00-00", nec46), (51, "16-00-20", nec16)):
        confere(f"Necessidade {prod} no resumo", nec, v(f"B{lin}"), aba, f"B{lin}")
        est, receber = v(f"C{lin}"), v(f"D{lin}")
        confere(f"Disponível {prod}", est + receber, v(f"E{lin}"), aba, f"E{lin}")
        confere(f"Diferença {prod}", est + receber - nec, v(f"F{lin}"), aba, f"F{lin}")
        balanco[prod] = {"necessidade_t": nec, "estoque_t": est, "a_receber_t": receber,
                         "disponivel_t": est + receber, "saldo_t": est + receber - nec}
    compra16 = -balanco["16-00-20"]["saldo_t"]
    confere("Necessidade de aquisição 16-00-20", compra16, v("B56"), aba, "B56")

    # Estoque em 16/10 (quadro I47:K58)
    estoque = []
    for lin in range(49, 59):
        prod, t = produto(v(f"I{lin}")), v(f"J{lin}")
        form = f(f"K{lin}")
        dose_t_ha = float(form.split("/")[1])
        confere(f"Área equivalente {prod}", t / dose_t_ha, v(f"K{lin}"), aba, f"K{lin}")
        grupo = "Socaria" if 49 <= lin <= 53 else "Plantio" if 54 <= lin <= 57 else "Não classificado"
        estoque.append({"produto": prod, "estoque_t": t, "dose_t_ha": dose_t_ha,
                        "area_eq_ha": t / dose_t_ha, "grupo": grupo, "linha": lin})
    confere("Área equivalente SOCARIA", sum(e["area_eq_ha"] for e in estoque if e["grupo"] == "Socaria"), v("K60"), aba, "K60")
    confere("Área equivalente PLANTIO", sum(e["area_eq_ha"] for e in estoque if e["grupo"] == "Plantio"), v("K61"), aba, "K61")
    confere("Estoque 16-00-20 (quadro × resumo)", v("J49"), v("C51"), aba, "J49/C51")
    confere("Estoque 46-00-00 (quadro × resumo)", v("J50"), v("C50"), aba, "J50/C50")

    return {
        "aba": aba, "nome": NOME_CENARIO[chave],
        "data_base": data_base.date().isoformat(), "data_limite_localizada": data_limite.date().isoformat(),
        "termino_safra": termino.date().isoformat(),
        "dias_limite_localizada": (data_limite - data_base).days, "dias_safra": (termino - data_base).days,
        "ritmo_localizada_ha_dia": ritmo, "dias_k_vinhaca": dias_k, "area_localizada_estimada_ha": area_loc_est,
        "realizado": real, "consumo_realizado": consumo_real, "quadro_localizada": quadro_localizada,
        "cargill": cargill,
        "crv": {"colhida_ha": v("B38"), "plantio_ha": plantio, "a_colher_ha": a_colher,
                "a_colher_partes": [11593, 9354, 8779], "tratos_ha": tratos, "realizado_ha": realizado_crv,
                "bisada_ha": bisada_ha, "tch_bisada": tch, "bisada_t": bisada_ha * tch,
                "a_tratar_localizada_ha": loc_crv, "a_tratar_mineral_ha": mineral,
                "h38_digitado": h38_digitado},
        "total_a_tratar_ha": total_tratar,
        "doses_kg_ha": {"46-00-00": dose46, "16-00-20": dose16},
        "balanco": balanco, "compra_16_t": compra16, "estoque": estoque,
    }


def derivados(c):
    """Indicadores calculados a partir dos números da aba (nada de fora da planilha)."""
    b16, b46 = c["balanco"]["16-00-20"], c["balanco"]["46-00-00"]
    nec_total = b16["necessidade_t"] + b46["necessidade_t"]
    dias = c["dias_safra"]
    base = c["data_base"]
    from datetime import date
    d0 = date.fromisoformat(base)
    diario16 = b16["necessidade_t"] / dias  # premissa ilustrativa: ritmo uniforme até o fim da safra
    dias_estoque = b16["estoque_t"] / diario16
    dias_disp = b16["disponivel_t"] / diario16
    tratos_total = c["crv"]["tratos_ha"] + c["cargill"]["tratos_ha"]
    realizado_total = c["crv"]["realizado_ha"] + c["cargill"]["realizado_ha"]
    n16, _, k16 = TEOR["16-00-20"]
    n46 = TEOR["46-00-00"][0]
    pontos = list(range(0, dias + 1, 5))  # passo constante: o eixo de datas não distorce a reta
    estoque = c["estoque"]
    return {
        "necessidade_total_t": nec_total,
        "participacao_16": b16["necessidade_t"] / nec_total,
        "cobertura_16_estoque": b16["estoque_t"] / b16["necessidade_t"],
        "cobertura_16_disponivel": b16["disponivel_t"] / b16["necessidade_t"],
        "deficit_16_ha": c["compra_16_t"] / (c["doses_kg_ha"]["16-00-20"] / 1000),
        "a_receber_total_t": b16["a_receber_t"] + b46["a_receber_t"],
        "nutrientes_t": {"N": b16["necessidade_t"] * n16 + b46["necessidade_t"] * n46,
                         "K2O": b16["necessidade_t"] * k16},
        "nutrientes_por_ha": {"16-00-20": {"N": c["doses_kg_ha"]["16-00-20"] * n16, "K2O": c["doses_kg_ha"]["16-00-20"] * k16},
                              "46-00-00": {"N": c["doses_kg_ha"]["46-00-00"] * n46, "K2O": 0}},
        "deficit_nutrientes_t": {"N": c["compra_16_t"] * n16, "K2O": c["compra_16_t"] * k16},
        "excedente_46_N_t": b46["saldo_t"] * n46,
        "area": {"tratos_total_ha": tratos_total, "realizado_total_ha": realizado_total,
                 "pct_realizado_total": realizado_total / tratos_total,
                 "pct_realizado_crv": c["crv"]["realizado_ha"] / c["crv"]["tratos_ha"],
                 "pct_realizado_cargill": c["cargill"]["realizado_ha"] / c["cargill"]["tratos_ha"],
                 "pct_mineral_no_restante": c["crv"]["a_tratar_mineral_ha"] / c["total_a_tratar_ha"],
                 "ritmo_para_fechar_na_safra_ha_dia": c["total_a_tratar_ha"] / dias},
        "curva_16": {
            "premissa": "aplicação uniforme do 16-00-20 entre a data-base e o término da safra",
            "consumo_diario_t": diario16,
            "dias_cobertura_estoque": dias_estoque,
            "dias_cobertura_disponivel": dias_disp,
            "data_esgota_estoque": (d0 + timedelta(days=int(dias_estoque))).isoformat(),
            "data_esgota_disponivel": (d0 + timedelta(days=int(round(dias_disp)))).isoformat(),
            "pontos": [{"dia": d, "data": (d0 + timedelta(days=d)).isoformat(),
                        "necessidade_acum_t": diario16 * d} for d in pontos],
        },
        "estoque_total_t": sum(e["estoque_t"] for e in estoque),
        "estoque_por_grupo_t": {g: sum(e["estoque_t"] for e in estoque if e["grupo"] == g)
                                for g in ("Socaria", "Plantio", "Não classificado")},
        "area_eq_por_grupo_ha": {g: sum(e["area_eq_ha"] for e in estoque if e["grupo"] == g)
                                 for g in ("Socaria", "Plantio", "Não classificado")},
        "sens_preco_por_100": {"compra_16_R$": c["compra_16_t"] * 100,
                               "custo_ha_16_R$": c["doses_kg_ha"]["16-00-20"] / 1000 * 100,
                               "excedente_46_R$": max(b46["saldo_t"], 0) * 100},
    }


def sensibilidade_dias(c, dias_lista):
    """Compra de 16-00-20 conforme os dias de vinhaça localizada com K adequado (mesma lógica da aba A)."""
    out = []
    for d in dias_lista:
        loc = max(0.0, d * c["ritmo_localizada_ha_dia"] - c["cargill"]["a_tratar_ha"])
        mineral = c["crv"]["tratos_ha"] - c["crv"]["realizado_ha"] - c["crv"]["bisada_ha"] - loc
        nec16 = mineral * c["doses_kg_ha"]["16-00-20"] / 1000
        nec46 = loc * c["doses_kg_ha"]["46-00-00"] / 1000
        b16 = c["balanco"]["16-00-20"]
        b46 = c["balanco"]["46-00-00"]
        out.append({"dias": d, "localizada_crv_ha": loc, "mineral_ha": mineral,
                    "necessidade_16_t": nec16, "necessidade_46_t": nec46,
                    "compra_16_t": max(0.0, nec16 - b16["disponivel_t"]),
                    "excedente_46_t": b46["disponivel_t"] - nec46})
    return out


def main():
    origem = Path(sys.argv[1])
    wb = openpyxl.load_workbook(origem, data_only=True)
    wf = openpyxl.load_workbook(origem, data_only=False)
    cen = {k: le_aba(wb[a], wf[a], k) for k, a in ABAS.items()}
    for k in cen:
        cen[k]["derivados"] = derivados(cen[k])
    sens = sensibilidade_dias(cen["A"], [0, 10, 20, 30, 45])
    # A lógica derivada precisa reproduzir os dois cenários da planilha.
    s20 = next(s for s in sens if s["dias"] == 20)
    confere("Sensibilidade 20 dias = cenário A", s20["compra_16_t"], cen["A"]["compra_16_t"], ABAS["A"], "B56")
    s0 = next(s for s in sens if s["dias"] == 0)
    confere("Sensibilidade 0 dias = cenário B", s0["compra_16_t"], cen["B"]["compra_16_t"], ABAS["B"], "B56")

    A, B = cen["A"], cen["B"]
    ql = A["quadro_localizada"]
    inconsistencias = [
        {"id": 1, "gravidade": "Alta", "onde": "Ambas as abas, linhas 26 a 29",
         "achado": f"Quadro de vinhaça localizada não fecha: total B26 = {ql['total_ha']:.2f} ha, mas B27+B28+B29 = {sum(ql['partes_ha']):.2f} ha "
                   f"(diferença {ql['total_ha'] - sum(ql['partes_ha']):.2f} ha). B27+B28 = {ql['partes_ha'][0] + ql['partes_ha'][1]:.2f} ha também difere de B22 = {A['realizado']['crv_localizada_ha']:.2f} ha. As linhas 27 a 29 não têm descrição na coluna A.",
         "efeito": "Não afeta a necessidade restante (calculada pelas áreas da seção 3), mas impede conciliar o consumo realizado por produto."},
        {"id": 2, "gravidade": "Média", "onde": f"{ABAS['B']}!H38",
         "achado": "H38 está digitado como 0; na aba A é a fórmula =C14-H36. Com C13 = 0 a fórmula daria -446,63 ha, por isso o valor fixo está certo neste cenário, mas não se atualiza se C13 mudar.",
         "efeito": "Se alguém alterar os dias de K na aba B, a necessidade não acompanha."},
        {"id": 3, "gravidade": "Baixa", "onde": "Ambas as abas, H36",
         "achado": "H36 = 1200 - F36 usa 1.200 digitado em vez de referenciar E36.", "efeito": "Risco de divergência se a área de tratos da Cargill mudar."},
        {"id": 4, "gravidade": "Média", "onde": "Ambas as abas, D38",
         "achado": "Área à colher da CRV = 11.593 + 9.354 + 8.779 ha digitados na fórmula, sem identificar a que corresponde cada parcela.",
         "efeito": "Sem rastreabilidade por bloco/fazenda; impede a visão de necessidade por área produtiva."},
        {"id": 5, "gravidade": "Média", "onde": "Ambas as abas, K58 × K60/K61",
         "achado": f"Oxyfertil (142,19 t; 284,38 ha equivalentes) não entra no subtotal SOCARIA nem no PLANTIO.",
         "efeito": "Estoque sem destino definido; pode reduzir compra se tiver uso agronômico validado."},
        {"id": 6, "gravidade": "Baixa", "onde": "Colunas C, I e A dos quadros",
         "achado": "A mesma fórmula aparece como '16-00-20.' e '16-00-20' (e '46-00-00.' e '46-00-00').",
         "efeito": "Quebra PROCV/tabela dinâmica que consolidem por produto."},
        {"id": 7, "gravidade": "Média", "onde": "Ambas as abas, quadro de estoque (I47:K58)",
         "achado": "21-00-00 (126 t), 13-00-34 (4 t), 25-00-00 (0,85 t) e Oxyfertil (142,19 t) estão em estoque mas fora do balanço de necessidade (B50:F51).",
         "efeito": "O saldo a comprar considera só 16-00-20 e 46-00-00."},
        {"id": 8, "gravidade": "Média", "onde": "Ambas as abas, A42",
         "achado": "Plantio projetado de 2.400 ha (1.400 fundação + 1.000 renovação) sem necessidade de fertilizante calculada; o estoque de fórmulas de plantio cobre 1.397,18 ha pela dose implícita de 600 kg/ha.",
         "efeito": "Possível necessidade adicional de fórmulas de plantio fora deste balanço — validar escopo e época."},
        {"id": 9, "gravidade": "Alta", "onde": "Planilha inteira",
         "achado": "Não há preços, datas de entrega do 'comprado a receber', prazo de fornecimento, distribuição mensal da demanda nem ritmo de aplicação do fertilizante mineral.",
         "efeito": "Orçamento, curva real de abastecimento e aderência ao cronograma não podem ser calculados."},
        {"id": 10, "gravidade": "Baixa", "onde": f"{ABAS['B']}, seções",
         "achado": "Seções numeradas '1. Projeção' e '1. Áreas adubadas' (numeração repetida); grafias 'Fertlizante' e 'Rítmo'.", "efeito": "Apenas apresentação."},
    ]
    out = {
        "fonte": {"arquivo": origem.name, "abas": ABAS},
        "usina": {"nome": "Usina CRV Industrial", "departamento": "Departamento Agrícola",
                  "unidade": "Unidade Capinópolis - MG"},
        "plantio_projetado_ha": {"fundacao": 1400, "renovacao": 1000, "fonte": "nota A42"},
        "cenarios": cen,
        "sensibilidade_dias": sens,
        "checks": checks,
        "inconsistencias": inconsistencias,
    }
    destino = Path(__file__).with_name("dados_consolidados.json")
    destino.write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"{len(checks)} conferências OK -> {destino.name}")
    for k in ("A", "B"):
        c = cen[k]
        d = c["derivados"]
        print(f"{k}: a tratar {c['total_a_tratar_ha']:.2f} ha | nec 16 {c['balanco']['16-00-20']['necessidade_t']:.3f} t | "
              f"nec 46 {c['balanco']['46-00-00']['necessidade_t']:.3f} t | compra 16 {c['compra_16_t']:.3f} t | "
              f"excedente 46 {c['balanco']['46-00-00']['saldo_t']:.3f} t | cobertura {d['cobertura_16_disponivel']:.4f} | "
              f"esgota estoque {d['curva_16']['data_esgota_estoque']} disp {d['curva_16']['data_esgota_disponivel']}")
    for s in sens:
        print(f"  {s['dias']:>2} dias: localizada CRV {s['localizada_crv_ha']:.2f} ha | compra 16 {s['compra_16_t']:.2f} t")


if __name__ == "__main__":
    main()
